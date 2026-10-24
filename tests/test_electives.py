from io import BytesIO
import unittest

from openpyxl import Workbook

from app.electives import (
    analyze_elective_subjects,
    parse_elective_excel,
    preview_elective_excel,
    suggest_elective_mapping,
)


class ElectiveExcelTest(unittest.TestCase):
    def _workbook_bytes(self) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Listado de electivas"])
        sheet.append([])
        sheet.append(
            [
                "Modalidad",
                "APELLIDOS",
                "NOMBRES",
                "DNI",
                "Codigo Moodle",
                "ASIGNATURA ELECTIVA",
                "MAIL",
            ]
        )
        sheet.append(["A DISTANCIA", "Lopez", "Ana", "401", "401", "Administracion Rural", "ana@example.com"])
        sheet.append(["PRESENCIAL", "Perez", "Beto", "402", "401", "Administracion Rural", ""])
        sheet.append(["PRESENCIAL", "Perez", "Beto", "402", "401", "Administracion Rural", ""])
        sheet.append(["PRESENCIAL", "Lopez", "Ana", "401", "401", "Administracion Rural", ""])
        stream = BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    def test_detects_header_after_intro_rows_and_maps_p_d(self):
        content = self._workbook_bytes()
        columns, sample = preview_elective_excel(content)
        self.assertIn("DNI", columns)
        self.assertEqual(sample[0]["APELLIDOS"], "Lopez")

        mapping = suggest_elective_mapping(columns)
        self.assertEqual(mapping["document"], "DNI")
        self.assertEqual(mapping["modality"], "Modalidad")
        self.assertEqual(mapping["course"], "Codigo Moodle")
        self.assertEqual(mapping["subject"], "ASIGNATURA ELECTIVA")

        rows, issues, total_rows, duplicates = parse_elective_excel(
            content,
            "eco.unrc.edu.ar",
            mapping,
        )
        self.assertEqual(total_rows, 4)
        self.assertEqual(duplicates, 1)
        self.assertEqual(len(rows), 2)
        self.assertEqual({row.group_name for row in rows}, {"D", "P"})
        self.assertEqual(rows[0].upn, "a401@eco.unrc.edu.ar")
        self.assertEqual(len(issues), 1)
        self.assertIn("modalidades distintas", issues[0].message)

    def test_infers_course_by_subject_and_accepts_a_default_group(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Nombre", "Apellidos"])
        sheet.append(["", "", "", "", "Electivas segundo cuatrimestre"])
        sheet.append([])
        sheet.append(["", "", "CARRERA", "APELLIDO Y NOMBRES", "DNI", "ASIGNTURA ELECTIVA"])
        sheet.append(["Ana", "Lopez", "C.P.", "Lopez Ana", "401", "Administracion Rural"])
        sheet.append(["Beto", "Perez", "C.P.", "Perez Beto", "402", "Administracion Rural (401)"])
        sheet.append(["Carla", "Diaz", "Lic.Econ.", "Diaz Carla", "403", "Seminario circular (345)"])
        stream = BytesIO()
        workbook.save(stream)
        content = stream.getvalue()

        columns, _ = preview_elective_excel(content)
        mapping = suggest_elective_mapping(columns)
        subjects = analyze_elective_subjects(content, mapping["subject"])
        destinations = {item["label"]: item for item in subjects}

        self.assertEqual(mapping["first_name"], "Nombre")
        self.assertEqual(mapping["last_name"], "Apellidos")
        self.assertNotIn("modality", mapping)
        self.assertEqual(destinations["Administracion Rural"]["row_count"], 2)
        self.assertEqual(destinations["Administracion Rural"]["suggested_course"], "401")
        self.assertEqual(destinations["Seminario circular"]["suggested_course"], "345")

        course_map = {
            str(item["key"]): str(item["suggested_course"])
            for item in subjects
        }
        rows, issues, total_rows, duplicates = parse_elective_excel(
            content,
            "eco.unrc.edu.ar",
            mapping,
            course_map,
            "P",
        )
        self.assertEqual(total_rows, 3)
        self.assertEqual(duplicates, 0)
        self.assertEqual(issues, [])
        self.assertEqual([row.course_shortname for row in rows], ["401", "401", "345"])
        self.assertEqual({row.group_name for row in rows}, {"P"})

    def test_maps_and_splits_apellido_y_nombres(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Titulo"])
        sheet.append([])
        sheet.append(["Modalidad", "APELLIDO Y NOMBRES", "DNI", "ASIGNTURA ELECTIVA"])
        sheet.append(["Presencial", "OLMOS Juan Ignacio", "43610388", "Administracion Rural (401)"])
        stream = BytesIO()
        workbook.save(stream)
        content = stream.getvalue()

        columns, _ = preview_elective_excel(content)
        mapping = suggest_elective_mapping(columns)
        self.assertEqual(mapping["full_name"], "APELLIDO Y NOMBRES")

        subjects = analyze_elective_subjects(content, mapping["subject"])
        rows, issues, _, _ = parse_elective_excel(
            content,
            "eco.unrc.edu.ar",
            mapping,
            {str(item["key"]): str(item["suggested_course"]) for item in subjects},
        )
        self.assertEqual(issues, [])
        self.assertEqual(rows[0].first_name, "Juan Ignacio")
        self.assertEqual(rows[0].last_name, "Olmos")


if __name__ == "__main__":
    unittest.main()
