import io
import re
import unittest
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook

from app import auditing as audit
from app.models import AuditDeletion, AuditDeletionRow
from app.services.graph import ServiceResult
from app.services.moodle import MoodleResult
import test_auditing as fixtures


class PreinscriptoAuditTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        fixtures.AuditTest.setUpClass()

    def setUp(self):
        self.case = fixtures.AuditTest()
        self.case.setUp()
        self.addCleanup(self.case.doCleanups)

    def test_any_course_or_unknown_courses_prevents_candidacy(self):
        for change in ({"moodle_cursos_cantidad": 1}, {"moodle_cursos_cantidad": 15},
                       {"moodle_cursos_estado": "error"}, {"moodle_cursos_estado": "pendiente"},
                       {"moodle_cursos_cantidad": None}, {"moodle_estado": "error"}):
            with self.subTest(change=change):
                self.assertFalse(audit.candidate_assessment({**self.case.p, **change}, self.case.run)[0])

    def test_confirmed_missing_moodle_is_allowed_but_missing_both_is_not(self):
        p = {**self.case.p, "moodle_estado": "sin_usuario", "moodle_id": None,
             "moodle_cursos_estado": "sin_usuario", "moodle_cursos_cantidad": None}
        self.assertEqual(audit.eligibility(p, self.case.run)[0], "preinscripto")
        p["microsoft_estado"] = "sin_usuario"
        self.assertFalse(audit.eligibility(p, self.case.run)[0])

    def test_preinscripto_requires_no_teams_and_verified_180_day_inactivity(self):
        with_teams = {**self.case.p, "teams_cantidad": 10}
        self.assertFalse(audit.candidate_assessment(with_teams, self.case.run)[0])

        for activity in ("Actividad registrada (180 dias)", "Actividad historica fuera de ventana", "Datos incompletos"):
            with self.subTest(activity=activity):
                p = {**self.case.p, "actividad": activity}
                self.assertEqual(audit.candidate_assessment(p, self.case.run)[0], "preinscripto")
                self.assertFalse(audit.eligibility(p, self.case.run)[0])

    def test_new_course_or_failed_live_read_blocks_both_deletions(self):
        for result in (MoodleResult("ok", data=[{"id": 1, "shortname": "2024-other-career"}]),
                       MoodleResult("error"), MoodleResult("ok", data=None), MoodleResult("ok", data={})):
            with self.subTest(result=result):
                m = MagicMock()
                m.get_user_courses.return_value = result
                batch = self.case.create_batch()
                m, g = self.case.execute(batch, moodle=m)
                m.get_user_courses.assert_called_once_with(7)
                m.delete_user_by_id.assert_not_called()
                g.delete_user.assert_not_called()
                with self.case.session() as db:
                    row = db.query(AuditDeletionRow).filter_by(deletion_id=batch).one()
                    self.assertIn("baja bloqueada en ambos sistemas", row.message)
                    self.assertEqual(row.moodle_status, "pending")
                    self.assertEqual(row.microsoft_status, "pending")

    def test_missing_moodle_can_delete_microsoft_without_course_call(self):
        m, g = MagicMock(), MagicMock()
        m.get_user_by_username.return_value = MoodleResult("missing")
        m.get_user_by_username(self.case.p["upn"])
        g.delete_user.return_value = ServiceResult("deleted")
        self.case.execute(self.case.create_batch(), moodle=m, graph=g)
        m.get_user_courses.assert_not_called()
        m.delete_user_by_id.assert_not_called()
        g.delete_user.assert_called_once_with("ms-test-id")

    def test_preinscripto_live_team_or_recent_activity_blocks_both_deletions(self):
        recent_usage = fixtures.usage(self.case.p)
        recent_usage[(self.case.p["dni"], "Teams")][0]["Call Count"] = "1"
        m, g = self.case.execute(self.case.create_batch(), report_usage=recent_usage)
        m.delete_user_by_id.assert_not_called()
        g.delete_user.assert_not_called()

        m, g = self.case.execute(self.case.create_batch(), teams=[{"id": "team-now"}])
        m.delete_user_by_id.assert_not_called()
        g.delete_user.assert_not_called()

    def add_protected_people(self):
        enrolled = {**self.case.p, "dni": "99888888", "upn": "a99888888@example.invalid",
                    "moodle_cursos_cantidad": 5}
        unknown = {**self.case.p, "dni": "99888999", "upn": "a99888999@example.invalid",
                   "moodle_cursos_estado": "error"}
        self.case.create_snapshot([enrolled, unknown])
        return enrolled, unknown

    def test_filters_excel_checkboxes_and_bulk_selection_exclude_protected_people(self):
        enrolled, unknown = self.add_protected_people()
        client = self.case.client()
        client.get("/test-session")
        with patch.object(audit, "MoodleClient") as moodle, patch.object(audit, "GraphClient") as graph:
            page = client.get("/auditoria?view=preinscriptos")
            self.assertEqual(re.findall(r'name="dni" value="(\d+)"', page.text), [self.case.p["dni"]])
            self.assertIn("Conservar: tiene matriculas", page.text)
            self.assertIn("Cursos Moodle sin verificar", page.text)
            token = re.search(r'name="token" value="([^"]+)"', page.text)[1]
            people = audit.people_for_run(self.case.run)
            self.assertEqual([p["dni"] for p in audit.filter_people(people, "pre_candidates")], [self.case.p["dni"]])
            self.assertEqual([p["dni"] for p in audit.filter_people(people, "pre_enrolled")], [enrolled["dni"]])
            protected_page = client.get("/auditoria?view=pre_enrolled")
            self.assertNotIn('name="dni"', protected_page.text)
            workbook = load_workbook(io.BytesIO(client.get(f"/auditoria/{self.case.run.id}/excel").content))
            self.assertEqual(workbook["Preins_sin_cursos_teams"].max_row, 2)
            self.assertEqual(workbook["Preins_sin_cursos_teams"]["A2"].value, self.case.p["dni"])
            self.assertEqual(workbook["Preinscriptos_con_cursos"]["A2"].value, enrolled["dni"])
            self.assertEqual(workbook["Bajas_preinscriptos"].max_row, 2)
            self.assertEqual(workbook["Bajas_preinscriptos"]["A2"].value, self.case.p["dni"])
            self.assertIn("Actividad Microsoft 365/Teams (180 dias)", [cell.value for cell in workbook["Personas"][1]])
            self.assertIn("Actividad 180 dias", [row[0].value for row in workbook["Alcance"].iter_rows(min_row=2)])
            for p in (enrolled, unknown):
                client.post(f"/auditoria/{self.case.run.id}/bajas/previa", data={"token": token, "dni": p["dni"]})
            with self.case.session() as db:
                self.assertEqual(db.query(AuditDeletion).count(), 0)
            client.post(f"/auditoria/{self.case.run.id}/bajas/previa", data={
                "token": token, "all_filtered": "yes", "view": "preinscriptos"})
            with self.case.session() as db:
                self.assertEqual([r.dni for r in db.query(AuditDeletionRow).all()], [self.case.p["dni"]])
            moodle.assert_not_called()
            graph.assert_not_called()

    def test_legacy_preview_with_courses_is_blocked_in_ui_confirm_and_worker(self):
        p = {**self.case.p, "moodle_cursos_cantidad": 5}
        batch = self.case.create_batch(p)
        client = self.case.client()
        client.get("/test-session")
        page = client.get(f"/auditoria/bajas/{batch}")
        self.assertIn("no cumplen el criterio vigente", page.text)
        self.assertNotIn('name="confirmation"', page.text)
        home = client.get("/auditoria")
        token = re.search(r'name="token" value="([^"]+)"', home.text)[1]
        with patch.object(audit, "deletion_job") as worker:
            response = client.post(f"/auditoria/bajas/{batch}/confirmar", data={
                "token": token, "confirmation": "ELIMINAR", "understood": "yes", "password": "test-password-only"})
            self.assertIn("no cumplen el criterio vigente", response.text)
            worker.assert_not_called()
        m, g = self.case.execute(batch, p=p)
        m.delete_user_by_id.assert_not_called()
        g.delete_user.assert_not_called()


if __name__ == "__main__":
    unittest.main()
