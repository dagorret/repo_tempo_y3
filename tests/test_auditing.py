from contextlib import closing
from datetime import date, datetime, timedelta
import io
import csv
import json
from pathlib import Path
import re
import tempfile
import unittest
from unittest.mock import MagicMock, patch

from fastapi import FastAPI, Request
from fastapi.templating import Jinja2Templates
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from starlette.middleware.sessions import SessionMiddleware

from app import auditing as audit, audit_sources as sources, audit_reports as reports
from app.audit_routes import create_audit_router
from app.config import Settings
from app.db import Base, get_db
from app.models import (AuditRun, AuditDeletion, AuditDeletionRow, AuditLease, ManagedUser,
                        ManagedEnrollment, MatriculationState, LocalUser)
from app.security import hash_password
from app.services.graph import ServiceResult
from app.services.moodle import MoodleClient, MoodleResult


def person(state="preinscriptos"):
    return dict(dni="99888777", upn="a99888777@example.invalid", nombre="Alumno", apellido="Prueba",
                estados_sial=state, moodle_estado="encontrado", moodle_id=7, moodle_cursos_estado="consultado",
                moodle_cursos_cantidad=0, moodle_primer_acceso=0, moodle_ultimo_acceso=0,
                microsoft_estado="encontrado", microsoft_id="ms-test-id", microsoft_creada="2025-01-01T00:00:00Z",
                teams_estado="consultado", teams_cantidad=0, actividad="Sin actividad informada (180 dias)",
                ultima_actividad=None, inscripciones=[{"fecha_inscripcion": "10/12/2025 12:00:00"}],
                candidate="preinscripto" if state == "preinscriptos" else "aspirante")


def usage(p, when=None):
    refreshed = (when or date.today()).isoformat()
    office = {"Report Refresh Date": refreshed, "Is Deleted": "False", **{k: "" for k in reports.ACTIVITY_FIELDS}}
    teams = {"Report Refresh Date": refreshed, "Is Deleted": "False", "Last Activity Date": "", "Report Period": "180",
             "Has Other Action": "False", **{k: "0" for k in reports.TEAM_METRICS}}
    return {(p["dni"], "Microsoft365"): [office], (p["dni"], "Teams"): [teams]}


class AuditTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.password_hash = hash_password("test-password-only")

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.settings = Settings(_env_file=None, audit_data_dir=str(Path(self.temp.name) / "snapshots"),
            crm_database_url="sqlite://", crm_dry_run=False, sial_username="test", sial_password="secret-test",
            azure_tenant_id="test", azure_client_id="test", azure_client_secret="secret-test",
            moodle_base_url="https://moodle.example.invalid", moodle_token="secret-test", crm_domain="example.invalid")
        self.engine = create_engine("sqlite:///" + str(Path(self.temp.name) / "crm.sqlite3"), connect_args={"check_same_thread": False})
        self.addCleanup(self.engine.dispose)
        Base.metadata.create_all(self.engine)
        self.session = sessionmaker(self.engine, expire_on_commit=False)
        for module in (audit, "app.audit_routes"):
            patcher = patch.object(module, "get_settings", return_value=self.settings) if not isinstance(module, str) else patch(module + ".get_settings", return_value=self.settings)
            patcher.start()
            self.addCleanup(patcher.stop)
        self.p = person()
        with self.session() as db:
            self.run = AuditRun(snapshot_key="a" * 32, campaign_start="2025-12-01", campaign_end="2026-03-01",
                environment=audit.environment_key(self.settings), actor="test", status="completed", progress=100)
            db.add(self.run)
            db.add(LocalUser(username="test", password_hash=self.password_hash, active=True))
            db.commit()
        self.create_snapshot([self.p])

    def create_snapshot(self, people):
        with closing(sources.open_db(audit.snapshot_path(self.run, self.settings).parent)) as db:
            reports.init_reports(db)
            columns = {r[1] for r in db.execute("PRAGMA table_info(personas)")}
            with db:
                for p in people:
                    keys = [k for k in p if k in columns]
                    db.execute("INSERT INTO revisiones(dni) VALUES (?)", (p["dni"],))
                    db.execute(f"INSERT INTO personas ({','.join(keys)}) VALUES ({','.join('?' for _ in keys)})", [p[k] for k in keys])
                    db.execute("INSERT INTO sial_inscripciones VALUES (?,?,?,?,?,?,?,?)", (p["dni"], p["estados_sial"], "1", "Carrera prueba", "1", "COMISION 1", "10/12/2025 12:00:00", sources.now()))
                    db.execute("INSERT INTO evaluacion_actividad VALUES (?,?,?,?,?,?)", (p["dni"], p["actividad"], None, "Sin acciones", "", sources.now()))

    def create_batch(self, p=None):
        p = p or self.p
        with self.session() as db:
            batch = AuditDeletion(run_id=self.run.id, actor="test")
            db.add(batch)
            db.flush()
            row = AuditDeletionRow(deletion_id=batch.id, dni=p["dni"], snapshot_json=json.dumps(p))
            db.add(row)
            db.commit()
            return batch.id

    def execute(self, batch_id, p=None, moodle=None, graph=None, states=None, report_usage=None, teams=None):
        p = p or self.p
        moodle = moodle or MagicMock()
        graph = graph or MagicMock()
        if isinstance(moodle.get_user_courses.return_value, MagicMock):
            moodle.get_user_courses.return_value = MoodleResult("ok", data=[])
        if not moodle.get_user_by_username.called:
            moodle.get_user_by_username.return_value = MoodleResult("ok", data={"id": 7, "username": p["upn"]})
        if not graph.get_user.called:
            graph.get_user.return_value = ServiceResult("exists", data={"id": "ms-test-id", "userPrincipalName": p["upn"]})
        reader = MagicMock()
        reader.get.return_value.status_code = 200
        reader.list_pages.return_value = teams if teams is not None else []
        state_data = states if states is not None else {p["dni"]: {"states": {p["estados_sial"]}, "careers": [{"fecha_inscri": "10/12/2025 12:00:00"}]}}
        with self.session() as db:
            audit.acquire_lease(db, "test-job")
            db.commit()
        with patch.object(audit, "MoodleClient", return_value=moodle), patch.object(audit, "GraphClient", return_value=graph), \
             patch.object(audit, "fresh_states", return_value=state_data), patch.object(audit, "fresh_usage", return_value=report_usage if report_usage is not None else usage(p)), \
             patch.object(sources, "GraphReadOnly", return_value=reader):
            audit.deletion_job(batch_id, "test-job", self.session, self.settings)
        return moodle, graph

    def test_eligibility_closed_campaign_and_activity_policy(self):
        self.p.update(moodle_cursos_cantidad=12, teams_cantidad=2, actividad="Actividad registrada (180 dias)")
        self.assertFalse(audit.eligibility(self.p, self.run)[0])
        self.p.update(moodle_cursos_cantidad=0, teams_cantidad=0, actividad="Sin actividad informada (180 dias)")
        self.assertEqual(audit.eligibility(self.p, self.run)[0], "preinscripto")
        self.p["actividad"] = "Actividad registrada (180 dias)"
        self.assertFalse(audit.eligibility(self.p, self.run)[0])
        self.run.campaign_end = ""
        self.assertFalse(audit.eligibility(self.p, self.run)[0])

    def test_new_campaign_unknown_dates_and_overlapping_states_blocked(self):
        for updates in ({"inscripciones": []}, {"inscripciones": [{"fecha_inscri": "01/12/2026"}]},
                        {"estados_sial": "preinscriptos; ingresantes"}, {"moodle_estado": "error"}):
            self.assertFalse(audit.eligibility({**self.p, **updates}, self.run)[0])

    def test_aspirante_old_moodle_access_is_allowed_but_unknown_is_not(self):
        p = person("aspirantes")
        p["moodle_ultimo_acceso"] = 1700000000
        self.assertEqual(audit.eligibility(p, self.run)[0], "aspirante")
        for updates in ({"moodle_cursos_estado": "error"}, {"teams_cantidad": 1}, {"actividad": "Datos incompletos"}):
            self.assertFalse(audit.eligibility({**p, **updates}, self.run)[0])

    def test_activity_checks_missing_stale_recent_and_new_account(self):
        p = person("aspirantes")
        audit.verify_inactivity(p, usage(p))
        for rows in ({}, usage(p, date.today() - timedelta(days=8))):
            with self.assertRaises(ValueError):
                audit.verify_inactivity(p, rows)
        rows = usage(p)
        rows[(p["dni"], "Teams")][0]["Call Count"] = "1"
        with self.assertRaises(ValueError):
            audit.verify_inactivity(p, rows)
        p["microsoft_creada"] = date.today().isoformat()
        with self.assertRaises(ValueError):
            audit.verify_inactivity(p, usage(p))

    def test_moodle_malformed_response_is_not_missing_or_empty(self):
        client = MoodleClient(self.settings)
        with patch.object(client, "_read_call", return_value=MoodleResult("ok", data={})), patch.object(client, "_call") as write:
            self.assertEqual(client.get_user_courses(7).status, "error")
            self.assertEqual(client.get_user_by_username(self.p["upn"]).status, "error")
            write.assert_not_called()

    def test_deletion_revalidates_and_updates_local_registry(self):
        batch = self.create_batch()
        with self.session() as db:
            user = ManagedUser(upn=self.p["upn"], document=self.p["dni"], person_type="student")
            db.add(user)
            db.flush()
            db.add(ManagedEnrollment(user_id=user.id, course_shortname="TEST"))
            db.add(MatriculationState(document=self.p["dni"], upn=self.p["upn"], course_shortname="TEST"))
            db.commit()
        m, g = MagicMock(), MagicMock()
        m.delete_user_by_id.return_value = MoodleResult("deleted")
        g.delete_user.return_value = ServiceResult("deleted")
        self.execute(batch, moodle=m, graph=g)
        m.delete_user_by_id.assert_called_once_with(7)
        g.delete_user.assert_called_once_with("ms-test-id")
        with self.session() as db:
            self.assertEqual(db.get(AuditDeletion, batch).status, "completed")
            self.assertFalse(db.query(ManagedUser).one().active)
            self.assertFalse(db.query(ManagedEnrollment).one().active)
            self.assertFalse(db.query(MatriculationState).one().active)

    def test_retry_only_failed_system(self):
        batch = self.create_batch()
        m, g = MagicMock(), MagicMock()
        m.delete_user_by_id.return_value = MoodleResult("deleted")
        g.delete_user.return_value = ServiceResult("error", "ServiceUnavailable")
        self.execute(batch, moodle=m, graph=g)
        m.get_user_by_username.return_value = MoodleResult("missing")
        g.delete_user.return_value = ServiceResult("deleted")
        self.execute(batch, moodle=m, graph=g)
        self.assertEqual(m.delete_user_by_id.call_count, 1)
        self.assertEqual(g.delete_user.call_count, 2)
        with self.session() as db:
            self.assertEqual(db.get(AuditDeletion, batch).status, "completed")

    def test_changed_sial_and_missing_sial_block_all_writes(self):
        for states in ({}, {self.p["dni"]: {"states": {"ingresantes"}, "careers": []}}):
            m, g = self.execute(self.create_batch(), states=states)
            m.delete_user_by_id.assert_not_called()
            g.delete_user.assert_not_called()

    def test_recreated_identity_blocks_before_any_delete(self):
        with self.assertRaises(ValueError):
            audit.check_identity(ServiceResult("exists", data={"id": "new-id", "userPrincipalName": self.p["upn"]}), "old-id", self.p["upn"], "Microsoft")

    def test_aspirante_with_new_courses_teams_or_activity_is_not_deleted(self):
        p = person("aspirantes")
        for courses, teams, report_usage in (([{"id": 1}], [], usage(p)), ([], [{"id": "team"}], usage(p)), ([], [], {})):
            m, g = MagicMock(), MagicMock()
            m.get_user_courses.return_value = MoodleResult("ok", data=courses)
            self.execute(self.create_batch(p), p, m, g, report_usage=report_usage, teams=teams)
            m.delete_user_by_id.assert_not_called()
            g.delete_user.assert_not_called()

    def test_dry_run_is_not_a_successful_deletion(self):
        self.settings.crm_dry_run = True
        batch = self.create_batch()
        m, g = self.execute(batch)
        m.delete_user_by_id.assert_not_called()
        g.delete_user.assert_not_called()
        with self.session() as db:
            self.assertEqual(db.get(AuditDeletion, batch).status, "error")
            self.assertEqual(db.query(AuditDeletionRow).one().moodle_status, "pending")

    def test_lease_serializes_workers_and_recovers_expired_tasks(self):
        with self.session() as db:
            audit.acquire_lease(db, "first")
            db.commit()
        with self.session() as db:
            with self.assertRaises(ValueError):
                audit.acquire_lease(db, "second")
            db.get(AuditRun, self.run.id).status = "running"
            db.get(AuditLease, 1).expires_at = datetime.utcnow() - timedelta(seconds=1)
            db.commit()
            audit.recover_stale(db)
            self.assertEqual(db.get(AuditRun, self.run.id).status, "interrupted")

    def client(self):
        app = FastAPI()
        app.add_middleware(SessionMiddleware, secret_key="test-only")
        @app.get("/test-session")
        def session(request: Request):
            request.session["username"] = "test"
        def require_user(request):
            from fastapi import HTTPException
            if not request.session.get("username"):
                raise HTTPException(403)
            return request.session["username"]
        templates = Jinja2Templates(directory="app/templates")
        templates.env.filters["local_datetime"] = lambda v: str(v or "")
        app.include_router(create_audit_router(templates, require_user))
        def db_override():
            with self.session() as db:
                yield db
        app.dependency_overrides[get_db] = db_override
        client = TestClient(app)
        self.addCleanup(client.close)
        return client

    def test_ui_export_and_preview_do_not_call_external_services(self):
        client = self.client()
        self.assertEqual(client.get("/auditoria").status_code, 403)
        client.get("/test-session")
        with patch.object(audit, "MoodleClient") as m, patch.object(audit, "GraphClient") as g:
            response = client.get("/auditoria")
            self.assertEqual(response.status_code, 200)
            token = re.search(r'name="token" value="([^"]+)"', response.text)[1]
            xlsx = client.get(f"/auditoria/{self.run.id}/excel")
            self.assertEqual(xlsx.status_code, 200)
            workbook = load_workbook(io.BytesIO(xlsx.content))
            self.assertIn("Uso_Teams", workbook.sheetnames)
            self.assertEqual(workbook["Personas"]["A2"].value, self.p["dni"])
            self.assertEqual(client.post(f"/auditoria/{self.run.id}/bajas/previa", data={"dni": self.p["dni"]}).status_code, 403)
            preview = client.post(f"/auditoria/{self.run.id}/bajas/previa", data={"token": token, "dni": self.p["dni"]})
            self.assertEqual(preview.status_code, 200)
            self.assertIn("Lote #", preview.text)
            m.assert_not_called()
            g.assert_not_called()

    def test_forged_selection_and_bad_confirmation_are_rejected(self):
        client = self.client()
        client.get("/test-session")
        page = client.get("/auditoria")
        token = re.search(r'name="token" value="([^"]+)"', page.text)[1]
        response = client.post(f"/auditoria/{self.run.id}/bajas/previa", data={"token": token, "dni": "11111111"})
        self.assertIn("no habilitadas", response.text)
        batch = self.create_batch()
        with patch.object(audit, "deletion_job") as job:
            client.post(f"/auditoria/bajas/{batch}/confirmar", data={"token": token, "confirmation": "ELIMINAR", "understood": "yes", "password": "bad-password"})
            job.assert_not_called()

    def test_filter_pagination_and_excel_formula_safety(self):
        for index in range(2, 30):
            p = {**self.p, "dni": str(99888777 + index), "upn": f"a{99888777 + index}@example.invalid", "nombre": "=2+2"}
            self.create_snapshot([p])
        client = self.client()
        client.get("/test-session")
        first = client.get("/auditoria?per_page=25&page=1")
        second = client.get("/auditoria?per_page=25&page=2")
        self.assertEqual(first.text.count('name="dni"'), 25)
        self.assertEqual(second.text.count('name="dni"'), 4)
        wb = load_workbook(audit.export_excel(self.run, audit.people_for_run(self.run)))
        cells = [r[2] for r in list(wb["Personas"])[1:] if r[2].value == "=2+2"]
        self.assertTrue(cells)
        self.assertTrue(all(c.data_type == "s" for c in cells))

    def test_full_background_read_pipeline_persists_snapshot_and_progress(self):
        with self.session() as db:
            run = AuditRun(snapshot_key="b" * 32, campaign_start="2025-12-01", campaign_end="2026-03-01",
                           environment=audit.environment_key(self.settings), actor="test")
            db.add(run)
            audit.acquire_lease(db, "read-test")
            db.commit()
            run_id = run.id
        moodle = MagicMock()
        moodle.get_users_by_usernames.return_value = MoodleResult("ok", data=[{
            "id": 7, "username": self.p["upn"], "firstname": "Alumno", "lastname": "Prueba", "firstaccess": 0, "lastaccess": 0}])
        moodle.get_user_courses.return_value = MoodleResult("ok", data=[])
        reader = MagicMock()
        def graph_batch(paths):
            if paths[0].endswith("joinedTeams"):
                return [(200, {"value": []})]
            return [(200, {"id": "ms-test-id", "userPrincipalName": self.p["upn"], "createdDateTime": "2025-01-01T00:00:00Z"})]
        reader.batch.side_effect = graph_batch
        reader.list_pages.return_value = []
        data = {"preinscriptos": [{"nrodoc": self.p["dni"], "carreras": [{"fecha_inscri": "10/12/2025 12:00:00"}]}],
                "aspirantes": [], "ingresantes": []}
        def report_download(reader, endpoint):
            label = "Teams" if "Teams" in endpoint else "Microsoft365"
            row = dict(usage(self.p)[(self.p["dni"], label)][0])
            row["User Principal Name"] = self.p["upn"]
            if label == "Teams":
                row["User Id"] = self.p["microsoft_id"]
            buffer = io.StringIO()
            writer = csv.DictWriter(buffer, fieldnames=list(row))
            writer.writeheader()
            writer.writerow(row)
            return buffer.getvalue().encode()
        with patch.object(sources, "sial_lists", return_value=data), patch.object(sources, "MoodleClient", return_value=moodle), \
             patch.object(sources, "GraphReadOnly", return_value=reader), patch.object(reports, "GraphReadOnly", return_value=reader), \
             patch.object(sources.time, "sleep"), patch.object(reports, "download_report", side_effect=report_download):
            audit.audit_job(run_id, "read-test", self.session, self.settings)
        with self.session() as db:
            result = db.get(AuditRun, run_id)
            self.assertEqual(result.status, "completed", result.message)
            self.assertEqual(result.progress, 100)
            self.assertEqual(json.loads(result.counts_json)["preinscriptos"], 1)
            self.assertEqual(audit.people_for_run(result)[0]["candidate"], "preinscripto")
        moodle.delete_user_by_id.assert_not_called()
        moodle.enrol_user.assert_not_called()


if __name__ == "__main__":
    unittest.main()
