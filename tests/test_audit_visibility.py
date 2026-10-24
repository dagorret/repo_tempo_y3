from datetime import date, timedelta
import hashlib
import io
import re
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from app import auditing as audit
from app.models import AuditRun, AuditDeletion
import test_auditing as fixtures


class AuditVisibilityTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        fixtures.AuditTest.setUpClass()

    def setUp(self):
        self.case = fixtures.AuditTest()
        self.case.setUp()
        self.addCleanup(self.case.doCleanups)

    def set_end(self, end):
        self.case.run.campaign_end = end
        with self.case.session() as db:
            db.get(AuditRun, self.case.run.id).campaign_end = end
            db.commit()

    def test_candidates_visible_with_today_future_or_missing_close(self):
        for end in (date.today().isoformat(), (date.today() + timedelta(days=1)).isoformat(), ""):
            with self.subTest(end=end):
                self.set_end(end)
                people = audit.people_for_run(self.case.run)
                self.assertEqual(audit.counts_for(people)["candidates"], 1)
                self.assertEqual(people[0]["candidate"], "preinscripto")
                self.assertFalse(people[0]["deletion_eligible"])
                self.assertIn("Campana abierta", people[0]["deletion_block_reason"])
                self.assertFalse(audit.eligibility(people[0], self.case.run)[0])

    def test_aspirante_is_visible_before_close(self):
        p = fixtures.person("aspirantes")
        p.update(dni="99888888", upn="a99888888@example.invalid")
        self.case.create_snapshot([p])
        self.set_end(date.today().isoformat())
        people = audit.people_for_run(self.case.run)
        self.assertEqual(audit.counts_for(people)["asp_candidates"], 1)
        self.assertEqual(audit.filter_people(people, "asp_candidates")[0]["dni"], p["dni"])
        self.assertTrue(all(not p["deletion_eligible"] for p in people))

    def test_period_identity_and_activity_checks_are_preserved(self):
        p = self.case.p
        run = self.case.run
        for change in ({"inscripciones": []}, {"estados_sial": "ingresantes"}, {"moodle_estado": "error"},
                       {"inscripciones": [{"fecha_inscri": "01/01/2027"}]},
                       {"estados_sial": "aspirantes", "actividad": "Datos incompletos"}):
            with self.subTest(change=change):
                self.assertFalse(audit.candidate_assessment({**p, **change}, run)[0])

    def test_closed_campaign_keeps_existing_behavior(self):
        self.set_end((date.today() - timedelta(days=1)).isoformat())
        p = audit.people_for_run(self.case.run)[0]
        self.assertTrue(p["deletion_eligible"])
        self.assertEqual(p["candidate"], "preinscripto")
        self.assertEqual(p["deletion_block_reason"], "")

    def test_incomplete_run_does_not_claim_deletion_is_enabled(self):
        self.case.run.status = "error"
        p = audit.people_for_run(self.case.run)[0]
        self.assertEqual(p["candidate"], "preinscripto")
        self.assertFalse(p["deletion_eligible"])
        self.assertIn("Auditoria sin finalizar", p["deletion_block_reason"])

    def test_view_excel_and_detail_read_saved_data_without_external_calls(self):
        self.set_end(date.today().isoformat())
        before = hashlib.sha256(audit.snapshot_path(self.case.run).read_bytes()).digest()
        client = self.case.client()
        client.get("/test-session")
        with patch.object(audit, "MoodleClient") as moodle, patch.object(audit, "GraphClient") as graph, \
             patch.object(audit.sources, "sial_lists") as sial, patch.object(audit.sources, "GraphReadOnly") as reader:
            page = client.get("/auditoria?view=candidates")
            self.assertEqual(page.status_code, 200)
            self.assertIn("1 candidatos en este filtro", page.text)
            self.assertIn("0 habilitados para preparar baja", page.text)
            self.assertNotIn('name="dni"', page.text)
            self.assertIn("Campa&ntilde;a sin cerrar", page.text)
            detail = client.get(f"/auditoria/{self.case.run.id}/persona/{self.case.p['dni']}")
            self.assertIn("Campana abierta", detail.text)
            result = client.get(f"/auditoria/{self.case.run.id}/excel")
            self.assertEqual(result.status_code, 200)
            workbook = load_workbook(io.BytesIO(result.content))
            sheet = workbook["Candidatos_a_revision"]
            self.assertEqual(sheet.max_row, 2)
            values = dict(zip([cell.value for cell in sheet[1]], [cell.value for cell in sheet[2]]))
            self.assertEqual(values["Candidato"], "preinscripto")
            self.assertEqual(values["Habilitado para preparar baja"], "No")
            self.assertIn("Campana abierta", values["Motivo del bloqueo"])
            for mock in (moodle, graph, sial, reader):
                mock.assert_not_called()
        after = hashlib.sha256(audit.snapshot_path(self.case.run).read_bytes()).digest()
        self.assertEqual(before, after)

    def test_open_campaign_blocks_individual_bulk_and_confirm_posts(self):
        self.set_end(date.today().isoformat())
        client = self.case.client()
        client.get("/test-session")
        page = client.get("/auditoria")
        token = re.search(r'name="token" value="([^"]+)"', page.text)[1]
        for selection in ({"dni": self.case.p["dni"]}, {"all_filtered": "yes", "view": "candidates"}):
            client.post(f"/auditoria/{self.case.run.id}/bajas/previa", data={"token": token, **selection})
        with self.case.session() as db:
            self.assertEqual(db.query(AuditDeletion).count(), 0)
        batch = self.case.create_batch()
        with patch.object(audit, "deletion_job") as job:
            response = client.post(f"/auditoria/bajas/{batch}/confirmar", data={
                "token": token, "confirmation": "ELIMINAR", "understood": "yes", "password": "test-password-only"})
            self.assertIn("Campana abierta", response.text)
            job.assert_not_called()

    def test_background_worker_rejects_open_campaign_even_with_candidate(self):
        self.set_end(date.today().isoformat())
        batch = self.case.create_batch()
        moodle, graph = self.case.execute(batch)
        moodle.get_user_by_username.assert_not_called()
        graph.get_user.assert_not_called()
        moodle.delete_user_by_id.assert_not_called()
        graph.delete_user.assert_not_called()
        with self.case.session() as db:
            self.assertEqual(db.get(AuditDeletion, batch).status, "error")


if __name__ == "__main__":
    unittest.main()
