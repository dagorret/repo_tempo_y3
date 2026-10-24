from pathlib import Path
from datetime import datetime, timezone
import csv
import io
import json
import re
import secrets
import time
import uuid
from urllib.parse import quote_plus, urlencode

from fastapi import BackgroundTasks, Depends, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from sqlalchemy import inspect, or_, text
from sqlalchemy.orm import Session
from starlette.middleware.sessions import SessionMiddleware

from app.account_rules import build_student_upn, clean_document, normalize_name, title_name
from app.config import get_settings, update_env_file
from app.audit_routes import create_audit_router
from app.ingreso_routes import create_ingreso_router
from app.auditing import safe_error
from app.auth_azure_ad import AzureADAuth
from app.batch_runner import run_batch_in_background
from app.db import Base, engine, get_db, SessionLocal
from app.importers import ImportedRow, parse_people_excel, preview_people_excel, suggest_column_mapping
from app.electives import (
    ElectiveIssue,
    ElectiveRow,
    analyze_elective_subjects,
    parse_elective_excel,
    preview_elective_excel,
    suggest_elective_mapping,
)
from app.lincu_processing import build_lincu_report, parse_ingresos_csv, parse_lincu_txt
from app.matriculation import (
    COURSE_RULES,
    CourseRule,
    build_course_load_alerts,
    parse_matriculation_lincu,
)
from app.matriculation_verification import course_usernames
from app.notifications import (
    NotificationRecipient,
    parse_notifications_excel,
    parse_notifications_txt,
    preview_notifications_excel,
)
from app.models import (
    ImportedPerson,
    ImportBatch,
    CourseConversionRule,
    LocalUser,
    MatriculationRun,
    MatriculationRunRow,
    MatriculationCourseAlert,
    MatriculationState,
    ManagedEnrollment,
    ManagedUser,
    OperationLog,
)
from app.security import generate_temporary_password, seed_local_users, verify_password
from app.services.graph import GraphClient
from app.services.mailer import Mailer
from app.services.moodle import MOODLE_FUNCTION_CATALOG, MoodleClient, MoodleResult
from app.user_registry import upsert_managed_enrollment, upsert_managed_user
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


settings = get_settings()
app = FastAPI(title="CRM-FCE")
app.add_middleware(SessionMiddleware, secret_key=settings.crm_secret_key)
templates = Jinja2Templates(directory=str(Path(__file__).parent / "templates"))
templates.env.globals["settings"] = get_settings()


def _format_local_datetime(value: datetime | None, format_string: str = "%Y-%m-%d %H:%M") -> str:
    if value is None:
        return ""
    source = value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value
    local_timezone = datetime.now().astimezone().tzinfo
    return source.astimezone(local_timezone).strftime(format_string)


templates.env.filters["local_datetime"] = _format_local_datetime
app.mount("/static", StaticFiles(directory=str(Path(__file__).parent / "static")), name="static")
PENDING_IMPORTS_DIR = Path("data/pending_imports")
PENDING_ELECTIVES_DIR = Path("data/pending_electives")
PENDING_NOTIFICATIONS_DIR = Path("data/pending_notifications")
PENDING_LINCU_DIR = Path("data/pending_lincu")
NOTIFICATION_JOBS: dict[str, dict] = {}
LINCU_JOBS: dict[str, dict] = {}
MANAGED_USERS_PAGE_SIZE = 25
MANAGED_USERS_PAGE_SIZES = {25, 50, 100}
MATRICULATION_MOODLE_BATCH_SIZE = 25
MATRICULATION_USER_LOOKUP_BATCH_SIZE = 100
NOTIFICATION_DEFAULT_SUBJECT = "Tìtulo del envio"
NOTIFICATION_DEFAULT_MESSAGE = "Escribir acá el mensaje que deseas enviar a los estudiantes."


@app.on_event("startup")
def startup() -> None:
    Base.metadata.create_all(bind=engine)
    ensure_schema()
    db = SessionLocal()
    try:
        seed_local_users(db, settings)
        seed_course_conversion_rules(db)
        backfill_managed_users(db)
        _recover_interrupted_matriculation_runs(db)
    finally:
        db.close()


def ensure_schema() -> None:
    inspector = inspect(engine)
    if "imported_people" not in inspector.get_table_names():
        return

    columns = {column["name"] for column in inspector.get_columns("imported_people")}
    with engine.begin() as connection:
        if "full_name" not in columns:
            connection.execute(text("ALTER TABLE imported_people ADD COLUMN full_name VARCHAR(255) DEFAULT ''"))
        if "azure_user_id" not in columns:
            connection.execute(text("ALTER TABLE imported_people ADD COLUMN azure_user_id VARCHAR(80) DEFAULT ''"))
        if "moodle_user_id" not in columns:
            connection.execute(text("ALTER TABLE imported_people ADD COLUMN moodle_user_id VARCHAR(80) DEFAULT ''"))
        connection.execute(
            text(
                """
                UPDATE imported_people
                SET full_name = trim(coalesce(first_name, '') || ' ' || coalesce(last_name, ''))
                WHERE full_name IS NULL OR full_name = ''
                """
            )
        )
    if "import_batches" in inspector.get_table_names():
        batch_columns = {column["name"] for column in inspector.get_columns("import_batches")}
        with engine.begin() as connection:
            if "email_message" not in batch_columns:
                connection.execute(text("ALTER TABLE import_batches ADD COLUMN email_message TEXT DEFAULT ''"))
            if "email_subject" not in batch_columns:
                connection.execute(text("ALTER TABLE import_batches ADD COLUMN email_subject VARCHAR(255) DEFAULT ''"))
            if "moodle_group_name" not in batch_columns:
                connection.execute(text("ALTER TABLE import_batches ADD COLUMN moodle_group_name VARCHAR(120) DEFAULT ''"))
            if "moodle_group_autocreate" not in batch_columns:
                connection.execute(text("ALTER TABLE import_batches ADD COLUMN moodle_group_autocreate BOOLEAN DEFAULT 1"))
            if "moodle_group_mode" not in batch_columns:
                connection.execute(text("ALTER TABLE import_batches ADD COLUMN moodle_group_mode VARCHAR(20) DEFAULT 'tolerant'"))
            if "run_status" not in batch_columns:
                connection.execute(text("ALTER TABLE import_batches ADD COLUMN run_status VARCHAR(30) DEFAULT 'pending'"))
            if "total_rows" not in batch_columns:
                connection.execute(text("ALTER TABLE import_batches ADD COLUMN total_rows INTEGER DEFAULT 0"))
            if "processed_rows" not in batch_columns:
                connection.execute(text("ALTER TABLE import_batches ADD COLUMN processed_rows INTEGER DEFAULT 0"))
            if "success_rows" not in batch_columns:
                connection.execute(text("ALTER TABLE import_batches ADD COLUMN success_rows INTEGER DEFAULT 0"))
            if "error_rows" not in batch_columns:
                connection.execute(text("ALTER TABLE import_batches ADD COLUMN error_rows INTEGER DEFAULT 0"))
            if "started_at" not in batch_columns:
                connection.execute(text("ALTER TABLE import_batches ADD COLUMN started_at DATETIME"))
            if "finished_at" not in batch_columns:
                connection.execute(text("ALTER TABLE import_batches ADD COLUMN finished_at DATETIME"))
            connection.execute(
                text(
                    """
                    UPDATE import_batches
                    SET total_rows = (
                        SELECT COUNT(1)
                        FROM imported_people p
                        WHERE p.batch_id = import_batches.id
                    )
                    WHERE total_rows IS NULL OR total_rows = 0
                    """
                )
            )
            connection.execute(
                text(
                    """
                    UPDATE import_batches
                    SET processed_rows = (
                        SELECT COUNT(1)
                        FROM imported_people p
                        WHERE p.batch_id = import_batches.id
                    )
                    WHERE processed_rows IS NULL OR processed_rows = 0
                    """
                )
            )
            connection.execute(
                text(
                    """
                    UPDATE import_batches
                    SET success_rows = (
                        SELECT COUNT(1)
                        FROM imported_people p
                        WHERE p.batch_id = import_batches.id
                        AND p.status <> 'error'
                    )
                    WHERE success_rows IS NULL OR success_rows = 0
                    """
                )
            )
            connection.execute(
                text(
                    """
                    UPDATE import_batches
                    SET error_rows = (
                        SELECT COUNT(1)
                        FROM imported_people p
                        WHERE p.batch_id = import_batches.id
                        AND p.status = 'error'
                    )
                    WHERE error_rows IS NULL OR error_rows = 0
                    """
                )
            )
            connection.execute(
                text(
                    """
                    UPDATE import_batches
                    SET run_status = 'completed'
                    WHERE (run_status IS NULL OR run_status = '' OR run_status = 'pending')
                    AND processed_rows >= total_rows
                    AND total_rows > 0
                    """
                )
            )

    if "matriculation_runs" in inspector.get_table_names():
        run_columns = {column["name"] for column in inspector.get_columns("matriculation_runs")}
        with engine.begin() as connection:
            for column, definition in {
                "verification_total_courses": "INTEGER DEFAULT 0",
                "verification_processed_courses": "INTEGER DEFAULT 0",
                "verification_scope": "VARCHAR(255) DEFAULT ''",
                "verification_started_at": "DATETIME",
                "verification_finished_at": "DATETIME",
            }.items():
                if column not in run_columns:
                    connection.execute(text(f"ALTER TABLE matriculation_runs ADD COLUMN {column} {definition}"))
            if "execution_total_rows" not in run_columns:
                connection.execute(
                    text("ALTER TABLE matriculation_runs ADD COLUMN execution_total_rows INTEGER DEFAULT 0")
                )
            if "execution_processed_rows" not in run_columns:
                connection.execute(
                    text("ALTER TABLE matriculation_runs ADD COLUMN execution_processed_rows INTEGER DEFAULT 0")
                )
            if "course_change_rows" not in run_columns:
                connection.execute(
                    text("ALTER TABLE matriculation_runs ADD COLUMN course_change_rows INTEGER DEFAULT 0")
                )
            if "course_alert_threshold" not in run_columns:
                connection.execute(
                    text("ALTER TABLE matriculation_runs ADD COLUMN course_alert_threshold INTEGER DEFAULT 5")
                )
            if "course_alert_count" not in run_columns:
                connection.execute(
                    text("ALTER TABLE matriculation_runs ADD COLUMN course_alert_count INTEGER DEFAULT 0")
                )
            if "source_type" not in run_columns:
                connection.execute(
                    text("ALTER TABLE matriculation_runs ADD COLUMN source_type VARCHAR(30) DEFAULT 'lincu'")
                )

    if "matriculation_run_rows" in inspector.get_table_names():
        row_columns = {column["name"] for column in inspector.get_columns("matriculation_run_rows")}
        with engine.begin() as connection:
            if "verification_status" not in row_columns:
                connection.execute(text("ALTER TABLE matriculation_run_rows ADD COLUMN verification_status VARCHAR(30) DEFAULT ''"))
            if "logical_course_code" not in row_columns:
                connection.execute(
                    text("ALTER TABLE matriculation_run_rows ADD COLUMN logical_course_code VARCHAR(40) DEFAULT ''")
                )
            if "previous_course_shortname" not in row_columns:
                connection.execute(
                    text("ALTER TABLE matriculation_run_rows ADD COLUMN previous_course_shortname VARCHAR(120) DEFAULT ''")
                )
            if "previous_moodle_course_id" not in row_columns:
                connection.execute(
                    text("ALTER TABLE matriculation_run_rows ADD COLUMN previous_moodle_course_id VARCHAR(80) DEFAULT ''")
                )
            connection.execute(
                text(
                    """
                    UPDATE matriculation_run_rows
                    SET logical_course_code = CASE
                        WHEN instr(course_shortname, '-') > 0
                        THEN substr(course_shortname, 1, instr(course_shortname, '-') - 1)
                        ELSE course_shortname
                    END
                    WHERE logical_course_code IS NULL OR logical_course_code = ''
                    """
                )
            )
            connection.execute(
                text(
                    "CREATE INDEX IF NOT EXISTS ix_matriculation_run_rows_logical_course_code "
                    "ON matriculation_run_rows (logical_course_code)"
                )
            )

    if "matriculation_states" in inspector.get_table_names():
        state_columns = {column["name"] for column in inspector.get_columns("matriculation_states")}
        with engine.begin() as connection:
            if "logical_course_code" not in state_columns:
                connection.execute(
                    text("ALTER TABLE matriculation_states ADD COLUMN logical_course_code VARCHAR(40) DEFAULT ''")
                )
            connection.execute(
                text(
                    """
                    UPDATE matriculation_states
                    SET logical_course_code = CASE
                        WHEN instr(course_shortname, '-') > 0
                        THEN substr(course_shortname, 1, instr(course_shortname, '-') - 1)
                        ELSE course_shortname
                    END
                    WHERE logical_course_code IS NULL OR logical_course_code = ''
                    """
                )
            )
            connection.execute(
                text(
                    "CREATE INDEX IF NOT EXISTS ix_matriculation_states_logical_course_code "
                    "ON matriculation_states (logical_course_code)"
                )
            )


def seed_course_conversion_rules(db: Session) -> None:
    seeded = (
        db.query(OperationLog)
        .filter(OperationLog.action == "course_conversion_rules_seeded")
        .first()
    )
    if seeded:
        return
    if db.query(CourseConversionRule).count() == 0:
        db.add_all(
            [
                CourseConversionRule(
                    actual_code=rule.actual,
                    old_code=rule.old,
                    category_rule=rule.category,
                    registration_rule="",
                    sort_order=index * 10,
                )
                for index, rule in enumerate(COURSE_RULES, start=1)
            ]
        )
    db.add(
        OperationLog(
            actor="system",
            action="course_conversion_rules_seeded",
            status="ok",
            detail=f"Tabla inicial del conversor cargada con {len(COURSE_RULES)} reglas.",
        )
    )
    db.commit()


def _normalize_conversion_rule_fields(
    actual_code: str,
    old_code: str,
    category_rule: str,
    registration_rule: str,
) -> tuple[str, str, str, str]:
    actual = str(actual_code or "").strip()
    old = str(old_code or "").strip()
    category = str(category_rule or "").strip().upper().replace(" ", "")
    registration = str(registration_rule or "").strip()
    if not actual or not actual.isdigit():
        raise ValueError("El codigo Moodle (Actual) debe ser numerico.")
    if old and not old.isdigit():
        raise ValueError("El codigo alternativo Lincu (Viejo) debe ser numerico.")
    if category and not re.fullmatch(r"(?:\d+|X\*?)", category):
        raise ValueError("La regla de Catedra debe quedar vacia, ser numerica, X o X*.")
    return str(int(actual)), str(int(old)) if old else "", category, registration[:40]


def _conversion_rule_result(rule: CourseConversionRule) -> str:
    category = str(rule.category_rule or "").strip().upper()
    if rule.actual_code == "105":
        return "105-1 / 105-2 / 105-3"
    if category == "X":
        return f"{rule.actual_code}-CATEDRA"
    if category == "X*":
        return f"{rule.actual_code}-CATEDRA-D/P"
    if category:
        return f"{rule.actual_code}-{category}"
    return rule.actual_code


templates.env.globals["conversion_result"] = _conversion_rule_result


def current_username(request: Request) -> str | None:
    return request.session.get("username")


def require_user(request: Request) -> str:
    username = current_username(request)
    if not username:
        raise HTTPException(status_code=303, headers={"Location": "/login"})
    return username


app.include_router(create_audit_router(templates, require_user))
app.include_router(create_ingreso_router(templates, require_user))


def render_dashboard(
    request: Request,
    db: Session,
    username: str,
    error: str = "",
    status_code: int = 200,
):
    batches = db.query(ImportBatch).order_by(ImportBatch.created_at.desc()).limit(20).all()
    return templates.TemplateResponse(
        "dashboard.html",
        {
            "request": request,
            "username": username,
            "batches": batches,
            "settings": settings,
            "error": error,
        },
        status_code=status_code,
    )


def render_mapping(
    request: Request,
    username: str,
    token: str,
    filename: str,
    person_type: str,
    course_shortname: str,
    moodle_group_name: str,
    moodle_group_autocreate: bool,
    moodle_group_mode: str,
    columns: list[str],
    sample_rows: list[dict[str, str]],
    mapping: dict[str, str],
    error: str = "",
    status_code: int = 200,
):
    return templates.TemplateResponse(
        "map_columns.html",
        {
            "request": request,
            "username": username,
            "token": token,
            "filename": filename,
            "person_type": person_type,
            "course_shortname": course_shortname,
            "moodle_group_name": moodle_group_name,
            "moodle_group_autocreate": moodle_group_autocreate,
            "moodle_group_mode": moodle_group_mode,
            "columns": columns,
            "sample_rows": sample_rows,
            "mapping": mapping,
            "error": error,
        },
        status_code=status_code,
    )


def _person_type_label(person_type: str) -> str:
    return "Estudiantes" if person_type == "student" else "Profesores"


def parse_course_shortnames(raw_value: str) -> list[str]:
    if not raw_value:
        return []
    values = [part.strip() for part in raw_value.replace(";", ",").split(",")]
    return [value for value in values if value]


def _notification_counts(recipients: list[NotificationRecipient]) -> dict[str, int]:
    counts = {"all": len(recipients), "distance": 0, "presential": 0, "other": 0}
    for recipient in recipients:
        group = recipient.modality_group
        counts[group] = counts.get(group, 0) + 1
    return counts


def _filter_notification_recipients(
    recipients: list[NotificationRecipient],
    target_mode: str,
) -> list[NotificationRecipient]:
    mode = (target_mode or "all").strip().lower()
    if mode == "all":
        return recipients
    return [row for row in recipients if row.modality_group == mode]


def _notification_file_path(token: str) -> Path:
    return PENDING_NOTIFICATIONS_DIR / f"{token}.json"


def _suggest_notification_mapping(columns: list[str]) -> dict[str, str]:
    aliases = {
        "email": ["mail", "email", "correo", "e-mail"],
        "modality": ["modalidad", "modo", "tipo_modalidad"],
        "dni": ["dni", "documento", "nrodoc", "nro_doc", "nro_documento"],
        "last_name": ["apellido", "ape", "ape.alumno", "apellidos"],
        "first_name": ["nombre", "nom", "nom.alumno", "nombres"],
    }
    normalized = {normalize_name(col): col for col in columns}
    result = {"email": "", "modality": "", "dni": "", "last_name": "", "first_name": ""}
    for key, values in aliases.items():
        for value in values:
            if normalize_name(value) in normalized:
                result[key] = normalized[normalize_name(value)]
                break
    return result


def _store_notification_preview(filename: str, recipients: list[NotificationRecipient]) -> str:
    token = secrets.token_urlsafe(16)
    payload = {
        "filename": filename,
        "created_at": datetime.utcnow().isoformat(),
        "counts": _notification_counts(recipients),
        "recipients": [
            {
                "dni": r.dni,
                "first_name": r.first_name,
                "last_name": r.last_name,
                "full_name": r.full_name,
                "email": r.email,
                "modality": r.modality_raw,
                "modality_group": r.modality_group,
            }
            for r in recipients
        ],
    }
    _notification_file_path(token).write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return token


def _lincu_file_path(job_id: str, kind: str) -> Path:
    return PENDING_LINCU_DIR / f"{job_id}_{kind}"


def _run_lincu_job(*, job_id: str, actor: str) -> None:
    job = LINCU_JOBS.get(job_id)
    if not job:
        return

    lincu_path = _lincu_file_path(job_id, "lincu.txt")
    ingresos_path = _lincu_file_path(job_id, "ingresos.csv")
    started_at = datetime.utcnow().isoformat()
    LINCU_JOBS[job_id].update(
        {
            "status": "running",
            "stage": "Leyendo Lincu",
            "progress": 5,
            "started_at": started_at,
            "finished_at": "",
            "error": "",
        }
    )

    try:
        lincu_map, lincu_rows = parse_lincu_txt(lincu_path.read_bytes())
        LINCU_JOBS[job_id].update(
            {
                "stage": "Leyendo CSV de ingresos",
                "progress": 40,
                "lincu_rows": lincu_rows,
            }
        )

        ingresos_map, ingresos_rows = parse_ingresos_csv(ingresos_path.read_bytes())
        LINCU_JOBS[job_id].update(
            {
                "stage": "Consolidando por DNI",
                "progress": 75,
                "ingresos_rows": ingresos_rows,
            }
        )

        rows = build_lincu_report(lincu_map, ingresos_map)
        LINCU_JOBS[job_id].update(
            {
                "status": "completed",
                "stage": "Completado",
                "progress": 100,
                "rows": rows,
                "summary": {
                    "total_lincu_rows": lincu_rows,
                    "total_ingresos_rows": ingresos_rows,
                    "students_with_ingresos": len(rows),
                },
                "finished_at": datetime.utcnow().isoformat(),
            }
        )
        db = SessionLocal()
        try:
            db.add(
                OperationLog(
                    actor=actor,
                    action="lincu_process",
                    status="ok",
                    detail=f"Lincu procesado. Filas resultado={len(rows)}.",
                )
            )
            db.commit()
        finally:
            db.close()
    except Exception as exc:
        LINCU_JOBS[job_id].update(
            {
                "status": "failed",
                "stage": "Error",
                "progress": 100,
                "error": str(exc),
                "finished_at": datetime.utcnow().isoformat(),
            }
        )
        db = SessionLocal()
        try:
            db.add(
                OperationLog(
                    actor=actor,
                    action="lincu_process",
                    status="error",
                    detail=f"Error Lincu: {exc}",
                )
            )
            db.commit()
        finally:
            db.close()
    finally:
        lincu_path.unlink(missing_ok=True)
        ingresos_path.unlink(missing_ok=True)


def _matriculation_recount(run: MatriculationRun, db: Session) -> None:
    rows = db.query(MatriculationRunRow).filter(MatriculationRunRow.run_id == run.id).all()
    run.enroll_rows = sum(1 for row in rows if row.action == "enroll")
    run.group_change_rows = sum(1 for row in rows if row.action == "group_change")
    run.course_change_rows = sum(1 for row in rows if row.action == "course_change")
    run.unenroll_rows = sum(1 for row in rows if row.action == "unenroll")
    run.unchanged_rows = sum(1 for row in rows if row.action == "unchanged")
    run.error_rows = sum(1 for row in rows if row.action == "error" or row.status in {"error", "unverified"})
    run.processed_rows = sum(1 for row in rows if row.status in {"done", "skipped", "error"})


def _recover_interrupted_matriculation_runs(db: Session) -> None:
    """Unlock local-only runs that were interrupted by a server restart."""
    interrupted_run_ids = {
        run_id
        for (run_id,) in db.query(MatriculationRunRow.run_id)
        .filter(
            MatriculationRunRow.action.in_(MATRICULATION_ACTIONS),
            MatriculationRunRow.status == "running",
        )
        .distinct()
        .all()
    }
    interrupted_runs = (
        db.query(MatriculationRun)
        .filter(
            or_(
                MatriculationRun.status.in_({"queued", "running", "verifying"}),
                MatriculationRun.id.in_(interrupted_run_ids) if interrupted_run_ids else False,
            )
        )
        .all()
    )
    for run in interrupted_runs:
        if run.status == "verifying":
            db.query(MatriculationRunRow).filter_by(run_id=run.id, verification_status="pending").update({
                "verification_status": "error", "status": "unverified",
                "message": "Comprobacion Moodle interrumpida. Volver a subir el Lincu para verificar; no se ejecutara esta fila.",
            })
            run.verification_finished_at = datetime.utcnow()
            _matriculation_recount(run, db)
            run.status = "completed_with_errors"
            run.message = "Comprobacion Moodle interrumpida. Volver a subir el Lincu; no se realizaron cambios en Moodle."
            continue
        interrupted_rows = (
            db.query(MatriculationRunRow)
            .filter(
                MatriculationRunRow.run_id == run.id,
                MatriculationRunRow.action.in_(MATRICULATION_ACTIONS),
                MatriculationRunRow.status == "running",
            )
            .all()
        )
        for row in interrupted_rows:
            row.status = "error"
            row.message = "Ejecucion interrumpida por reinicio del servidor. Reintenta esta fila."
        db.flush()
        _matriculation_recount(run, db)
        recoverable_rows = (
            db.query(MatriculationRunRow)
            .filter(
                MatriculationRunRow.run_id == run.id,
                MatriculationRunRow.action.in_(MATRICULATION_ACTIONS),
                MatriculationRunRow.status.in_({"preview", "error"}),
            )
            .count()
        )
        run.status = "completed_with_errors" if run.error_rows else "preview"
        run.finished_at = datetime.utcnow()
        run.message = (
            "Ejecucion interrumpida por reinicio del servidor. "
            f"Hay {recoverable_rows} filas disponibles para reintentar."
        )
    if interrupted_runs:
        db.commit()


def _load_course_conversion_rules(db: Session) -> list[CourseRule]:
    stored_rules = (
        db.query(CourseConversionRule)
        .order_by(CourseConversionRule.sort_order.asc(), CourseConversionRule.id.asc())
        .all()
    )
    return [
        CourseRule(
            actual=rule.actual_code,
            old=rule.old_code,
            category=rule.category_rule,
        )
        for rule in stored_rules
    ]


def _parse_matriculation_verification_scope(value: str) -> tuple[str, ...]:
    values: list[str] = []
    for raw_code in str(value or "").split(","):
        code = raw_code.strip()
        if not code:
            continue
        if not code.isdigit():
            raise ValueError("Los cursos a comprobar deben ser codigos numericos separados por coma. Ejemplo: 230, 105")
        normalized = str(int(code))
        if normalized not in values:
            values.append(normalized)
    if not values:
        raise ValueError("Indica al menos un codigo de curso para comprobar en Moodle.")
    return tuple(values)


def _is_in_matriculation_verification_scope(item, scope: tuple[str, ...]) -> bool:
    return item.logical_course_code in scope or item.source_code in scope


def _create_matriculation_preview(
    content: bytes,
    filename: str,
    actor: str,
    db: Session,
    *,
    verification_scope: tuple[str, ...] = (),
) -> MatriculationRun:
    conversion_rules = _load_course_conversion_rules(db)
    if not conversion_rules:
        raise ValueError("El conversor no tiene reglas cargadas.")
    desired, total_rows, issues = parse_matriculation_lincu(
        content,
        settings.crm_domain,
        rules=conversion_rules,
    )
    if verification_scope and not any(
        _is_in_matriculation_verification_scope(item, verification_scope)
        for item in desired.values()
    ):
        raise ValueError(
            "Ninguno de los codigos indicados aparece en este Lincu: "
            + ", ".join(verification_scope)
        )
    course_alert_threshold = max(
        1,
        min(int(settings.matriculation_course_alert_threshold or 5), 30),
    )
    course_alerts = build_course_load_alerts(
        content,
        threshold=course_alert_threshold,
        rules=conversion_rules,
    )
    active_states = (
        db.query(MatriculationState)
        .filter(MatriculationState.active.is_(True))
        .all()
    )
    states_by_key: dict[tuple[str, str], list[MatriculationState]] = {}
    for state in active_states:
        logical_code = state.logical_course_code or state.course_shortname.split("-", 1)[0]
        states_by_key.setdefault((state.document, logical_code), []).append(state)

    run = MatriculationRun(
        filename=filename,
        created_by=actor,
        status="verifying" if verification_scope else "preview",
        total_source_rows=total_rows,
        desired_rows=len(desired),
        course_alert_threshold=course_alert_threshold,
        course_alert_count=len(course_alerts),
        verification_scope=", ".join(verification_scope),
        message=(
            "Comprobacion Moodle en cola para: " + ", ".join(verification_scope) + ". Sin cambios externos."
            if verification_scope else
                 "Vista previa local, sin comprobar matriculas actuales en Moodle. Revisa los cambios antes de ejecutar."),
    )
    db.add(run)
    db.commit()
    db.refresh(run)

    db.add_all(
        [
            MatriculationCourseAlert(
                run_id=run.id,
                document=alert.document,
                first_name=alert.first_name,
                last_name=alert.last_name,
                phone=alert.phone,
                email=alert.email,
                modality=alert.modality,
                course_count=alert.course_count,
                course_codes=", ".join(code for code, _ in alert.courses),
                course_names="; ".join(
                    f"{code} - {name}" if name else code
                    for code, name in alert.courses
                ),
            )
            for alert in course_alerts
        ]
    )

    preview_rows: list[MatriculationRunRow] = []
    for key, item in desired.items():
        states = states_by_key.get(key, [])
        if len(states) > 1:
            desired_states = [state for state in states if state.course_shortname == item.course_shortname]
            previous_states = [state for state in states if state.course_shortname != item.course_shortname]
            if len(desired_states) == 1 and len(previous_states) == 1:
                # A previous course change completed the new enrolment but not the old removal.
                state = previous_states[0]
            else:
                courses = ", ".join(sorted(state.course_shortname for state in states))
                preview_rows.append(
                    MatriculationRunRow(
                        run_id=run.id,
                        action="error",
                        status="preview",
                        document=item.document,
                        upn=item.upn,
                        first_name=item.first_name,
                        last_name=item.last_name,
                        email=item.email,
                        source_code=item.source_code,
                        source_course=item.source_course,
                        logical_course_code=item.logical_course_code,
                        course_shortname=item.course_shortname,
                        group_name=item.group_name,
                        message=f"Hay mas de una matricula activa para la materia {item.logical_course_code}: {courses}.",
                    )
                )
                continue
        else:
            state = states[0] if states else None
        if state is None:
            action = "enroll"
            previous_group = ""
            previous_course = ""
            message = "Nueva matriculacion."
        elif state.course_shortname != item.course_shortname:
            action = "course_change"
            previous_group = state.group_name
            previous_course = state.course_shortname
            message = (
                f"Cambia curso {state.course_shortname} -> {item.course_shortname} "
                f"para la materia {item.logical_course_code}."
            )
        elif state.group_name != item.group_name:
            action = "group_change"
            previous_group = state.group_name
            previous_course = state.course_shortname
            message = f"Cambia grupo {state.group_name or '-'} -> {item.group_name}."
        else:
            action = "unchanged"
            previous_group = state.group_name
            previous_course = state.course_shortname
            message = "Sin cambios."

        preview_rows.append(
            MatriculationRunRow(
                run_id=run.id,
                action=action,
                status="preview",
                document=item.document,
                upn=item.upn,
                first_name=item.first_name,
                last_name=item.last_name,
                email=item.email,
                source_code=item.source_code,
                source_course=item.source_course,
                logical_course_code=item.logical_course_code,
                course_shortname=item.course_shortname,
                previous_course_shortname=previous_course,
                group_name=item.group_name,
                previous_group_name=previous_group,
                moodle_user_id=state.moodle_user_id if state else "",
                moodle_course_id=(
                    state.moodle_course_id
                    if state and state.course_shortname == item.course_shortname
                    else ""
                ),
                previous_moodle_course_id=state.moodle_course_id if state else "",
                message=message,
            )
        )

    for issue in issues:
        preview_rows.append(
            MatriculationRunRow(
                run_id=run.id,
                action="error",
                status="preview",
                document=issue.document,
                source_code=issue.source_code,
                source_course=issue.source_course,
                message=issue.message,
            )
        )

    if verification_scope:
        for row in preview_rows:
            if (
                row.action in MATRICULATION_ACTIONS | {"unchanged"}
                and (row.logical_course_code in verification_scope or row.source_code in verification_scope)
            ):
                row.verification_status = "pending"
        run.verification_total_courses = len({row.course_shortname for row in preview_rows if row.verification_status == "pending"})
    db.add_all(preview_rows)
    db.commit()
    _matriculation_recount(run, db)
    db.commit()
    db.refresh(run)
    return run


def _run_matriculation_verification(run_id: int, actor: str, session_factory=SessionLocal, moodle=None) -> None:
    with session_factory() as db:
        run = db.get(MatriculationRun, run_id)
        if not run or run.status != "verifying":
            return
        run.verification_started_at = datetime.utcnow()
        db.commit()
        try:
            client = moodle if moodle is not None else MoodleClient(settings)
            rows = db.query(MatriculationRunRow).filter_by(run_id=run.id, verification_status="pending").order_by(MatriculationRunRow.id).all()
            by_course = {}
            for row in rows:
                by_course.setdefault(row.course_shortname, []).append(row)
            for code, course_rows in by_course.items():
                try:
                    usernames = course_usernames(client, code)
                except Exception as exc:
                    for row in course_rows:
                        row.verification_status, row.status = "error", "unverified"
                        row.message = "No se pudo comprobar Moodle. " + safe_error(exc, settings) + " Volver a subir el Lincu para verificar."
                else:
                    for row in course_rows:
                        present = row.upn.strip().lower() in usernames
                        row.verification_status = "verified" if present else "missing"
                        if present:
                            row.message += " Matricula comprobada en Moodle."
                        else:
                            if row.action == "unchanged":
                                row.action = "enroll"
                                row.message = "La matricula guardada localmente ya no existe en Moodle. Pendiente de volver a matricular y agregar al grupo."
                            else:
                                row.message += " No figura matriculado en el curso de destino en Moodle."
                            # The execution resolves identities again; never reuse a deleted Moodle user's ID.
                            row.moodle_user_id, row.moodle_course_id = "", ""
                run.verification_processed_courses += 1
                run.message = f"Comprobando Moodle: {run.verification_processed_courses}/{run.verification_total_courses} cursos. Sin cambios externos."
                db.commit()
            _matriculation_recount(run, db)
            missing = db.query(MatriculationRunRow).filter_by(run_id=run.id, verification_status="missing").count()
            failed = db.query(MatriculationRunRow).filter_by(run_id=run.id, verification_status="error").count()
            run.status = "preview"
            run.message = (
                f"Comprobacion Moodle finalizada para {run.verification_scope}: {missing} matriculas faltantes; "
                f"{failed} filas no verificadas. Revisa la vista previa antes de ejecutar."
            )
        except Exception as exc:
            db.rollback()
            db.query(MatriculationRunRow).filter_by(run_id=run.id, verification_status="pending").update({
                "verification_status": "error", "status": "unverified",
                "message": "Comprobacion incompleta. Volver a subir el Lincu; esta fila no se puede ejecutar.",
            })
            _matriculation_recount(run, db)
            run.status = "completed_with_errors"
            run.message = "Comprobacion Moodle incompleta: " + safe_error(exc, settings)
        finally:
            run.verification_finished_at = datetime.utcnow()
            db.add(OperationLog(actor=actor, action="matriculation_verify", status=run.status,
                                detail=f"Corrida #{run.id}: {run.message}"))
            db.commit()


def _moodle_users_by_upn(moodle: MoodleClient, upns: list[str]) -> dict[str, dict]:
    users_by_upn: dict[str, dict] = {}
    unique_upns = sorted({upn.strip().lower() for upn in upns if upn.strip()})
    for upn_chunk in _chunked(unique_upns, MATRICULATION_USER_LOOKUP_BATCH_SIZE):
        result = moodle.get_users_by_usernames(upn_chunk)
        if result.status != "ok" or not isinstance(result.data, list):
            raise ValueError(f"No se pudieron verificar usuarios Moodle: {result.message}")
        for user in result.data:
            if not isinstance(user, dict):
                continue
            username = str(user.get("username", "")).strip().lower()
            if username:
                users_by_upn[username] = user
    return users_by_upn


def _create_elective_preview(
    content: bytes,
    filename: str,
    actor: str,
    mapping: dict[str, str],
    subject_course_map: dict[str, str],
    default_group: str,
    db: Session,
) -> MatriculationRun:
    records, parse_issues, total_rows, duplicates = parse_elective_excel(
        content,
        settings.crm_domain,
        mapping,
        subject_course_map,
        default_group,
    )
    if not records and not parse_issues:
        raise ValueError("El Excel no contiene filas con datos.")

    moodle = MoodleClient(settings)
    if not moodle.enabled:
        raise ValueError("Moodle no esta configurado para validar la matriculacion.")

    users_by_upn = _moodle_users_by_upn(moodle, [record.upn for record in records])
    courses: dict[str, MoodleResult] = {}
    groups: dict[tuple[str, str], MoodleResult] = {}
    for shortname in sorted({record.course_shortname for record in records}):
        course = moodle.get_course_by_shortname(shortname)
        courses[shortname] = course
        if course.status != "ok" or not isinstance(course.data, dict):
            continue
        course_id = int(course.data.get("id", 0) or 0)
        if course_id <= 0:
            courses[shortname] = MoodleResult("error", "Curso Moodle sin id.")
            continue
        for group_name in sorted(
            {record.group_name for record in records if record.course_shortname == shortname}
        ):
            groups[(shortname, group_name)] = moodle.get_group_in_course_by_name(
                course_id,
                group_name,
            )

    known_contacts = {
        user.upn.lower(): user.contact_email
        for user in db.query(ManagedUser).filter(ManagedUser.contact_email != "").all()
    }
    run = MatriculationRun(
        filename=f"Electiva - {filename}",
        created_by=actor,
        source_type="elective",
        status="preview",
        total_source_rows=total_rows,
        desired_rows=len(records),
        message=(
            "Vista previa Electiva generada. Moodle no se modifica hasta ejecutar. "
            "Este modulo nunca crea usuarios Moodle."
        ),
    )
    db.add(run)
    db.commit()
    db.refresh(run)

    preview_rows: list[MatriculationRunRow] = []
    for record in records:
        course = courses[record.course_shortname]
        user = users_by_upn.get(record.upn.lower())
        first_name = record.first_name
        last_name = record.last_name
        if user:
            first_name = title_name(user.get("firstname", "")) or first_name
            last_name = title_name(user.get("lastname", "")) or last_name
        course_name = record.source_course
        if course.status == "ok" and isinstance(course.data, dict):
            course_name = course_name or str(course.data.get("fullname", "") or "")
        base = {
            "run_id": run.id,
            "document": record.document,
            "upn": record.upn,
            "first_name": first_name,
            "last_name": last_name,
            "email": record.email or known_contacts.get(record.upn.lower(), ""),
            "source_code": record.course_shortname,
            "source_course": course_name,
            "logical_course_code": record.course_shortname.split("-", 1)[0],
            "course_shortname": record.course_shortname,
            "group_name": record.group_name,
        }
        if not user:
            preview_rows.append(
                MatriculationRunRow(
                    action="error",
                    status="preview",
                    message="El usuario no existe en Moodle. Electiva no crea usuarios.",
                    **base,
                )
            )
            continue
        if course.status != "ok" or not isinstance(course.data, dict):
            preview_rows.append(
                MatriculationRunRow(
                    action="error",
                    status="preview",
                    message=f"Curso {record.course_shortname}: {course.message}",
                    **base,
                )
            )
            continue
        group = groups.get((record.course_shortname, record.group_name))
        if group and group.status not in {"ok", "missing"}:
            preview_rows.append(
                MatriculationRunRow(
                    action="error",
                    status="preview",
                    message=f"Grupo {record.group_name}: {group.message}",
                    **base,
                )
            )
            continue
        group_notice = (
            f" Se creara el grupo {record.group_name} si todavia no existe."
            if group and group.status == "missing"
            else ""
        )
        preview_rows.append(
            MatriculationRunRow(
                action="enroll",
                status="preview",
                moodle_user_id=str(user.get("id", "") or ""),
                moodle_course_id=str(course.data.get("id", "") or ""),
                message=f"Matriculacion Electiva pendiente.{group_notice}",
                **base,
            )
        )

    for issue in parse_issues:
        preview_rows.append(
            MatriculationRunRow(
                run_id=run.id,
                action="error",
                status="preview",
                document=issue.document,
                first_name=issue.first_name,
                last_name=issue.last_name,
                source_code=issue.course_shortname,
                source_course=issue.source_course,
                course_shortname=issue.course_shortname,
                logical_course_code=issue.course_shortname.split("-", 1)[0] if issue.course_shortname else "",
                message=f"Fila Excel {issue.row_number}: {issue.message}",
            )
        )

    db.add_all(preview_rows)
    _matriculation_recount(run, db)
    if duplicates:
        run.message += f" Se omitieron {duplicates} duplicados exactos del mismo Excel."
    db.commit()
    db.refresh(run)
    return run


def _friendly_moodle_ok(result_status: str, message: str = "") -> bool:
    text = str(message or "").lower()
    return result_status in {"ok", "enrolled", "added", "created"} or "already" in text or "ya existe" in text


def _moodle_error_detail(result: MoodleResult) -> str:
    detail = str(result.message or "").strip()
    if detail:
        return detail
    return f"Moodle no informo el detalle (estado {result.status or 'desconocido'})."


MATRICULATION_ACTIONS = {"enroll", "group_change", "course_change"}


def _resolve_matriculation_user(
    moodle: MoodleClient,
    row: MatriculationRunRow,
    context: dict | None = None,
    create_if_missing: bool = True,
) -> MoodleResult:
    cache = context.setdefault("users", {}) if context is not None else {}
    cache_key = row.upn.lower()
    if cache_key in cache:
        return cache[cache_key]
    user = moodle.get_user_by_username(row.upn.lower())
    if user.status == "missing" and create_if_missing:
        user = moodle.create_user(
            row.upn.lower(),
            row.first_name or row.upn,
            row.last_name or "-",
            row.upn.lower(),
        )
    if user.status in {"ok", "created"}:
        cache[cache_key] = user
    return user


def _resolve_matriculation_course(
    moodle: MoodleClient,
    shortname: str,
    context: dict | None = None,
) -> MoodleResult:
    cache = context.setdefault("courses", {}) if context is not None else {}
    if shortname in cache:
        return cache[shortname]
    course = moodle.get_course_by_shortname(shortname)
    if course.status == "ok":
        cache[shortname] = course
    return course


def _resolve_group(
    moodle: MoodleClient,
    course_id: int,
    group_name: str,
    context: dict | None = None,
) -> MoodleResult:
    cache = context.setdefault("groups", {}) if context is not None else {}
    cache_key = (course_id, group_name.strip().lower())
    if cache_key in cache:
        return cache[cache_key]
    group = moodle.get_group_in_course_by_name(course_id, group_name)
    if group.status == "missing":
        group = moodle.create_group(course_id, group_name)
    if group.status in {"ok", "created"}:
        cache[cache_key] = group
    return group


def _group_member_ids(result: MoodleResult) -> set[int]:
    member_ids: set[int] = set()
    if result.status != "ok" or not isinstance(result.data, list):
        return member_ids
    for item in result.data:
        if not isinstance(item, dict):
            continue
        user_ids = item.get("userids", [])
        if isinstance(user_ids, list):
            for user_id in user_ids:
                try:
                    member_ids.add(int(user_id))
                except (TypeError, ValueError):
                    continue
        for key in ("userid", "id"):
            try:
                value = int(item.get(key, 0) or 0)
            except (TypeError, ValueError):
                value = 0
            if value > 0:
                member_ids.add(value)
    return member_ids


def _ensure_consult_group_membership(
    moodle: MoodleClient,
    course_id: int,
    group_name: str,
    moodle_user_id: int,
) -> MoodleResult:
    group = _resolve_group(moodle, course_id, group_name)
    if group.status not in {"ok", "created"} or not isinstance(group.data, dict):
        return MoodleResult("error", f"Grupo {group_name}: {_moodle_error_detail(group)}")

    group_id = int(group.data.get("id", 0) or 0)
    if group_id <= 0:
        return MoodleResult("error", f"Grupo {group_name}: Moodle no devolvio un id valido.")

    created = group.status == "created"
    if not created:
        members = moodle.get_group_members(group_id)
        if moodle_user_id in _group_member_ids(members):
            return MoodleResult(
                "exists",
                f"El usuario ya pertenecia al grupo {group_name}.",
                {"group_id": group_id, "group_created": False},
            )

    added = moodle.add_group_member(group_id, moodle_user_id)
    if not _friendly_moodle_ok(added.status, added.message):
        return MoodleResult("error", f"Grupo {group_name}: {_moodle_error_detail(added)}")
    return MoodleResult(
        "added",
        f"Usuario agregado al grupo {group_name}.",
        {"group_id": group_id, "group_created": created},
    )


def _get_course_enrolled_users(
    moodle: MoodleClient,
    course_id: int,
    context: dict | None = None,
) -> dict[str, dict] | None:
    if context is None:
        return None
    cache = context.setdefault("enrolled_users", {})
    if course_id in cache:
        return cache[course_id]
    result = moodle.get_enrolled_users(course_id)
    if result.status != "ok" or not isinstance(result.data, list):
        cache[course_id] = None
        return None
    users = {
        str(user.get("username", "")).strip().lower(): user
        for user in result.data
        if isinstance(user, dict) and str(user.get("username", "")).strip()
    }
    cache[course_id] = users
    return users


def _chunked(values: list, size: int):
    for index in range(0, len(values), size):
        yield values[index : index + size]


def _prime_matriculation_users(
    moodle: MoodleClient,
    rows: list[MatriculationRunRow],
    context: dict,
) -> None:
    cache = context.setdefault("users", {})
    usernames: list[str] = []
    seen: set[str] = set()
    for row in rows:
        username = row.upn.strip().lower()
        if username and username not in cache and username not in seen:
            usernames.append(username)
            seen.add(username)

    for usernames_chunk in _chunked(usernames, MATRICULATION_USER_LOOKUP_BATCH_SIZE):
        result = moodle.get_users_by_usernames(usernames_chunk)
        if result.status != "ok" or not isinstance(result.data, list):
            continue
        for user in result.data:
            if not isinstance(user, dict):
                continue
            username = str(user.get("username", "")).strip().lower()
            if username:
                cache[username] = MoodleResult("ok", "Usuario Moodle encontrado.", user)


def _fast_enrolment_candidate(
    moodle: MoodleClient,
    row: MatriculationRunRow,
    context: dict,
) -> dict | None:
    course = _resolve_matriculation_course(moodle, row.course_shortname, context)
    if course.status != "ok" or not isinstance(course.data, dict):
        return None
    course_id = int(course.data.get("id", 0) or 0)
    if course_id <= 0:
        return None

    enrolled_users = _get_course_enrolled_users(moodle, course_id, context)
    if enrolled_users is None:
        return None
    existing_user = enrolled_users.get(row.upn.lower())
    already_enrolled = isinstance(existing_user, dict)
    if already_enrolled:
        user = existing_user
    else:
        user_result = context.setdefault("users", {}).get(row.upn.lower())
        user = user_result.data if isinstance(user_result, MoodleResult) else None
    if not isinstance(user, dict):
        return None
    moodle_user_id = int(user.get("id", 0) or 0)
    if moodle_user_id <= 0:
        return None

    group = _resolve_group(moodle, course_id, row.group_name, context)
    if group.status not in {"ok", "created"} or not isinstance(group.data, dict):
        return None
    group_id = int(group.data.get("id", 0) or 0)
    if group_id <= 0:
        return None
    user_groups = user.get("groups")
    already_grouped = isinstance(user_groups, list) and any(
        int(item.get("id", 0) or 0) == group_id
        for item in user_groups
        if isinstance(item, dict)
    )
    return {
        "row": row,
        "user": user,
        "moodle_user_id": moodle_user_id,
        "course": course.data,
        "course_id": course_id,
        "group_id": group_id,
        "already_enrolled": already_enrolled,
        "already_grouped": already_grouped,
        "enrolled_users": enrolled_users,
    }


def _complete_fast_enrolment(db: Session, candidate: dict) -> None:
    row: MatriculationRunRow = candidate["row"]
    row.status = "done"
    row.moodle_user_id = str(candidate["moodle_user_id"])
    row.moodle_course_id = str(candidate["course_id"])
    if candidate["already_enrolled"] and candidate["already_grouped"]:
        row.message = "Ya estaba matriculado y en el grupo; estado local sincronizado."
    elif candidate["already_enrolled"]:
        row.message = "Ya estaba matriculado; agregado al grupo."
    else:
        row.message = "Matriculado y agregado al grupo."
    course_name = str(candidate["course"].get("fullname", "") or "")
    _upsert_matriculation_state(
        db,
        row,
        str(candidate["moodle_user_id"]),
        str(candidate["course_id"]),
        course_name,
    )
    _sync_managed_matriculation(
        db,
        row,
        str(candidate["moodle_user_id"]),
        str(candidate["course_id"]),
        course_name,
    )


def _advance_matriculation_progress(run: MatriculationRun, db: Session, count: int) -> None:
    if count <= 0:
        return
    run.execution_processed_rows += count
    db.commit()


def _process_fast_matriculation_enrolments(
    db: Session,
    moodle: MoodleClient,
    run: MatriculationRun,
    rows: list[MatriculationRunRow],
    context: dict,
) -> list[MatriculationRunRow]:
    """Batch only simple enrolments; complex changes retain the individual path."""
    if not rows:
        return []

    _prime_matriculation_users(moodle, rows, context)
    candidates: list[dict] = []
    fallback: dict[int, MatriculationRunRow] = {}
    for row in rows:
        candidate = _fast_enrolment_candidate(moodle, row, context)
        if candidate is None:
            fallback[row.id] = row
        else:
            candidates.append(candidate)

    failed_candidate_ids: set[int] = set()
    candidates_by_course: dict[int, list[dict]] = {}
    for candidate in candidates:
        if not candidate["already_enrolled"]:
            candidates_by_course.setdefault(candidate["course_id"], []).append(candidate)

    for course_candidates in candidates_by_course.values():
        for candidate_chunk in _chunked(course_candidates, MATRICULATION_MOODLE_BATCH_SIZE):
            for candidate in candidate_chunk:
                candidate["row"].status = "running"
            db.commit()
            enrolled = moodle.enrol_users(
                [
                    (candidate["moodle_user_id"], candidate["course_id"], "student")
                    for candidate in candidate_chunk
                ]
            )
            if _friendly_moodle_ok(enrolled.status, enrolled.message):
                for candidate in candidate_chunk:
                    candidate["enrolled_users"][candidate["row"].upn.lower()] = candidate["user"]
                    candidate["row"].status = "preview"
            else:
                for candidate in candidate_chunk:
                    failed_candidate_ids.add(candidate["row"].id)
                    fallback[candidate["row"].id] = candidate["row"]
                    candidate["row"].status = "preview"
            db.commit()

    already_complete = [
        candidate
        for candidate in candidates
        if candidate["row"].id not in failed_candidate_ids and candidate["already_grouped"]
    ]
    for candidate in already_complete:
        _complete_fast_enrolment(db, candidate)
    _advance_matriculation_progress(run, db, len(already_complete))

    candidates_by_group: dict[int, list[dict]] = {}
    for candidate in candidates:
        if candidate["row"].id in failed_candidate_ids or candidate["already_grouped"]:
            continue
        candidates_by_group.setdefault(candidate["group_id"], []).append(candidate)

    for group_candidates in candidates_by_group.values():
        for candidate_chunk in _chunked(group_candidates, MATRICULATION_MOODLE_BATCH_SIZE):
            for candidate in candidate_chunk:
                candidate["row"].status = "running"
            db.commit()
            grouped = moodle.add_group_members(
                [
                    (candidate["group_id"], candidate["moodle_user_id"])
                    for candidate in candidate_chunk
                ]
            )
            if _friendly_moodle_ok(grouped.status, grouped.message):
                for candidate in candidate_chunk:
                    _complete_fast_enrolment(db, candidate)
                _advance_matriculation_progress(run, db, len(candidate_chunk))
            else:
                for candidate in candidate_chunk:
                    fallback[candidate["row"].id] = candidate["row"]
                    candidate["row"].status = "preview"
                db.commit()

    return sorted(fallback.values(), key=lambda row: row.id)


def _upsert_matriculation_state(
    db: Session,
    row: MatriculationRunRow,
    moodle_user_id: str,
    moodle_course_id: str,
    course_name: str = "",
) -> None:
    state = (
        db.query(MatriculationState)
        .filter(
            MatriculationState.document == row.document,
            MatriculationState.course_shortname == row.course_shortname,
        )
        .first()
    )
    if not state:
        state = MatriculationState(
            document=row.document,
            course_shortname=row.course_shortname,
        )
        db.add(state)
    state.upn = row.upn
    state.first_name = row.first_name
    state.last_name = row.last_name
    state.email = row.email
    state.logical_course_code = row.logical_course_code or row.course_shortname.split("-", 1)[0]
    state.group_name = row.group_name
    state.moodle_user_id = moodle_user_id
    state.moodle_course_id = moodle_course_id
    state.course_name = course_name
    state.active = True


def _sync_managed_matriculation(
    db: Session,
    row: MatriculationRunRow,
    moodle_user_id: str,
    moodle_course_id: str,
    course_name: str = "",
) -> None:
    imported = ImportedRow(
        document=row.document,
        first_name=row.first_name,
        last_name=row.last_name,
        full_name=f"{row.first_name} {row.last_name}".strip(),
        email=row.email,
        upn=row.upn,
    )
    user = upsert_managed_user(db, imported, "student", moodle_user_id=moodle_user_id)
    upsert_managed_enrollment(
        db,
        user,
        course_shortname=row.course_shortname,
        role="student",
        course_name=course_name,
        moodle_course_id=moodle_course_id,
    )


def _process_matriculation_enroll(
    db: Session,
    moodle: MoodleClient,
    row: MatriculationRunRow,
    context: dict | None = None,
    create_if_missing: bool = True,
) -> None:
    course = _resolve_matriculation_course(moodle, row.course_shortname, context)
    if course.status != "ok" or not isinstance(course.data, dict):
        row.status = "error"
        row.message = f"Curso {row.course_shortname}: {_moodle_error_detail(course)}"
        return
    course_id = int(course.data.get("id", 0) or 0)
    if course_id <= 0:
        row.status = "error"
        row.message = f"Curso {row.course_shortname} sin id."
        return

    enrolled_users = _get_course_enrolled_users(moodle, course_id, context)
    existing_user = enrolled_users.get(row.upn.lower()) if enrolled_users is not None else None
    already_enrolled = isinstance(existing_user, dict)
    if already_enrolled:
        user = MoodleResult("ok", "Usuario ya matriculado.", existing_user)
        if context is not None:
            context.setdefault("users", {})[row.upn.lower()] = user
    else:
        user = _resolve_matriculation_user(
            moodle,
            row,
            context,
            create_if_missing=create_if_missing,
        )

    if user.status not in {"ok", "created"} or not isinstance(user.data, dict):
        row.status = "error"
        row.message = f"Usuario Moodle: {_moodle_error_detail(user)}"
        return
    moodle_user_id = int(user.data.get("id", 0) or 0)
    if moodle_user_id <= 0:
        row.status = "error"
        row.message = "Usuario Moodle sin id."
        return

    if not already_enrolled:
        enrollment = moodle.enrol_user(moodle_user_id, course_id, "student")
        if not _friendly_moodle_ok(enrollment.status, enrollment.message):
            row.status = "error"
            row.message = f"Matriculacion: {_moodle_error_detail(enrollment)}"
            return
        if enrolled_users is not None:
            enrolled_users[row.upn.lower()] = user.data

    group = _resolve_group(moodle, course_id, row.group_name, context)
    if group.status not in {"ok", "created"} or not isinstance(group.data, dict):
        row.status = "error"
        row.message = f"Grupo {row.group_name}: {_moodle_error_detail(group)}"
        return
    group_id = int(group.data.get("id", 0) or 0)
    user_groups = user.data.get("groups") if isinstance(user.data, dict) else None
    already_grouped = isinstance(user_groups, list) and any(
        int(item.get("id", 0) or 0) == group_id
        for item in user_groups
        if isinstance(item, dict)
    )
    if not already_grouped:
        grouped = moodle.add_group_member(group_id, moodle_user_id)
        if not _friendly_moodle_ok(grouped.status, grouped.message):
            row.status = "error"
            row.message = f"Grupo {row.group_name}: {_moodle_error_detail(grouped)}"
            return
        if isinstance(user_groups, list):
            user_groups.append({"id": group_id, "name": row.group_name})

    row.status = "done"
    row.moodle_user_id = str(moodle_user_id)
    row.moodle_course_id = str(course_id)
    if already_enrolled and already_grouped:
        row.message = "Ya estaba matriculado y en el grupo; estado local sincronizado."
    elif already_enrolled:
        row.message = "Ya estaba matriculado; agregado al grupo."
    else:
        row.message = "Matriculado y agregado al grupo."
    course_name = str(course.data.get("fullname", "") or "")
    _upsert_matriculation_state(db, row, str(moodle_user_id), str(course_id), course_name)
    _sync_managed_matriculation(db, row, str(moodle_user_id), str(course_id), course_name)


def _process_matriculation_group_change(
    db: Session,
    moodle: MoodleClient,
    row: MatriculationRunRow,
    context: dict | None = None,
) -> None:
    _process_matriculation_enroll(db, moodle, row, context)
    if row.status != "done":
        return
    if not row.previous_group_name:
        row.message = "Grupo actualizado."
        return
    old_group = moodle.get_group_in_course_by_name(int(row.moodle_course_id), row.previous_group_name)
    if old_group.status == "ok" and isinstance(old_group.data, dict):
        old_group_id = int(old_group.data.get("id", 0) or 0)
        if old_group_id > 0:
            removed = moodle.delete_group_member(old_group_id, int(row.moodle_user_id))
            if removed.status == "removed":
                row.message = f"Grupo actualizado: {row.previous_group_name} -> {row.group_name}."
            else:
                row.message = f"Grupo actualizado; aviso al quitar de {row.previous_group_name}: {_moodle_error_detail(removed)}"
    else:
        row.message = f"Grupo actualizado; no se encontro grupo anterior {row.previous_group_name}."


def _process_matriculation_course_change(
    db: Session,
    moodle: MoodleClient,
    row: MatriculationRunRow,
    context: dict | None = None,
) -> None:
    previous_shortname = row.previous_course_shortname.strip()
    if not previous_shortname or previous_shortname == row.course_shortname:
        row.status = "error"
        row.message = "Cambio de curso sin curso anterior valido."
        return

    _process_matriculation_enroll(db, moodle, row, context)
    if row.status != "done":
        return

    user_id = int(row.moodle_user_id or 0)
    previous_course_id = int(row.previous_moodle_course_id or 0)
    if previous_course_id <= 0:
        previous_course = _resolve_matriculation_course(moodle, previous_shortname, context)
        if previous_course.status == "ok" and isinstance(previous_course.data, dict):
            previous_course_id = int(previous_course.data.get("id", 0) or 0)
        else:
            row.status = "error"
            row.message = (
                f"Matriculado en {row.course_shortname}, pero no se pudo resolver "
                f"el curso anterior {previous_shortname}: {_moodle_error_detail(previous_course)}"
            )
            return

    if user_id <= 0 or previous_course_id <= 0:
        row.status = "error"
        row.message = "Matricula nueva creada, pero faltan ids para retirar el curso anterior."
        return

    if row.previous_group_name:
        old_group = moodle.get_group_in_course_by_name(previous_course_id, row.previous_group_name)
        if old_group.status == "ok" and isinstance(old_group.data, dict):
            old_group_id = int(old_group.data.get("id", 0) or 0)
            if old_group_id > 0:
                moodle.delete_group_member(old_group_id, user_id)

    removed = moodle.unenrol_user(user_id, previous_course_id, "student")
    if removed.status not in {"unenrolled", "missing"} and not _friendly_moodle_ok(
        removed.status,
        removed.message,
    ):
        row.status = "error"
        row.message = (
            f"Matriculado en {row.course_shortname}, pero no se pudo retirar "
            f"{previous_shortname}: {_moodle_error_detail(removed)}"
        )
        return

    previous_state = (
        db.query(MatriculationState)
        .filter(
            MatriculationState.document == row.document,
            MatriculationState.course_shortname == previous_shortname,
        )
        .first()
    )
    if previous_state:
        previous_state.active = False

    managed_user = db.query(ManagedUser).filter(ManagedUser.upn == row.upn).first()
    if managed_user:
        previous_enrollment = (
            db.query(ManagedEnrollment)
            .filter(
                ManagedEnrollment.user_id == managed_user.id,
                ManagedEnrollment.course_shortname == previous_shortname,
            )
            .first()
        )
        if previous_enrollment:
            previous_enrollment.active = False

    row.status = "done"
    row.message = f"Curso actualizado: {previous_shortname} -> {row.course_shortname}."


def _process_matriculation_unenroll(
    db: Session,
    moodle: MoodleClient,
    row: MatriculationRunRow,
    context: dict | None = None,
) -> None:
    user_id = int(row.moodle_user_id or 0)
    course_id = int(row.moodle_course_id or 0)
    if user_id <= 0:
        user = _resolve_matriculation_user(moodle, row, context, create_if_missing=False)
        if user.status == "missing":
            user_id = 0
        elif user.status == "ok" and isinstance(user.data, dict):
            user_id = int(user.data.get("id", 0) or 0)
        else:
            row.status = "error"
            row.message = f"Usuario Moodle: {_moodle_error_detail(user)}"
            return
    if course_id <= 0:
        course = _resolve_matriculation_course(moodle, row.course_shortname, context)
        if course.status == "ok" and isinstance(course.data, dict):
            course_id = int(course.data.get("id", 0) or 0)
        else:
            row.status = "error"
            row.message = f"Curso {row.course_shortname}: {_moodle_error_detail(course)}"
            return

    if user_id > 0 and row.previous_group_name:
        old_group = moodle.get_group_in_course_by_name(course_id, row.previous_group_name)
        if old_group.status == "ok" and isinstance(old_group.data, dict):
            old_group_id = int(old_group.data.get("id", 0) or 0)
            if old_group_id > 0:
                moodle.delete_group_member(old_group_id, user_id)

    if user_id > 0 and course_id > 0:
        result = moodle.unenrol_user(user_id, course_id, "student")
        if result.status not in {"unenrolled", "missing"} and not _friendly_moodle_ok(result.status, result.message):
            row.status = "error"
            row.message = f"Desmatriculacion: {_moodle_error_detail(result)}"
            return

    state = (
        db.query(MatriculationState)
        .filter(
            MatriculationState.document == row.document,
            MatriculationState.course_shortname == row.course_shortname,
        )
        .first()
    )
    if state:
        state.active = False

    managed_user = db.query(ManagedUser).filter(ManagedUser.upn == row.upn).first()
    if managed_user:
        enrollment = (
            db.query(ManagedEnrollment)
            .filter(
                ManagedEnrollment.user_id == managed_user.id,
                ManagedEnrollment.course_shortname == row.course_shortname,
            )
            .first()
        )
        if enrollment:
            enrollment.active = False

    row.status = "done"
    row.message = "Desmatriculado porque no figura en el Lincu actual."


def _run_matriculation_job(run_id: int, actor: str, row_ids: list[int]) -> None:
    db = SessionLocal()
    try:
        run = db.query(MatriculationRun).filter(MatriculationRun.id == run_id).first()
        if not run:
            return
        if run.status != "queued":
            return
        run.status = "running"
        run.started_at = datetime.utcnow()
        run.finished_at = None
        run.execution_processed_rows = 0
        run.message = "Ejecutando cambios en Moodle."
        db.commit()

        moodle = MoodleClient(settings)
        context: dict = {}
        allow_user_creation = run.source_type != "elective"
        rows = (
            db.query(MatriculationRunRow)
            .filter(
                MatriculationRunRow.run_id == run_id,
                MatriculationRunRow.id.in_(row_ids),
                MatriculationRunRow.action.in_(MATRICULATION_ACTIONS),
                MatriculationRunRow.status.in_({"preview", "error"}),
            )
            .order_by(MatriculationRunRow.id.asc())
            .all()
        )
        run.execution_total_rows = len(rows)
        db.commit()
        simple_enrolments = [row for row in rows if row.action == "enroll"]
        fallback_enrolments = _process_fast_matriculation_enrolments(
            db,
            moodle,
            run,
            simple_enrolments,
            context,
        )
        slow_rows = [row for row in rows if row.action != "enroll"] + fallback_enrolments
        for row in sorted(slow_rows, key=lambda current: current.id):
            row.status = "running"
            db.commit()
            try:
                if row.action == "enroll":
                    _process_matriculation_enroll(
                        db,
                        moodle,
                        row,
                        context,
                        create_if_missing=allow_user_creation,
                    )
                elif row.action == "group_change":
                    _process_matriculation_group_change(db, moodle, row, context)
                elif row.action == "course_change":
                    _process_matriculation_course_change(db, moodle, row, context)
                else:
                    row.status = "error"
                    row.message = f"Accion desconocida: {row.action}"
            except Exception as exc:
                row.status = "error"
                row.message = str(exc)
            run.execution_processed_rows += 1
            db.commit()

        _matriculation_recount(run, db)
        pending_rows = (
            db.query(MatriculationRunRow)
            .filter(
                MatriculationRunRow.run_id == run.id,
                MatriculationRunRow.action.in_(MATRICULATION_ACTIONS),
                MatriculationRunRow.status == "preview",
            )
            .count()
        )
        if pending_rows:
            run.status = "preview"
        else:
            run.status = "completed_with_errors" if run.error_rows else "completed"
        run.finished_at = datetime.utcnow()
        elapsed = max(0.0, (run.finished_at - run.started_at).total_seconds())
        if pending_rows:
            run.message = (
                f"Lote finalizado: {len(rows)} filas en {elapsed:.1f}s. "
                f"Quedan {pending_rows} filas pendientes."
            )
        else:
            run.message = f"Matriculacion diaria finalizada: {len(rows)} filas en {elapsed:.1f}s."
        db.add(
            OperationLog(
                actor=actor,
                action="matriculation_lincu",
                status="ok" if not run.error_rows else "error",
                detail=(
                    f"Corrida {run.id}: lote={len(rows)}, altas={run.enroll_rows}, "
                    f"cambios_grupo={run.group_change_rows}, cambios_curso={run.course_change_rows}, "
                    f"bajas={run.unenroll_rows}, "
                    f"errores={run.error_rows}, segundos={elapsed:.1f}"
                ),
            )
        )
        db.commit()
    except Exception as exc:
        run = db.query(MatriculationRun).filter(MatriculationRun.id == run_id).first()
        if run:
            run.status = "failed"
            run.finished_at = datetime.utcnow()
            run.message = str(exc)
            db.commit()
    finally:
        db.close()


def _queue_matriculation_execution(
    run: MatriculationRun,
    db: Session,
    requested_row_ids: list[int] | None = None,
) -> list[int]:
    if run.status in {"queued", "running", "verifying"}:
        raise ValueError("La corrida se esta ejecutando o comprobando contra Moodle. Espera a que termine.")
    query = db.query(MatriculationRunRow.id).filter(
        MatriculationRunRow.run_id == run.id,
        MatriculationRunRow.action.in_(MATRICULATION_ACTIONS),
        MatriculationRunRow.status.in_({"preview", "error"}),
    )
    if requested_row_ids is not None:
        clean_ids = sorted({int(row_id) for row_id in requested_row_ids if int(row_id) > 0})
        if not clean_ids:
            raise ValueError("Selecciona al menos una fila.")
        query = query.filter(MatriculationRunRow.id.in_(clean_ids))
    row_ids = [row_id for (row_id,) in query.order_by(MatriculationRunRow.id.asc()).all()]
    if not row_ids:
        raise ValueError("No hay filas pendientes para ejecutar.")

    run.status = "queued"
    run.execution_total_rows = len(row_ids)
    run.execution_processed_rows = 0
    run.started_at = None
    run.finished_at = None
    run.message = f"En cola: {len(row_ids)} filas."
    db.commit()
    return row_ids


def _teacher_upn_from_query(value: str, domain: str) -> str:
    raw = str(value or "").strip().lower()
    if not raw:
        return ""
    if "@" in raw:
        local_part, separator, supplied_domain = raw.partition("@")
        if (
            not separator
            or supplied_domain != str(domain or "").strip().lower()
            or not re.fullmatch(r"[a-z0-9](?:[a-z0-9._-]*[a-z0-9])?", local_part)
        ):
            return ""
        return f"{local_part}@{supplied_domain}"
    local_part = normalize_name(raw)
    return f"{local_part}@{domain}" if local_part else ""


def _teacher_lookup_key(value: str, domain: str) -> str:
    """Preserve an explicit institutional UPN; normalize only abbreviated keys."""
    upn = _teacher_upn_from_query(value, domain)
    if not upn:
        return ""
    return upn if "@" in str(value or "") else normalize_name(value)


def _redirect_consult(
    dni: str = "",
    teacher_key: str = "",
    message: str = "",
    error: str = "",
) -> RedirectResponse:
    params: list[str] = []
    if dni:
        params.append(f"dni={quote_plus(dni)}")
    if teacher_key:
        params.append(f"teacher_key={quote_plus(teacher_key)}")
    if message:
        params.append(f"message={quote_plus(message)}")
    if error:
        params.append(f"error={quote_plus(error)}")
    query = f"?{'&'.join(params)}" if params else ""
    return RedirectResponse(f"/consultas{query}", status_code=303)


def _managed_users_redirect_url(
    person_type: str,
    *,
    page: int = 1,
    per_page: int = MANAGED_USERS_PAGE_SIZE,
    dni: str = "",
    name_query: str = "",
    course_shortname: str = "",
    course_query: str = "",
    message: str = "",
    error: str = "",
) -> str:
    params: dict[str, str | int] = {}
    if page > 1:
        params["page"] = page
    if per_page != MANAGED_USERS_PAGE_SIZE:
        params["per_page"] = per_page
    if dni:
        params["dni"] = dni
    if name_query:
        params["name_query"] = name_query
    if course_shortname:
        params["course_shortname"] = course_shortname
    if course_query:
        params["course_query"] = course_query
    if message:
        params["message"] = message
    if error:
        params["error"] = error

    query = urlencode(params)
    return f"/users/{person_type}" + (f"?{query}" if query else "")


def _pagination_window(current_page: int, total_pages: int, radius: int = 2) -> list[int | None]:
    if total_pages <= 1:
        return []

    pages = {1, total_pages}
    for value in range(max(1, current_page - radius), min(total_pages, current_page + radius) + 1):
        pages.add(value)

    ordered = sorted(pages)
    result: list[int | None] = []
    previous = 0
    for value in ordered:
        if previous and value - previous > 1:
            result.append(None)
        result.append(value)
        previous = value
    return result


def _deactivate_local_enrollment_state(
    db: Session,
    user: ManagedUser,
    enrollment: ManagedEnrollment,
) -> None:
    enrollment.active = False
    identity_filter = MatriculationState.upn == user.upn
    if user.document:
        identity_filter = or_(identity_filter, MatriculationState.document == user.document)
    db.query(MatriculationState).filter(
        identity_filter,
        MatriculationState.course_shortname == enrollment.course_shortname,
        MatriculationState.active.is_(True),
    ).update({"active": False}, synchronize_session=False)


def _deactivate_local_user_state(db: Session, user: ManagedUser) -> None:
    user.active = False
    db.query(ManagedEnrollment).filter(ManagedEnrollment.user_id == user.id).update(
        {"active": False},
        synchronize_session=False,
    )
    identity_filter = MatriculationState.upn == user.upn
    if user.document:
        identity_filter = or_(identity_filter, MatriculationState.document == user.document)
    db.query(MatriculationState).filter(identity_filter).update(
        {"active": False},
        synchronize_session=False,
    )


def backfill_managed_users(db: Session) -> None:
    exists = db.query(ManagedUser.id).first()
    if exists:
        return
    people = db.query(ImportedPerson).order_by(ImportedPerson.id.asc()).all()
    for person in people:
        row = ImportedRow(
            document=person.document,
            first_name=person.first_name,
            last_name=person.last_name,
            full_name=person.full_name or f"{person.first_name} {person.last_name}".strip(),
            email=person.email,
            upn=person.upn,
        )
        managed_user = upsert_managed_user(
            db,
            row,
            person_type=person.person_type,
            azure_user_id=person.azure_user_id,
            moodle_user_id=person.moodle_user_id,
        )
        batch = db.query(ImportBatch).filter(ImportBatch.id == person.batch_id).first()
        if batch:
            for course_code in parse_course_shortnames(batch.course_shortname):
                upsert_managed_enrollment(
                    db,
                    managed_user,
                    course_shortname=course_code,
                    role=person.person_type,
                )


@app.get("/", response_class=HTMLResponse)
def root(request: Request):
    require_user(request)
    return RedirectResponse("/consultas", status_code=303)


@app.get("/imports", response_class=HTMLResponse)
def dashboard(request: Request, db: Session = Depends(get_db)):
    username = require_user(request)
    return render_dashboard(request, db, username)


@app.get("/consultas", response_class=HTMLResponse)
def consultas(
    request: Request,
    dni: str = "",
    teacher_key: str = "",
    db: Session = Depends(get_db),
):
    username = require_user(request)
    cleaned_dni = clean_document(dni)
    cleaned_teacher_key = _teacher_lookup_key(teacher_key, settings.crm_domain)
    query_error = ""
    if dni.strip() and not cleaned_dni:
        query_error = "DNI invalido."
    if teacher_key.strip() and not cleaned_teacher_key:
        query_error = "Identificador de docente invalido."
    if cleaned_dni and cleaned_teacher_key:
        query_error = "Ingresa solo DNI o solo identificador de docente."

    course_rows = (
        db.query(ManagedEnrollment.course_shortname)
        .filter(ManagedEnrollment.active.is_(True))
        .distinct()
        .order_by(ManagedEnrollment.course_shortname.asc())
        .all()
    )
    available_courses = [row[0] for row in course_rows if row[0]]

    result: dict[str, object] | None = None
    person_type = ""
    lookup_key = ""
    if cleaned_dni and not cleaned_teacher_key:
        person_type = "student"
        lookup_key = cleaned_dni
        upn = build_student_upn(cleaned_dni, settings.crm_domain)
    elif cleaned_teacher_key and not cleaned_dni:
        person_type = "teacher"
        lookup_key = cleaned_teacher_key
        upn = _teacher_upn_from_query(cleaned_teacher_key, settings.crm_domain)
    else:
        upn = ""

    if upn:
        graph = GraphClient(settings)
        azure = graph.get_user(upn)
        azure_exists = azure.status == "exists"

        managed_user = (
            db.query(ManagedUser)
            .filter(ManagedUser.upn == upn, ManagedUser.active.is_(True))
            .first()
        )
        managed_enrollments = []
        if managed_user:
            managed_enrollments = (
                db.query(ManagedEnrollment)
                .filter(ManagedEnrollment.user_id == managed_user.id, ManagedEnrollment.active.is_(True))
                .order_by(ManagedEnrollment.course_shortname.asc())
                .all()
            )

        moodle = MoodleClient(settings)
        moodle_user = moodle.get_user_by_username(upn.lower())
        moodle_exists = moodle_user.status == "ok" and isinstance(moodle_user.data, dict)
        moodle_courses: list[dict[str, str]] = []
        moodle_courses_error = ""
        if moodle_exists:
            moodle_user_id = int(moodle_user.data.get("id", 0))
            if moodle_user_id > 0:
                courses = moodle.get_user_courses(moodle_user_id)
                if courses.status == "ok" and isinstance(courses.data, list):
                    for course in courses.data:
                        moodle_courses.append(
                            {
                                "shortname": str(course.get("shortname", "") or ""),
                                "fullname": str(course.get("fullname", "") or ""),
                            }
                        )
                else:
                    moodle_courses_error = courses.message

        enrolled_shortnames = {course["shortname"] for course in moodle_courses if course.get("shortname")}
        result = {
            "person_type": person_type,
            "lookup_key": lookup_key,
            "upn": upn,
            "azure_status": azure.status,
            "azure_message": azure.message,
            "azure_exists": azure_exists,
            "moodle_exists": moodle_exists,
            "moodle_message": moodle_user.message,
            "moodle_courses": moodle_courses,
            "moodle_courses_error": moodle_courses_error,
            "enrolled_shortnames": enrolled_shortnames,
            "managed_user": managed_user,
            "managed_enrollments": managed_enrollments,
        }

    return templates.TemplateResponse(
        "consultas.html",
        {
            "request": request,
            "username": username,
            "dni": cleaned_dni or dni.strip(),
            "teacher_key": cleaned_teacher_key or teacher_key.strip(),
            "result": result,
            "available_courses": available_courses,
            "message": request.query_params.get("message", ""),
            "error": request.query_params.get("error", "") or query_error,
        },
    )


@app.post("/consultas/create-account")
def consultas_create_account(
    request: Request,
    person_type: str = Form("student"),
    dni: str = Form(""),
    teacher_key: str = Form(""),
    first_name: str = Form(""),
    last_name: str = Form(""),
    contact_email: str = Form(""),
    db: Session = Depends(get_db),
):
    username = require_user(request)
    if person_type not in {"student", "teacher"}:
        return _redirect_consult(error="Tipo de persona invalido.")

    cleaned_dni = clean_document(dni)
    cleaned_teacher_key = _teacher_lookup_key(teacher_key, settings.crm_domain)
    if person_type == "student" and not cleaned_dni:
        return _redirect_consult(dni=dni, error="DNI invalido.")
    if person_type == "teacher" and not cleaned_teacher_key:
        return _redirect_consult(teacher_key=teacher_key, error="Identificador de docente invalido.")

    first_name = title_name(first_name)
    last_name = title_name(last_name)
    if not first_name or not last_name:
        return _redirect_consult(
            dni=cleaned_dni,
            teacher_key=cleaned_teacher_key,
            error="Nombre y apellido son obligatorios para crear la cuenta.",
        )

    if person_type == "student":
        upn = build_student_upn(cleaned_dni, settings.crm_domain)
    else:
        upn = _teacher_upn_from_query(cleaned_teacher_key, settings.crm_domain)

    graph = GraphClient(settings)
    existing = graph.get_user(upn)
    if existing.status == "exists":
        return _redirect_consult(
            dni=cleaned_dni,
            teacher_key=cleaned_teacher_key,
            message="La cuenta ya existe en Microsoft 365.",
        )
    if existing.status == "error":
        return _redirect_consult(
            dni=cleaned_dni,
            teacher_key=cleaned_teacher_key,
            error=f"Error consultando Microsoft 365: {existing.message}",
        )

    if settings.crm_dry_run:
        return _redirect_consult(
            dni=cleaned_dni,
            teacher_key=cleaned_teacher_key,
            message="Simulacion activa: no se creo la cuenta.",
        )

    password = generate_temporary_password()
    created = graph.create_user(upn, first_name, last_name, password)
    if created.status not in {"created", "exists"}:
        return _redirect_consult(
            dni=cleaned_dni,
            teacher_key=cleaned_teacher_key,
            error=f"No se pudo crear la cuenta: {created.message}",
        )

    azure_user_id = ""
    if isinstance(created.data, dict):
        azure_user_id = str(created.data.get("id", "") or "")

    license_target = azure_user_id or upn
    licensed = graph.assign_license(license_target, person_type)
    if licensed.status not in {"licensed", "skipped"}:
        return _redirect_consult(
            dni=cleaned_dni,
            teacher_key=cleaned_teacher_key,
            error=f"Cuenta creada, pero fallo la licencia: {licensed.message}",
        )

    email_value = contact_email.strip().lower()
    row = ImportedRow(
        document=cleaned_dni if person_type == "student" else "",
        first_name=first_name,
        last_name=last_name,
        full_name=f"{first_name} {last_name}".strip(),
        email=email_value,
        upn=upn,
    )
    upsert_managed_user(db, row, person_type=person_type, azure_user_id=azure_user_id)

    mail_status = "sin envio"
    if email_value:
        mail = Mailer(settings).send_credentials(email_value, upn, password, full_name=row.full_name)
        mail_status = mail.status

    db.add(
        OperationLog(
            actor=username,
            action="consult_create_account",
            status="ok",
            detail=f"{upn} creado desde consultas (correo: {mail_status})",
        )
    )
    db.commit()
    return _redirect_consult(
        dni=cleaned_dni,
        teacher_key=cleaned_teacher_key,
        message=f"Cuenta creada. Licencia: {licensed.status}. Correo: {mail_status}.",
    )


@app.post("/consultas/enroll")
def consultas_enroll(
    request: Request,
    person_type: str = Form("student"),
    dni: str = Form(""),
    teacher_key: str = Form(""),
    course_shortname: str = Form(""),
    group_name: str = Form(""),
    first_name: str = Form(""),
    last_name: str = Form(""),
    contact_email: str = Form(""),
    db: Session = Depends(get_db),
):
    username = require_user(request)
    if person_type not in {"student", "teacher"}:
        return _redirect_consult(error="Tipo de persona invalido.")

    cleaned_dni = clean_document(dni)
    cleaned_teacher_key = _teacher_lookup_key(teacher_key, settings.crm_domain)
    if person_type == "student" and not cleaned_dni:
        return _redirect_consult(dni=dni, error="DNI invalido.")
    if person_type == "teacher" and not cleaned_teacher_key:
        return _redirect_consult(teacher_key=teacher_key, error="Identificador de docente invalido.")

    course_codes = parse_course_shortnames(course_shortname)
    if not course_codes:
        return _redirect_consult(
            dni=cleaned_dni,
            teacher_key=cleaned_teacher_key,
            error="Informar al menos un curso (shortname).",
        )
    group_name = " ".join(group_name.strip().split())
    if len(group_name) > 120:
        return _redirect_consult(
            dni=cleaned_dni,
            teacher_key=cleaned_teacher_key,
            error="El nombre del grupo no puede superar 120 caracteres.",
        )

    if person_type == "student":
        upn = build_student_upn(cleaned_dni, settings.crm_domain)
    else:
        upn = _teacher_upn_from_query(cleaned_teacher_key, settings.crm_domain)

    graph = GraphClient(settings)
    existing = graph.get_user(upn)
    if existing.status != "exists":
        return _redirect_consult(
            dni=cleaned_dni,
            teacher_key=cleaned_teacher_key,
            error="La cuenta Microsoft no existe. Primero crear la cuenta.",
        )

    moodle = MoodleClient(settings)
    moodle_user = moodle.get_user_by_username(upn.lower())
    if moodle_user.status == "missing":
        managed_user = db.query(ManagedUser).filter(ManagedUser.upn == upn, ManagedUser.active.is_(True)).first()
        resolved_first = title_name(first_name) or (managed_user.first_name if managed_user else "")
        resolved_last = title_name(last_name) or (managed_user.last_name if managed_user else "")
        if not resolved_first or not resolved_last:
            return _redirect_consult(
                dni=cleaned_dni,
                teacher_key=cleaned_teacher_key,
                error="Para crear usuario Moodle faltan nombre y apellido.",
            )
        moodle_user = moodle.create_user(
            upn,
            resolved_first,
            resolved_last,
            (contact_email or (managed_user.contact_email if managed_user else "")).strip().lower(),
        )
    if moodle_user.status not in {"ok", "created"} or not isinstance(moodle_user.data, dict):
        return _redirect_consult(
            dni=cleaned_dni,
            teacher_key=cleaned_teacher_key,
            error=f"No se pudo resolver usuario Moodle: {moodle_user.message}",
        )

    moodle_user_id = int(moodle_user.data.get("id", 0))
    if moodle_user_id <= 0:
        return _redirect_consult(
            dni=cleaned_dni,
            teacher_key=cleaned_teacher_key,
            error="Moodle no devolvio un id de usuario valido.",
        )

    courses_info = moodle.get_user_courses(moodle_user_id)
    enrolled_shortnames: set[str] = set()
    if courses_info.status == "ok" and isinstance(courses_info.data, list):
        for row in courses_info.data:
            code = str(row.get("shortname", "") or "").strip()
            if code:
                enrolled_shortnames.add(code)

    managed_user = db.query(ManagedUser).filter(ManagedUser.upn == upn).first()
    if not managed_user:
        normalized_first = title_name(first_name)
        normalized_last = title_name(last_name)
        if isinstance(moodle_user.data, dict):
            normalized_first = normalized_first or title_name(str(moodle_user.data.get("firstname", "") or ""))
            normalized_last = normalized_last or title_name(str(moodle_user.data.get("lastname", "") or ""))
        managed_user = upsert_managed_user(
            db,
            ImportedRow(
                document=cleaned_dni if person_type == "student" else "",
                first_name=normalized_first,
                last_name=normalized_last,
                full_name=f"{normalized_first} {normalized_last}".strip(),
                email=(contact_email or "").strip().lower(),
                upn=upn,
            ),
            person_type=person_type,
            azure_user_id=str(existing.data.get("id", "") or "") if isinstance(existing.data, dict) else "",
            moodle_user_id=str(moodle_user_id),
        )
    else:
        managed_user.moodle_user_id = str(moodle_user_id)
        db.commit()

    enrolled_now = 0
    skipped = 0
    groups_added = 0
    groups_skipped = 0
    groups_created = 0
    errors: list[str] = []
    for code in course_codes:
        already_enrolled = code in enrolled_shortnames
        if already_enrolled and not group_name:
            skipped += 1
            continue
        course = moodle.get_course_by_shortname(code)
        if course.status != "ok" or not isinstance(course.data, dict):
            errors.append(f"{code}: {course.message}")
            continue
        course_id = int(course.data.get("id", 0))
        if course_id <= 0:
            errors.append(f"{code}: curso sin id valido")
            continue
        if not already_enrolled:
            enrollment = moodle.enrol_user(moodle_user_id, course_id, person_type)
            if not _friendly_moodle_ok(enrollment.status, enrollment.message):
                errors.append(f"{code}: {enrollment.message}")
                continue
            enrolled_now += 1
            enrolled_shortnames.add(code)
        else:
            skipped += 1
        upsert_managed_enrollment(
            db,
            managed_user,
            course_shortname=code,
            role=person_type,
            course_name=str(course.data.get("fullname") or course.data.get("shortname") or ""),
            moodle_course_id=str(course_id),
        )
        if group_name:
            grouped = _ensure_consult_group_membership(
                moodle,
                course_id,
                group_name,
                moodle_user_id,
            )
            if grouped.status == "error":
                errors.append(f"{code}: {grouped.message}")
            elif grouped.status == "exists":
                groups_skipped += 1
            else:
                groups_added += 1
                if isinstance(grouped.data, dict) and grouped.data.get("group_created"):
                    groups_created += 1

    summary = f"Matriculados: {enrolled_now}, matriculas existentes: {skipped}."
    if group_name:
        summary += (
            f" Grupo {group_name}: agregados {groups_added}, ya asignados {groups_skipped}, "
            f"creados en cursos {groups_created}."
        )

    status = "ok" if not errors else "error"
    db.add(
        OperationLog(
            actor=username,
            action="consult_enroll",
            status=status,
            detail=(
                f"{upn}: matriculados={enrolled_now}, existentes={skipped}, "
                f"grupo={group_name or '-'}, grupos_agregados={groups_added}, "
                f"grupos_existentes={groups_skipped}, grupos_creados={groups_created}, "
                f"errores={len(errors)}"
            ),
        )
    )
    db.commit()
    if errors:
        return _redirect_consult(
            dni=cleaned_dni,
            teacher_key=cleaned_teacher_key,
            error=f"Operacion parcial. {summary} Errores: {' | '.join(errors)}",
        )
    return _redirect_consult(
        dni=cleaned_dni,
        teacher_key=cleaned_teacher_key,
        message=f"Matriculacion completada. {summary}",
    )


def _send_notification_job(
    *,
    job_id: str,
    actor: str,
    recipients: list[dict[str, str]],
    subject: str,
    message_template: str,
    batch_size: int,
    pause_seconds: int,
) -> None:
    mailer = Mailer(settings)
    sent = 0
    failed = 0
    errors: list[str] = []
    total = len(recipients)
    started_at = datetime.utcnow().isoformat()
    NOTIFICATION_JOBS[job_id] = {
        "status": "running",
        "total": total,
        "processed": 0,
        "sent": 0,
        "failed": 0,
        "started_at": started_at,
        "finished_at": "",
        "errors": [],
    }

    for idx, recipient in enumerate(recipients, start=1):
        personalized = message_template
        for key in ("full_name", "first_name", "last_name", "dni", "email", "modality"):
            personalized = personalized.replace("{" + key + "}", str(recipient.get(key, "")))

        result = mailer.send_message(recipient.get("email", ""), subject, personalized)
        if result.status == "sent":
            sent += 1
        else:
            failed += 1
            if len(errors) < 20:
                errors.append(f"{recipient.get('email', '')}: {result.message}")

        NOTIFICATION_JOBS[job_id].update(
            {
                "processed": idx,
                "sent": sent,
                "failed": failed,
                "errors": errors,
            }
        )

        if batch_size > 0 and idx < total and idx % batch_size == 0 and pause_seconds > 0:
            time.sleep(pause_seconds)

    finished_at = datetime.utcnow().isoformat()
    status = "completed" if failed == 0 else "completed_with_errors"
    NOTIFICATION_JOBS[job_id].update(
        {
            "status": status,
            "processed": total,
            "sent": sent,
            "failed": failed,
            "finished_at": finished_at,
            "errors": errors,
        }
    )

    db = SessionLocal()
    try:
        db.add(
            OperationLog(
                actor=actor,
                action="notifications_send",
                status="ok" if failed == 0 else "error",
                detail=f"Notificaciones: total={total}, enviados={sent}, errores={failed}",
            )
        )
        db.commit()
    finally:
        db.close()


@app.get("/notifications", response_class=HTMLResponse)
def notifications_page(request: Request):
    username = require_user(request)
    token = request.query_params.get("token", "").strip()
    job_id = request.query_params.get("job_id", "").strip()
    message = request.query_params.get("message", "")
    error = request.query_params.get("error", "")

    preview = None
    if token:
        path = _notification_file_path(token)
        if path.exists():
            data = json.loads(path.read_text(encoding="utf-8"))
            preview = data

    job = NOTIFICATION_JOBS.get(job_id) if job_id else None
    return templates.TemplateResponse(
        "notifications.html",
        {
            "request": request,
            "username": username,
            "preview": preview,
            "token": token,
            "job_id": job_id,
            "job": job,
            "message": message,
            "error": error,
            "settings": settings,
            "default_notification_subject": NOTIFICATION_DEFAULT_SUBJECT,
            "default_notification_message": NOTIFICATION_DEFAULT_MESSAGE,
        },
    )


@app.post("/notifications/preview")
async def notifications_preview(
    request: Request,
    file: UploadFile = File(...),
):
    require_user(request)
    ext = Path(file.filename or "").suffix.lower()
    if ext not in {".txt", ".xlsx"}:
        return RedirectResponse("/notifications?error=Solo+se+aceptan+archivos+.txt+o+.xlsx", status_code=303)

    content = await file.read()
    PENDING_NOTIFICATIONS_DIR.mkdir(parents=True, exist_ok=True)

    if ext == ".txt":
        try:
            recipients = parse_notifications_txt(content)
        except Exception as exc:
            return RedirectResponse(f"/notifications?error={quote_plus(str(exc))}", status_code=303)
        if not recipients:
            return RedirectResponse("/notifications?error=No+se+encontraron+emails+validos", status_code=303)
        token = _store_notification_preview(file.filename, recipients)
        return RedirectResponse(
            f"/notifications?token={quote_plus(token)}&message=Archivo+procesado",
            status_code=303,
        )

    try:
        columns, sample_rows = preview_notifications_excel(content)
    except Exception as exc:
        return RedirectResponse(f"/notifications?error={quote_plus(f'No se pudo leer el Excel: {exc}')}", status_code=303)

    mapping_token = secrets.token_urlsafe(16)
    pending_path = PENDING_NOTIFICATIONS_DIR / f"{mapping_token}.xlsx"
    pending_path.write_bytes(content)
    pending = request.session.get("pending_notifications", {})
    pending[mapping_token] = {"path": str(pending_path), "filename": file.filename}
    request.session["pending_notifications"] = pending

    return templates.TemplateResponse(
        "notifications_map_columns.html",
        {
            "request": request,
            "username": current_username(request),
            "token": mapping_token,
            "filename": file.filename,
            "columns": columns,
            "sample_rows": sample_rows,
            "mapping": _suggest_notification_mapping(columns),
            "error": "",
        },
    )


@app.post("/notifications/confirm-mapping")
def notifications_confirm_mapping(
    request: Request,
    token: str = Form(...),
    email_column: str = Form(""),
    modality_column: str = Form(""),
    dni_column: str = Form(""),
    last_name_column: str = Form(""),
    first_name_column: str = Form(""),
):
    require_user(request)
    pending = request.session.get("pending_notifications", {})
    item = pending.get(token)
    if not item:
        return RedirectResponse("/notifications?error=La+previsualizacion+vencio", status_code=303)

    path = Path(item["path"])
    if not path.exists():
        return RedirectResponse("/notifications?error=No+se+encontro+el+archivo+temporal", status_code=303)

    content = path.read_bytes()
    mapping = {
        "email": email_column,
        "modality": modality_column,
        "dni": dni_column,
        "last_name": last_name_column,
        "first_name": first_name_column,
    }

    try:
        recipients = parse_notifications_excel(content, mapping)
    except Exception as exc:
        columns, sample_rows = preview_notifications_excel(content)
        return templates.TemplateResponse(
            "notifications_map_columns.html",
            {
                "request": request,
                "username": current_username(request),
                "token": token,
                "filename": item["filename"],
                "columns": columns,
                "sample_rows": sample_rows,
                "mapping": mapping,
                "error": str(exc),
            },
            status_code=400,
        )

    if not recipients:
        return RedirectResponse("/notifications?error=No+se+encontraron+emails+validos", status_code=303)

    preview_token = _store_notification_preview(item["filename"], recipients)
    pending.pop(token, None)
    request.session["pending_notifications"] = pending
    path.unlink(missing_ok=True)
    return RedirectResponse(
        f"/notifications?token={quote_plus(preview_token)}&message=Archivo+procesado",
        status_code=303,
    )


@app.post("/notifications/send")
def notifications_send(
    request: Request,
    background_tasks: BackgroundTasks,
    token: str = Form(""),
    target_mode: str = Form("all"),
    subject: str = Form(""),
    message_body: str = Form(""),
    batch_size: int = Form(25),
    pause_seconds: int = Form(30),
):
    actor = require_user(request)
    if not token:
        return RedirectResponse("/notifications?error=Falta+token+de+previsualizacion", status_code=303)
    path = _notification_file_path(token)
    if not path.exists():
        return RedirectResponse("/notifications?error=Previsualizacion+expirada", status_code=303)

    data = json.loads(path.read_text(encoding="utf-8"))
    recipients = data.get("recipients", [])
    if not isinstance(recipients, list) or not recipients:
        return RedirectResponse("/notifications?error=No+hay+destinatarios+para+enviar", status_code=303)

    mode = (target_mode or "all").strip().lower()
    if mode not in {"all", "distance", "presential", "other"}:
        mode = "all"
    filtered = [row for row in recipients if mode == "all" or row.get("modality_group") == mode]
    if not filtered:
        return RedirectResponse(
            f"/notifications?token={quote_plus(token)}&error=No+hay+destinatarios+para+ese+filtro",
            status_code=303,
        )

    normalized_subject = (subject or "").strip() or NOTIFICATION_DEFAULT_SUBJECT
    normalized_body = (message_body or "").strip() or NOTIFICATION_DEFAULT_MESSAGE
    if not normalized_subject:
        return RedirectResponse(
            f"/notifications?token={quote_plus(token)}&error=Completar+asunto",
            status_code=303,
        )
    if not normalized_body:
        return RedirectResponse(
            f"/notifications?token={quote_plus(token)}&error=Completar+mensaje",
            status_code=303,
        )

    batch_size = max(1, min(int(batch_size), 200))
    pause_seconds = max(0, min(int(pause_seconds), 180))

    job_id = uuid.uuid4().hex[:12]
    NOTIFICATION_JOBS[job_id] = {
        "status": "queued",
        "total": len(filtered),
        "processed": 0,
        "sent": 0,
        "failed": 0,
        "started_at": "",
        "finished_at": "",
        "errors": [],
    }
    background_tasks.add_task(
        _send_notification_job,
        job_id=job_id,
        actor=actor,
        recipients=filtered,
        subject=normalized_subject,
        message_template=normalized_body,
        batch_size=batch_size,
        pause_seconds=pause_seconds,
    )
    return RedirectResponse(
        f"/notifications?token={quote_plus(token)}&job_id={quote_plus(job_id)}&message=Envio+iniciado",
        status_code=303,
    )


@app.get("/notifications/jobs/{job_id}")
def notifications_job_status(job_id: str, request: Request):
    require_user(request)
    job = NOTIFICATION_JOBS.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Job no encontrado")
    return JSONResponse(job)


@app.get("/converter", response_class=HTMLResponse)
def converter_page(request: Request, q: str = "", db: Session = Depends(get_db)):
    username = require_user(request)
    search = str(q or "").strip()
    course_names: dict[str, str] = {}
    known_courses = (
        db.query(MatriculationRunRow.course_shortname, MatriculationRunRow.source_course)
        .filter(MatriculationRunRow.source_course != "")
        .order_by(MatriculationRunRow.id.desc())
        .distinct()
        .all()
    )
    for shortname, course_name in known_courses:
        base_code = str(shortname or "").split("-", 1)[0]
        if base_code and course_name and base_code not in course_names:
            course_names[base_code] = str(course_name).strip()
    rules = (
        db.query(CourseConversionRule)
        .order_by(CourseConversionRule.sort_order.asc(), CourseConversionRule.id.asc())
        .all()
    )
    if search:
        needle = search.casefold()
        rules = [
            rule
            for rule in rules
            if needle
            in " ".join(
                [
                    rule.actual_code,
                    rule.old_code,
                    rule.category_rule,
                    rule.registration_rule,
                    course_names.get(rule.actual_code, ""),
                ]
            ).casefold()
        ]
    return templates.TemplateResponse(
        "converter.html",
        {
            "request": request,
            "username": username,
            "rules": rules,
            "rule_count": db.query(CourseConversionRule).count(),
            "course_names": course_names,
            "q": search,
            "message": request.query_params.get("message", ""),
            "error": request.query_params.get("error", ""),
        },
    )


@app.post("/converter/add")
def converter_add(
    request: Request,
    actual_code: str = Form(...),
    old_code: str = Form(""),
    category_rule: str = Form(""),
    registration_rule: str = Form(""),
    db: Session = Depends(get_db),
):
    actor = require_user(request)
    try:
        actual, old, category, registration = _normalize_conversion_rule_fields(
            actual_code,
            old_code,
            category_rule,
            registration_rule,
        )
    except ValueError as exc:
        return RedirectResponse(f"/converter?error={quote_plus(str(exc))}", status_code=303)
    max_order = db.query(CourseConversionRule.sort_order).order_by(
        CourseConversionRule.sort_order.desc()
    ).first()
    sort_order = (max_order[0] if max_order else 0) + 10
    rule = CourseConversionRule(
        actual_code=actual,
        old_code=old,
        category_rule=category,
        registration_rule=registration,
        sort_order=sort_order,
    )
    db.add(rule)
    db.flush()
    db.add(
        OperationLog(
            actor=actor,
            action="course_conversion_add",
            status="ok",
            detail=f"Regla {rule.id}: Lincu {old or actual} -> Moodle {actual}.",
        )
    )
    db.commit()
    return RedirectResponse(
        f"/converter?q={quote_plus(old or actual)}&message=Regla+agregada",
        status_code=303,
    )


@app.post("/converter/{rule_id}/update")
def converter_update(
    rule_id: int,
    request: Request,
    actual_code: str = Form(...),
    old_code: str = Form(""),
    category_rule: str = Form(""),
    registration_rule: str = Form(""),
    q: str = Form(""),
    db: Session = Depends(get_db),
):
    actor = require_user(request)
    rule = db.query(CourseConversionRule).filter(CourseConversionRule.id == rule_id).first()
    if not rule:
        return RedirectResponse("/converter?error=Regla+no+encontrada", status_code=303)
    try:
        actual, old, category, registration = _normalize_conversion_rule_fields(
            actual_code,
            old_code,
            category_rule,
            registration_rule,
        )
    except ValueError as exc:
        return RedirectResponse(
            f"/converter?q={quote_plus(q)}&error={quote_plus(str(exc))}",
            status_code=303,
        )
    rule.actual_code = actual
    rule.old_code = old
    rule.category_rule = category
    rule.registration_rule = registration
    db.add(
        OperationLog(
            actor=actor,
            action="course_conversion_update",
            status="ok",
            detail=f"Regla {rule.id}: Lincu {old or actual} -> Moodle {actual}.",
        )
    )
    db.commit()
    return RedirectResponse(
        f"/converter?q={quote_plus(q)}&message=Regla+guardada",
        status_code=303,
    )


@app.post("/converter/{rule_id}/delete")
def converter_delete(
    rule_id: int,
    request: Request,
    q: str = Form(""),
    db: Session = Depends(get_db),
):
    actor = require_user(request)
    rule = db.query(CourseConversionRule).filter(CourseConversionRule.id == rule_id).first()
    if not rule:
        return RedirectResponse("/converter?error=Regla+no+encontrada", status_code=303)
    detail = f"Regla {rule.id}: Lincu {rule.old_code or rule.actual_code} -> Moodle {rule.actual_code}."
    db.delete(rule)
    db.add(
        OperationLog(
            actor=actor,
            action="course_conversion_delete",
            status="ok",
            detail=detail,
        )
    )
    db.commit()
    return RedirectResponse(
        f"/converter?q={quote_plus(q)}&message=Regla+eliminada",
        status_code=303,
    )


@app.get("/electives", response_class=HTMLResponse)
def electives_page(
    request: Request,
    run_id: int | None = None,
    db: Session = Depends(get_db),
):
    username = require_user(request)
    runs = (
        db.query(MatriculationRun)
        .filter(MatriculationRun.source_type == "elective")
        .order_by(MatriculationRun.created_at.desc())
        .limit(20)
        .all()
    )
    selected_run = None
    if run_id:
        selected_run = (
            db.query(MatriculationRun)
            .filter(
                MatriculationRun.id == run_id,
                MatriculationRun.source_type == "elective",
            )
            .first()
        )
    elif runs:
        selected_run = runs[0]
    actionable_count = 0
    if selected_run:
        actionable_count = (
            db.query(MatriculationRunRow)
            .filter(
                MatriculationRunRow.run_id == selected_run.id,
                MatriculationRunRow.action.in_(MATRICULATION_ACTIONS),
                MatriculationRunRow.status.in_({"preview", "error"}),
            )
            .count()
        )
    return templates.TemplateResponse(
        "electives.html",
        {
            "request": request,
            "username": username,
            "runs": runs,
            "selected_run": selected_run,
            "actionable_count": actionable_count,
            "message": request.query_params.get("message", ""),
            "error": request.query_params.get("error", ""),
        },
    )


@app.post("/electives/upload")
async def electives_upload(
    request: Request,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    username = require_user(request)
    filename = file.filename or "electiva.xlsx"
    if Path(filename).suffix.lower() not in {".xlsx", ".xls"}:
        return RedirectResponse("/electives?error=Solo+se+aceptan+archivos+Excel", status_code=303)
    content = await file.read()
    if not content:
        return RedirectResponse("/electives?error=El+archivo+esta+vacio", status_code=303)
    try:
        columns, sample_rows = preview_elective_excel(content)
    except Exception as exc:
        return RedirectResponse(
            f"/electives?error={quote_plus(f'No se pudo leer el Excel: {exc}')}",
            status_code=303,
        )
    if not columns:
        return RedirectResponse("/electives?error=No+se+encontraron+columnas", status_code=303)

    PENDING_ELECTIVES_DIR.mkdir(parents=True, exist_ok=True)
    token = secrets.token_urlsafe(24)
    pending_path = PENDING_ELECTIVES_DIR / f"{token}.xlsx"
    pending_path.write_bytes(content)
    pending = request.session.get("pending_electives", {})
    pending[token] = {"path": str(pending_path), "filename": filename}
    request.session["pending_electives"] = pending
    mapping = suggest_elective_mapping(columns)
    subject_summaries = analyze_elective_subjects(content, mapping.get("subject", ""))
    return templates.TemplateResponse(
        "elective_mapping.html",
        {
            "request": request,
            "username": username,
            "token": token,
            "filename": filename,
            "columns": columns,
            "sample_rows": sample_rows,
            "mapping": mapping,
            "subject_summaries": subject_summaries,
            "default_group": "",
            "error": "",
        },
    )


@app.post("/electives/preview")
async def electives_preview(
    request: Request,
    token: str = Form(...),
    document_column: str = Form(""),
    first_name_column: str = Form(""),
    last_name_column: str = Form(""),
    full_name_column: str = Form(""),
    email_column: str = Form(""),
    modality_column: str = Form(""),
    course_column: str = Form(""),
    subject_column: str = Form(""),
    default_group: str = Form(""),
    subject_keys: list[str] = Form([]),
    subject_courses: list[str] = Form([]),
    db: Session = Depends(get_db),
):
    username = require_user(request)
    pending = request.session.get("pending_electives", {})
    item = pending.get(token)
    if not item:
        return RedirectResponse(
            "/electives?error=La+previsualizacion+vencio.+Subi+el+Excel+nuevamente",
            status_code=303,
        )
    pending_path = Path(item["path"])
    if not pending_path.exists():
        return RedirectResponse(
            "/electives?error=No+se+encontro+el+archivo+temporal.+Subi+el+Excel+nuevamente",
            status_code=303,
        )
    content = pending_path.read_bytes()
    mapping = {
        "document": document_column,
        "first_name": first_name_column,
        "last_name": last_name_column,
        "full_name": full_name_column,
        "email": email_column,
        "modality": modality_column,
        "course": course_column,
        "subject": subject_column,
    }
    subject_course_map = {
        key: course
        for key, course in zip(subject_keys, subject_courses)
        if key.strip() and course.strip()
    }
    try:
        run = _create_elective_preview(
            content,
            item["filename"],
            username,
            mapping,
            subject_course_map,
            default_group,
            db,
        )
    except ValueError as exc:
        columns, sample_rows = preview_elective_excel(content)
        subject_summaries = analyze_elective_subjects(content, subject_column)
        for summary in subject_summaries:
            if summary["key"] in subject_course_map:
                summary["suggested_course"] = subject_course_map[summary["key"]]
        return templates.TemplateResponse(
            "elective_mapping.html",
            {
                "request": request,
                "username": username,
                "token": token,
                "filename": item["filename"],
                "columns": columns,
                "sample_rows": sample_rows,
                "mapping": mapping,
                "subject_summaries": subject_summaries,
                "default_group": default_group,
                "error": str(exc),
            },
            status_code=400,
        )
    except Exception as exc:
        return RedirectResponse(
            f"/electives?error={quote_plus(f'No se pudo generar la previa: {exc}')}",
            status_code=303,
        )
    finally:
        # The pending file is retained only while the user is associating columns.
        if "run" in locals():
            pending_path.unlink(missing_ok=True)
            pending.pop(token, None)
            request.session["pending_electives"] = pending
    return RedirectResponse(
        f"/electives?run_id={run.id}&message=Vista+previa+Electiva+generada",
        status_code=303,
    )


@app.post("/electives/runs/{run_id}/execute")
def electives_execute(
    run_id: int,
    request: Request,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
):
    actor = require_user(request)
    run = (
        db.query(MatriculationRun)
        .filter(MatriculationRun.id == run_id, MatriculationRun.source_type == "elective")
        .first()
    )
    if not run:
        return RedirectResponse("/electives?error=Vista+previa+Electiva+no+encontrada", status_code=303)
    try:
        row_ids = _queue_matriculation_execution(run, db)
    except ValueError as exc:
        return RedirectResponse(
            f"/electives?run_id={run.id}&error={quote_plus(str(exc))}",
            status_code=303,
        )
    background_tasks.add_task(_run_matriculation_job, run.id, actor, row_ids)
    return RedirectResponse(
        f"/electives?run_id={run.id}&message=Ejecucion+iniciada+en+Moodle",
        status_code=303,
    )


@app.get("/matriculation", response_class=HTMLResponse)
def matriculation_page(
    request: Request,
    run_id: int | None = None,
    show_all: int = 0,
    view: str = "",
    page: int = 1,
    per_page: int = 200,
    db: Session = Depends(get_db),
):
    username = require_user(request)
    runs = db.query(MatriculationRun).order_by(MatriculationRun.created_at.desc()).limit(20).all()
    selected_run = None
    if run_id:
        selected_run = db.query(MatriculationRun).filter(MatriculationRun.id == run_id).first()
    elif runs:
        selected_run = runs[0]

    rows = []
    total_rows = 0
    filtered_count = 0
    current_page = max(1, int(page or 1))
    page_size = max(50, min(int(per_page or 200), 500))
    total_pages = 1
    pagination_pages: list[int | None] = []
    pagination_urls: dict[int, str] = {}
    actionable_count = 0
    filtered_actionable_count = 0
    pending_count = 0
    unique_student_count = 0
    execution_seconds = 0.0
    valid_views = {
        "changes",
        "missing",
        "group_changes",
        "course_changes",
        "unenrollments",
        "all",
        "unchanged",
        "errors",
    }
    detail_view = view.strip().lower() if view else ("all" if show_all else "changes")
    if detail_view not in valid_views:
        detail_view = "changes"
    view_labels = {
        "changes": "Cambios",
        "missing": "Faltantes en Moodle",
        "group_changes": "Cambios de grupo",
        "course_changes": "Cambios de curso",
        "unenrollments": "Bajas",
        "all": "Todos",
        "unchanged": "Sin cambios",
        "errors": "Errores",
    }
    if selected_run:
        query = db.query(MatriculationRunRow).filter(MatriculationRunRow.run_id == selected_run.id)
        total_rows = query.count()
        unique_student_count = (
            query.filter(
                MatriculationRunRow.document != "",
                MatriculationRunRow.upn != "",
            )
            .with_entities(MatriculationRunRow.document)
            .distinct()
            .count()
        )
        actionable_count = query.filter(
            MatriculationRunRow.action.in_(MATRICULATION_ACTIONS),
            MatriculationRunRow.status.in_({"preview", "error"}),
        ).count()
        pending_count = query.filter(
            MatriculationRunRow.action.in_(MATRICULATION_ACTIONS),
            MatriculationRunRow.status == "preview",
        ).count()
        if selected_run.started_at:
            execution_end = selected_run.finished_at or datetime.utcnow()
            execution_seconds = max(0.0, (execution_end - selected_run.started_at).total_seconds())
        if detail_view == "changes":
            query = query.filter(or_(MatriculationRunRow.action != "unchanged", MatriculationRunRow.status == "unverified"))
        elif detail_view == "missing":
            query = query.filter(MatriculationRunRow.verification_status == "missing")
        elif detail_view == "group_changes":
            query = query.filter(MatriculationRunRow.action == "group_change")
        elif detail_view == "course_changes":
            query = query.filter(MatriculationRunRow.action == "course_change")
        elif detail_view == "unenrollments":
            query = query.filter(MatriculationRunRow.action == "unenroll")
        elif detail_view == "unchanged":
            query = query.filter(MatriculationRunRow.action == "unchanged")
        elif detail_view == "errors":
            query = query.filter(
                or_(
                    MatriculationRunRow.action == "error",
                    MatriculationRunRow.status.in_({"error", "unverified"}),
                )
            )
        filtered_actionable_count = query.filter(
            MatriculationRunRow.action.in_(MATRICULATION_ACTIONS),
            MatriculationRunRow.status.in_({"preview", "error"}),
        ).count()
        filtered_count = query.count()
        total_pages = max(1, (filtered_count + page_size - 1) // page_size)
        current_page = min(current_page, total_pages)
        rows = (
            query.order_by(MatriculationRunRow.action.asc(), MatriculationRunRow.id.asc())
            .offset((current_page - 1) * page_size)
            .limit(page_size)
            .all()
        )
        pagination_pages = _pagination_window(current_page, total_pages)
        pagination_urls = {
            page_number: "/matriculation?"
            + urlencode(
                {
                    "run_id": selected_run.id,
                    "view": detail_view,
                    "page": page_number,
                    "per_page": page_size,
                }
            )
            for page_number in pagination_pages
            if page_number is not None
        }

    return templates.TemplateResponse(
        "matriculation.html",
        {
            "request": request,
            "username": username,
            "runs": runs,
            "selected_run": selected_run,
            "rows": rows,
            "total_rows": total_rows,
            "filtered_count": filtered_count,
            "current_page": current_page,
            "per_page": page_size,
            "total_pages": total_pages,
            "pagination_pages": pagination_pages,
            "pagination_urls": pagination_urls,
            "show_all": detail_view == "all",
            "detail_view": detail_view,
            "detail_view_label": view_labels[detail_view],
            "actionable_count": actionable_count,
            "filtered_actionable_count": filtered_actionable_count,
            "pending_count": pending_count,
            "unique_student_count": unique_student_count,
            "missing_moodle_count": db.query(MatriculationRunRow).filter_by(run_id=selected_run.id, verification_status="missing").count() if selected_run else 0,
            "execution_seconds": execution_seconds,
            "message": request.query_params.get("message", ""),
            "error": request.query_params.get("error", ""),
        },
    )


@app.post("/matriculation/preview")
async def matriculation_preview(
    request: Request,
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    verify_moodle: str = Form(""),
    verification_courses: str = Form(""),
    db: Session = Depends(get_db),
):
    actor = require_user(request)
    ext = Path(file.filename or "").suffix.lower()
    if ext != ".txt":
        return RedirectResponse("/matriculation?error=Solo+se+acepta+archivo+Lincu+.txt", status_code=303)
    content = await file.read()
    if not content:
        return RedirectResponse("/matriculation?error=El+archivo+esta+vacio", status_code=303)
    try:
        verification_scope = _parse_matriculation_verification_scope(verification_courses) if verify_moodle == "1" else ()
        run = _create_matriculation_preview(
            content,
            file.filename or "lincu.txt",
            actor,
            db,
            verification_scope=verification_scope,
        )
    except Exception as exc:
        return RedirectResponse(f"/matriculation?error={quote_plus(str(exc))}", status_code=303)
    if run.status == "verifying":
        background_tasks.add_task(_run_matriculation_verification, run.id, actor)
    return RedirectResponse(
        f"/matriculation?run_id={run.id}&message=Vista+previa+generada",
        status_code=303,
    )


@app.post("/matriculation/runs/{run_id}/execute")
def matriculation_execute(
    run_id: int,
    request: Request,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
):
    actor = require_user(request)
    run = db.query(MatriculationRun).filter(MatriculationRun.id == run_id).first()
    if not run:
        return RedirectResponse("/matriculation?error=Corrida+no+encontrada", status_code=303)
    try:
        row_ids = _queue_matriculation_execution(run, db)
    except ValueError as exc:
        return RedirectResponse(
            f"/matriculation?run_id={run.id}&error={quote_plus(str(exc))}",
            status_code=303,
        )
    background_tasks.add_task(_run_matriculation_job, run.id, actor, row_ids)
    return RedirectResponse(
        f"/matriculation?run_id={run.id}&message=Ejecucion+iniciada",
        status_code=303,
    )


@app.post("/matriculation/runs/{run_id}/execute-selected")
def matriculation_execute_selected(
    run_id: int,
    request: Request,
    background_tasks: BackgroundTasks,
    row_ids: list[int] = Form([]),
    page: int = Form(1),
    per_page: int = Form(200),
    show_all: int = Form(0),
    view: str = Form(""),
    db: Session = Depends(get_db),
):
    actor = require_user(request)
    run = db.query(MatriculationRun).filter(MatriculationRun.id == run_id).first()
    if not run:
        return RedirectResponse("/matriculation?error=Corrida+no+encontrada", status_code=303)
    valid_views = {
        "changes", "group_changes", "course_changes", "unenrollments", "all", "unchanged", "errors", "missing"
    }
    return_view = view if view in valid_views else ("all" if show_all else "changes")
    return_url = (
        f"/matriculation?run_id={run.id}&page={max(1, page)}"
        f"&per_page={max(50, min(per_page, 500))}&view={quote_plus(return_view)}"
    )
    try:
        selected_ids = _queue_matriculation_execution(run, db, row_ids)
    except ValueError as exc:
        return RedirectResponse(f"{return_url}&error={quote_plus(str(exc))}", status_code=303)
    background_tasks.add_task(_run_matriculation_job, run.id, actor, selected_ids)
    return RedirectResponse(f"{return_url}&message=Lote+seleccionado+iniciado", status_code=303)


@app.post("/matriculation/runs/{run_id}/rows/{row_id}/execute")
def matriculation_execute_row(
    run_id: int,
    row_id: int,
    request: Request,
    background_tasks: BackgroundTasks,
    page: int = Form(1),
    per_page: int = Form(200),
    show_all: int = Form(0),
    view: str = Form(""),
    db: Session = Depends(get_db),
):
    actor = require_user(request)
    run = db.query(MatriculationRun).filter(MatriculationRun.id == run_id).first()
    if not run:
        return RedirectResponse("/matriculation?error=Corrida+no+encontrada", status_code=303)
    valid_views = {
        "changes", "group_changes", "course_changes", "unenrollments", "all", "unchanged", "errors", "missing"
    }
    return_view = view if view in valid_views else ("all" if show_all else "changes")
    return_url = (
        f"/matriculation?run_id={run.id}&page={max(1, page)}"
        f"&per_page={max(50, min(per_page, 500))}&view={quote_plus(return_view)}"
    )
    try:
        selected_ids = _queue_matriculation_execution(run, db, [row_id])
    except ValueError as exc:
        return RedirectResponse(f"{return_url}&error={quote_plus(str(exc))}", status_code=303)
    background_tasks.add_task(_run_matriculation_job, run.id, actor, selected_ids)
    return RedirectResponse(f"{return_url}&message=Registro+enviado+a+Moodle", status_code=303)


@app.post("/matriculation/runs/{run_id}/delete")
def matriculation_delete_run(run_id: int, request: Request, db: Session = Depends(get_db)):
    actor = require_user(request)
    run = db.query(MatriculationRun).filter(MatriculationRun.id == run_id).first()
    if not run:
        return RedirectResponse("/matriculation?error=Corrida+no+encontrada", status_code=303)
    if run.status in {"queued", "running", "verifying"}:
        return RedirectResponse(
            f"/matriculation?run_id={run.id}&error=No+se+puede+eliminar+una+corrida+en+ejecucion",
            status_code=303,
        )
    filename = run.filename
    db.query(MatriculationRunRow).filter(MatriculationRunRow.run_id == run.id).delete(
        synchronize_session=False
    )
    db.query(MatriculationCourseAlert).filter(MatriculationCourseAlert.run_id == run.id).delete(
        synchronize_session=False
    )
    db.delete(run)
    db.add(
        OperationLog(
            actor=actor,
            action="matriculation_history_delete",
            status="ok",
            detail=f"Corrida {run_id} eliminada de la bitacora ({filename}).",
        )
    )
    db.commit()
    return RedirectResponse("/matriculation?message=Historial+eliminado", status_code=303)


@app.get("/matriculation/runs/{run_id}/course-load.xlsx")
def matriculation_course_load_export(
    run_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    require_user(request)
    run = db.query(MatriculationRun).filter(MatriculationRun.id == run_id).first()
    if not run:
        raise HTTPException(status_code=404, detail="Corrida no encontrada")
    alerts = (
        db.query(MatriculationCourseAlert)
        .filter(MatriculationCourseAlert.run_id == run.id)
        .order_by(
            MatriculationCourseAlert.course_count.desc(),
            MatriculationCourseAlert.last_name.asc(),
            MatriculationCourseAlert.first_name.asc(),
        )
        .all()
    )
    if not alerts:
        return RedirectResponse(
            f"/matriculation?run_id={run.id}&error=No+hay+alumnos+que+alcancen+el+umbral",
            status_code=303,
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Carga academica"
    headers = [
        "DNI",
        "APELLIDO",
        "NOMBRE",
        "TELEFONO",
        "EMAIL PERSONAL",
        "MODALIDAD",
        "CANTIDAD MATERIAS",
        "MATERIAS",
    ]
    sheet.append(headers)
    for alert in alerts:
        sheet.append(
            [
                alert.document,
                alert.last_name,
                alert.first_name,
                alert.phone,
                alert.email,
                alert.modality,
                alert.course_count,
                alert.course_names,
            ]
        )

    header_fill = PatternFill("solid", fgColor="1B6CA8")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = {"A": 14, "B": 24, "C": 24, "D": 18, "E": 34, "F": 18, "G": 20, "H": 80}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        row[7].alignment = Alignment(wrap_text=True, vertical="top")

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"carga_academica_lincu_corrida_{run.id}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/matriculation/runs/{run_id}/progress")
def matriculation_progress(run_id: int, request: Request, db: Session = Depends(get_db)):
    require_user(request)
    run = db.query(MatriculationRun).filter(MatriculationRun.id == run_id).first()
    if not run:
        raise HTTPException(status_code=404, detail="Corrida no encontrada")
    verification = run.status == "verifying" or (
        run.status not in {"queued", "running"}
        and bool(run.verification_finished_at and not run.started_at)
    )
    total = (run.verification_total_courses if verification else run.execution_total_rows) or 0
    processed = (run.verification_processed_courses if verification else run.execution_processed_rows) or 0
    percent = int((processed / total) * 100) if total else 0
    if run.status not in {"queued", "running", "verifying"} and total:
        percent = 100
    elapsed_seconds = 0.0
    started = run.verification_started_at if verification else run.started_at
    finished = run.verification_finished_at if verification else run.finished_at
    if started:
        elapsed_seconds = max(0.0, ((finished or datetime.utcnow()) - started).total_seconds())
    rows_per_minute = (processed / elapsed_seconds) * 60 if elapsed_seconds > 0 else 0.0
    eta_seconds = (
        ((total - processed) / (processed / elapsed_seconds))
        if processed > 0 and elapsed_seconds > 0 and total > processed
        else 0.0
    )
    return JSONResponse(
        {
            "id": run.id,
            "status": run.status,
            "unit": "cursos" if verification else "filas",
            "processed": processed,
            "total": total,
            "percent": max(0, min(percent, 100)),
            "elapsed_seconds": round(elapsed_seconds, 1),
            "rows_per_minute": round(rows_per_minute, 1),
            "eta_seconds": round(eta_seconds, 1),
            "enroll_rows": run.enroll_rows,
            "group_change_rows": run.group_change_rows,
            "course_change_rows": run.course_change_rows,
            "unenroll_rows": run.unenroll_rows,
            "unchanged_rows": run.unchanged_rows,
            "error_rows": run.error_rows,
            "message": run.message,
        }
    )


@app.get("/lincu", response_class=HTMLResponse)
def lincu_page(request: Request):
    username = require_user(request)
    job_id = request.query_params.get("job_id", "").strip()
    job = LINCU_JOBS.get(job_id) if job_id else None
    message = request.query_params.get("message", "")
    error = request.query_params.get("error", "")
    return templates.TemplateResponse(
        "lincu.html",
        {
            "request": request,
            "username": username,
            "job_id": job_id,
            "job": job,
            "message": message,
            "error": error,
        },
    )


@app.post("/lincu/process")
async def lincu_process(
    request: Request,
    background_tasks: BackgroundTasks,
    lincu_file: UploadFile = File(...),
    ingresos_file: UploadFile = File(...),
):
    actor = require_user(request)
    lincu_ext = Path(lincu_file.filename or "").suffix.lower()
    ingresos_ext = Path(ingresos_file.filename or "").suffix.lower()
    if lincu_ext != ".txt":
        return RedirectResponse("/lincu?error=El+archivo+Lincu+debe+ser+.txt", status_code=303)
    if ingresos_ext != ".csv":
        return RedirectResponse("/lincu?error=El+archivo+de+ingresos+debe+ser+.csv", status_code=303)

    lincu_bytes = await lincu_file.read()
    ingresos_bytes = await ingresos_file.read()
    if not lincu_bytes:
        return RedirectResponse("/lincu?error=El+archivo+Lincu+esta+vacio", status_code=303)
    if not ingresos_bytes:
        return RedirectResponse("/lincu?error=El+archivo+CSV+esta+vacio", status_code=303)

    PENDING_LINCU_DIR.mkdir(parents=True, exist_ok=True)
    job_id = uuid.uuid4().hex[:12]
    _lincu_file_path(job_id, "lincu.txt").write_bytes(lincu_bytes)
    _lincu_file_path(job_id, "ingresos.csv").write_bytes(ingresos_bytes)
    LINCU_JOBS[job_id] = {
        "status": "queued",
        "stage": "En cola",
        "progress": 0,
        "rows": [],
        "summary": {},
        "error": "",
        "started_at": "",
        "finished_at": "",
    }
    background_tasks.add_task(_run_lincu_job, job_id=job_id, actor=actor)
    return RedirectResponse(f"/lincu?job_id={quote_plus(job_id)}&message=Proceso+iniciado", status_code=303)


@app.get("/lincu/jobs/{job_id}")
def lincu_job_status(job_id: str, request: Request):
    require_user(request)
    job = LINCU_JOBS.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Job no encontrado")
    return JSONResponse(
        {
            "status": job.get("status", "queued"),
            "stage": job.get("stage", ""),
            "progress": job.get("progress", 0),
            "error": job.get("error", ""),
            "summary": job.get("summary", {}),
        }
    )


@app.get("/lincu/jobs/{job_id}/export.csv")
def lincu_export_csv(job_id: str, request: Request):
    require_user(request)
    job = LINCU_JOBS.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Job no encontrado")
    if job.get("status") != "completed":
        return RedirectResponse(f"/lincu?job_id={quote_plus(job_id)}&error=El+proceso+aun+no+termino", status_code=303)

    rows = job.get("rows", [])
    if not isinstance(rows, list) or not rows:
        return RedirectResponse(f"/lincu?job_id={quote_plus(job_id)}&error=No+hay+resultados+para+exportar", status_code=303)

    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(
        [
            "DNI",
            "APELLIDO",
            "NOMBRE",
            "UPN",
            "INGRESOS",
            "CURSOS_CSV",
            "MATERIAS_LINCU",
            "MODALIDAD",
            "EMAIL_PERSONAL",
        ]
    )
    for row in rows:
        writer.writerow(
            [
                row.get("dni", ""),
                row.get("last_name", ""),
                row.get("first_name", ""),
                row.get("upn", ""),
                row.get("ingresos", ""),
                row.get("cursos_csv", 0),
                row.get("materias_lincu", 0),
                row.get("modality", "-"),
                row.get("email_personal", ""),
            ]
        )

    content = buffer.getvalue().encode("utf-8-sig")
    filename = f"lincu_resultado_{job_id}.csv"
    return StreamingResponse(
        iter([content]),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/lincu/jobs/{job_id}/export.xlsx")
def lincu_export_xlsx(job_id: str, request: Request):
    require_user(request)
    job = LINCU_JOBS.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Job no encontrado")
    if job.get("status") != "completed":
        return RedirectResponse(f"/lincu?job_id={quote_plus(job_id)}&error=El+proceso+aun+no+termino", status_code=303)

    rows = job.get("rows", [])
    if not isinstance(rows, list) or not rows:
        return RedirectResponse(f"/lincu?job_id={quote_plus(job_id)}&error=No+hay+resultados+para+exportar", status_code=303)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Lincu"
    sheet.append(
        [
            "DNI",
            "APELLIDO",
            "NOMBRE",
            "UPN",
            "INGRESOS",
            "CURSOS_CSV",
            "MATERIAS_LINCU",
            "MODALIDAD",
            "EMAIL_PERSONAL",
        ]
    )

    for row in rows:
        sheet.append(
            [
                row.get("dni", ""),
                row.get("last_name", ""),
                row.get("first_name", ""),
                row.get("upn", ""),
                row.get("ingresos", ""),
                row.get("cursos_csv", 0),
                row.get("materias_lincu", 0),
                row.get("modality", "-"),
                row.get("email_personal", ""),
            ]
        )

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"lincu_resultado_{job_id}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/users/{person_type}", response_class=HTMLResponse)
def managed_users(
    person_type: str,
    request: Request,
    dni: str = "",
    name_query: str = "",
    course_shortname: str = "",
    course_query: str = "",
    page: int = 1,
    per_page: int = MANAGED_USERS_PAGE_SIZE,
    db: Session = Depends(get_db),
):
    username = require_user(request)
    if person_type not in {"student", "teacher"}:
        raise HTTPException(status_code=404, detail="Tipo invalido")

    query = db.query(ManagedUser).filter(ManagedUser.person_type == person_type, ManagedUser.active.is_(True))
    if dni.strip():
        query = query.filter(ManagedUser.document.contains(dni.strip()))
    if name_query.strip():
        q = name_query.strip()
        query = query.filter(
            or_(
                ManagedUser.first_name.contains(q),
                ManagedUser.last_name.contains(q),
                ManagedUser.full_name.contains(q),
                ManagedUser.upn.contains(q),
            )
        )
    if course_shortname.strip():
        query = query.join(ManagedEnrollment).filter(
            ManagedEnrollment.course_shortname == course_shortname.strip(),
            ManagedEnrollment.active.is_(True),
        )
    elif course_query.strip():
        q = course_query.strip()
        query = query.join(ManagedEnrollment).filter(
            ManagedEnrollment.active.is_(True),
            or_(
                ManagedEnrollment.course_shortname.contains(q),
                ManagedEnrollment.course_name.contains(q),
            ),
        )

    current_page = max(page, 1)
    page_size = per_page if per_page in MANAGED_USERS_PAGE_SIZES else MANAGED_USERS_PAGE_SIZE
    user_ids_query = query.with_entities(ManagedUser.id).distinct()
    total_users = user_ids_query.count()
    total_pages = max(1, (total_users + page_size - 1) // page_size) if total_users else 1
    if current_page > total_pages:
        current_page = total_pages

    offset = (current_page - 1) * page_size
    paged_user_ids = [
        row[0]
        for row in (
            user_ids_query
            .order_by(ManagedUser.full_name.asc(), ManagedUser.id.asc())
            .offset(offset)
            .limit(page_size)
            .all()
        )
    ]

    users = []
    if paged_user_ids:
        users = (
            db.query(ManagedUser)
            .filter(ManagedUser.id.in_(paged_user_ids))
            .order_by(ManagedUser.full_name.asc(), ManagedUser.id.asc())
            .all()
        )
    user_ids = [user.id for user in users]

    enrollments = []
    if user_ids:
        enrollments_query = db.query(ManagedEnrollment).filter(
            ManagedEnrollment.user_id.in_(user_ids),
            ManagedEnrollment.active.is_(True),
        )
        if course_shortname.strip():
            enrollments_query = enrollments_query.filter(
                ManagedEnrollment.course_shortname == course_shortname.strip()
            )
        elif course_query.strip():
            q = course_query.strip()
            enrollments_query = enrollments_query.filter(
                or_(
                    ManagedEnrollment.course_shortname.contains(q),
                    ManagedEnrollment.course_name.contains(q),
                )
            )
        enrollments = enrollments_query.order_by(ManagedEnrollment.course_shortname.asc()).all()
    enrollments_by_user: dict[int, list[ManagedEnrollment]] = {}
    for enrollment in enrollments:
        enrollments_by_user.setdefault(enrollment.user_id, []).append(enrollment)

    course_rows = (
        db.query(ManagedEnrollment.course_shortname)
        .join(ManagedUser, ManagedUser.id == ManagedEnrollment.user_id)
        .filter(ManagedUser.person_type == person_type, ManagedEnrollment.active.is_(True))
        .distinct()
        .order_by(ManagedEnrollment.course_shortname.asc())
        .all()
    )
    available_courses = [row[0] for row in course_rows if row[0]]

    start_row = offset + 1 if total_users else 0
    end_row = min(offset + len(users), total_users)
    pagination_pages = _pagination_window(current_page, total_pages)
    pagination_urls = {
        page_number: _managed_users_redirect_url(
            person_type,
            page=page_number,
            per_page=page_size,
            dni=dni,
            name_query=name_query,
            course_shortname=course_shortname,
            course_query=course_query,
        )
        for page_number in pagination_pages
        if page_number is not None
    }
    prev_page_url = (
        _managed_users_redirect_url(
            person_type,
            page=current_page - 1,
            per_page=page_size,
            dni=dni,
            name_query=name_query,
            course_shortname=course_shortname,
            course_query=course_query,
        )
        if current_page > 1
        else ""
    )
    next_page_url = (
        _managed_users_redirect_url(
            person_type,
            page=current_page + 1,
            per_page=page_size,
            dni=dni,
            name_query=name_query,
            course_shortname=course_shortname,
            course_query=course_query,
        )
        if current_page < total_pages
        else ""
    )

    return templates.TemplateResponse(
        "managed_users.html",
        {
            "request": request,
            "username": username,
            "person_type": person_type,
            "person_type_label": _person_type_label(person_type),
            "users": users,
            "enrollments_by_user": enrollments_by_user,
            "dni": dni,
            "name_query": name_query,
            "course_shortname": course_shortname,
            "course_query": course_query,
            "available_courses": available_courses,
            "current_page": current_page,
            "page_size": page_size,
            "total_users": total_users,
            "total_pages": total_pages,
            "start_row": start_row,
            "end_row": end_row,
            "pagination_pages": pagination_pages,
            "pagination_urls": pagination_urls,
            "prev_page_url": prev_page_url,
            "next_page_url": next_page_url,
            "message": request.query_params.get("message", ""),
            "error": request.query_params.get("error", ""),
        },
    )


@app.post("/users/{person_type}/bulk-delete")
def managed_users_bulk_delete(
    person_type: str,
    request: Request,
    selected_user_ids: list[str] = Form([]),
    action: str = Form(""),
    return_page: int = Form(1),
    return_dni: str = Form(""),
    return_name_query: str = Form(""),
    return_course_shortname: str = Form(""),
    return_course_query: str = Form(""),
    db: Session = Depends(get_db),
):
    username = require_user(request)
    if person_type not in {"student", "teacher"}:
        raise HTTPException(status_code=404, detail="Tipo invalido")
    if action not in {"azure", "moodle", "both"}:
        return RedirectResponse(
            _managed_users_redirect_url(
                person_type,
                page=return_page,
                dni=return_dni,
                name_query=return_name_query,
                course_shortname=return_course_shortname,
                course_query=return_course_query,
                error="Accion invalida.",
            ),
            status_code=303,
        )
    if not selected_user_ids:
        return RedirectResponse(
            _managed_users_redirect_url(
                person_type,
                page=return_page,
                dni=return_dni,
                name_query=return_name_query,
                course_shortname=return_course_shortname,
                course_query=return_course_query,
                error="No seleccionaste usuarios.",
            ),
            status_code=303,
        )

    ids: list[int] = []
    for value in selected_user_ids:
        try:
            ids.append(int(value))
        except ValueError:
            continue
    if not ids:
        return RedirectResponse(
            _managed_users_redirect_url(
                person_type,
                page=return_page,
                dni=return_dni,
                name_query=return_name_query,
                course_shortname=return_course_shortname,
                course_query=return_course_query,
                error="Seleccion invalida.",
            ),
            status_code=303,
        )

    users = (
        db.query(ManagedUser)
        .filter(ManagedUser.id.in_(ids), ManagedUser.person_type == person_type, ManagedUser.active.is_(True))
        .all()
    )

    graph = GraphClient(settings)
    moodle = MoodleClient(settings)
    deleted_azure = 0
    deleted_moodle = 0
    errors: list[str] = []

    for user in users:
        azure_ok = True
        moodle_ok = True

        if action in {"azure", "both"}:
            azure = graph.delete_user(user.azure_user_id or user.upn)
            if azure.status == "deleted":
                deleted_azure += 1
            elif azure.status not in {"missing", "skipped"}:
                azure_ok = False
                errors.append(f"{user.upn}: Azure {azure.message}")

        if action in {"moodle", "both"}:
            moodle_delete = moodle.delete_user_by_username(user.upn)
            if moodle_delete.status == "deleted":
                deleted_moodle += 1
            elif moodle_delete.status not in {"missing", "skipped"}:
                moodle_ok = False
                errors.append(f"{user.upn}: Moodle {moodle_delete.message}")

        if action == "both" and azure_ok and moodle_ok:
            user.active = False
            db.query(ManagedEnrollment).filter(ManagedEnrollment.user_id == user.id).update(
                {"active": False},
                synchronize_session=False,
            )

        db.add(
            OperationLog(
                actor=username,
                action=f"bulk_delete_{action}",
                status="ok" if (azure_ok and moodle_ok) else "error",
                detail=f"{user.upn} (azure_ok={azure_ok}, moodle_ok={moodle_ok})",
            )
        )

    db.commit()
    if errors:
        return RedirectResponse(
            _managed_users_redirect_url(
                person_type,
                page=return_page,
                dni=return_dni,
                name_query=return_name_query,
                course_shortname=return_course_shortname,
                course_query=return_course_query,
                error=f"Se procesaron con errores: {len(errors)}. Revisar logs.",
            ),
            status_code=303,
        )

    message = f"Eliminacion completada. Azure: {deleted_azure}, Moodle: {deleted_moodle}."
    return RedirectResponse(
        _managed_users_redirect_url(
            person_type,
            page=return_page,
            dni=return_dni,
            name_query=return_name_query,
            course_shortname=return_course_shortname,
            course_query=return_course_query,
            message=message,
        ),
        status_code=303,
    )


@app.post("/users/{person_type}/unenroll-local/{enrollment_id}")
def managed_user_unenroll_local(
    person_type: str,
    enrollment_id: int,
    request: Request,
    return_page: int = Form(1),
    return_dni: str = Form(""),
    return_name_query: str = Form(""),
    return_course_shortname: str = Form(""),
    return_course_query: str = Form(""),
    db: Session = Depends(get_db),
):
    username = require_user(request)
    if person_type not in {"student", "teacher"}:
        raise HTTPException(status_code=404, detail="Tipo invalido")

    enrollment = (
        db.query(ManagedEnrollment)
        .join(ManagedUser, ManagedUser.id == ManagedEnrollment.user_id)
        .filter(
            ManagedEnrollment.id == enrollment_id,
            ManagedUser.person_type == person_type,
            ManagedUser.active.is_(True),
            ManagedEnrollment.active.is_(True),
        )
        .first()
    )
    user = db.query(ManagedUser).filter(ManagedUser.id == enrollment.user_id).first() if enrollment else None
    if not enrollment or not user:
        return RedirectResponse(
            _managed_users_redirect_url(
                person_type,
                page=return_page,
                dni=return_dni,
                name_query=return_name_query,
                course_shortname=return_course_shortname,
                course_query=return_course_query,
                error="Matricula local no encontrada.",
            ),
            status_code=303,
        )

    _deactivate_local_enrollment_state(db, user, enrollment)
    db.add(
        OperationLog(
            actor=username,
            action="cleanup_local_enrollment",
            status="ok",
            detail=f"{user.upn}: matricula local {enrollment.course_shortname} desactivada sin tocar Moodle",
        )
    )
    db.commit()
    return RedirectResponse(
        _managed_users_redirect_url(
            person_type,
            page=return_page,
            dni=return_dni,
            name_query=return_name_query,
            course_shortname=return_course_shortname,
            course_query=return_course_query,
            message="Matricula eliminada solo de la base local.",
        ),
        status_code=303,
    )


@app.post("/users/{person_type}/unenroll/{enrollment_id}")
def managed_user_unenroll(
    person_type: str,
    enrollment_id: int,
    request: Request,
    return_page: int = Form(1),
    return_dni: str = Form(""),
    return_name_query: str = Form(""),
    return_course_shortname: str = Form(""),
    return_course_query: str = Form(""),
    db: Session = Depends(get_db),
):
    username = require_user(request)
    if person_type not in {"student", "teacher"}:
        raise HTTPException(status_code=404, detail="Tipo invalido")

    enrollment = (
        db.query(ManagedEnrollment)
        .join(ManagedUser, ManagedUser.id == ManagedEnrollment.user_id)
        .filter(
            ManagedEnrollment.id == enrollment_id,
            ManagedUser.person_type == person_type,
            ManagedUser.active.is_(True),
            ManagedEnrollment.active.is_(True),
        )
        .first()
    )
    if not enrollment:
        return RedirectResponse(
            _managed_users_redirect_url(
                person_type,
                page=return_page,
                dni=return_dni,
                name_query=return_name_query,
                course_shortname=return_course_shortname,
                course_query=return_course_query,
                error="Matricula no encontrada.",
            ),
            status_code=303,
        )

    user = db.query(ManagedUser).filter(ManagedUser.id == enrollment.user_id).first()
    if not user:
        return RedirectResponse(
            _managed_users_redirect_url(
                person_type,
                page=return_page,
                dni=return_dni,
                name_query=return_name_query,
                course_shortname=return_course_shortname,
                course_query=return_course_query,
                error="Usuario no encontrado.",
            ),
            status_code=303,
        )

    moodle = MoodleClient(settings)
    moodle_user = moodle.get_user_by_username(user.upn.lower())
    if moodle_user.status == "missing":
        _deactivate_local_user_state(db, user)
        db.add(
            OperationLog(
                actor=username,
                action="cleanup_missing_moodle_user",
                status="ok",
                detail=f"{user.upn}: usuario ausente en Moodle; estado local desactivado",
            )
        )
        db.commit()
        return RedirectResponse(
            _managed_users_redirect_url(
                person_type,
                page=return_page,
                dni=return_dni,
                name_query=return_name_query,
                course_shortname=return_course_shortname,
                course_query=return_course_query,
                message="El usuario ya no existia en Moodle. Se limpio su estado local.",
            ),
            status_code=303,
        )
    if moodle_user.status != "ok" or not isinstance(moodle_user.data, dict):
        return RedirectResponse(
            _managed_users_redirect_url(
                person_type,
                page=return_page,
                dni=return_dni,
                name_query=return_name_query,
                course_shortname=return_course_shortname,
                course_query=return_course_query,
                error=f"No se pudo consultar usuario Moodle: {moodle_user.message}",
            ),
            status_code=303,
        )

    moodle_user_id = int(moodle_user.data.get("id", 0) or 0)
    if moodle_user_id <= 0:
        return RedirectResponse(
            _managed_users_redirect_url(
                person_type,
                page=return_page,
                dni=return_dni,
                name_query=return_name_query,
                course_shortname=return_course_shortname,
                course_query=return_course_query,
                error="Moodle devolvio un usuario sin id valido.",
            ),
            status_code=303,
        )
    user.moodle_user_id = str(moodle_user_id)

    user_courses = moodle.get_user_courses(moodle_user_id)
    if user_courses.status != "ok" or not isinstance(user_courses.data, list):
        return RedirectResponse(
            _managed_users_redirect_url(
                person_type,
                page=return_page,
                dni=return_dni,
                name_query=return_name_query,
                course_shortname=return_course_shortname,
                course_query=return_course_query,
                error=f"No se pudieron consultar las matriculas Moodle: {user_courses.message}",
            ),
            status_code=303,
        )
    enrolled_shortnames = {
        str(course.get("shortname", "") or "").strip()
        for course in user_courses.data
        if isinstance(course, dict)
    }
    if enrollment.course_shortname not in enrolled_shortnames:
        _deactivate_local_enrollment_state(db, user, enrollment)
        db.add(
            OperationLog(
                actor=username,
                action="cleanup_missing_moodle_enrollment",
                status="ok",
                detail=f"{user.upn}: matricula {enrollment.course_shortname} ausente; estado local desactivado",
            )
        )
        db.commit()
        return RedirectResponse(
            _managed_users_redirect_url(
                person_type,
                page=return_page,
                dni=return_dni,
                name_query=return_name_query,
                course_shortname=return_course_shortname,
                course_query=return_course_query,
                message="La matricula ya no existia en Moodle. Se limpio el registro local.",
            ),
            status_code=303,
        )

    course_id = 0
    if enrollment.moodle_course_id.strip().isdigit():
        course_id = int(enrollment.moodle_course_id)
    else:
        course = moodle.get_course_by_shortname(enrollment.course_shortname)
        if course.status == "ok" and isinstance(course.data, dict):
            course_id = int(course.data.get("id", 0))
            enrollment.moodle_course_id = str(course_id)
            if not enrollment.course_name:
                enrollment.course_name = str(
                    course.data.get("fullname")
                    or course.data.get("displayname")
                    or course.data.get("shortname")
                    or ""
                )
        else:
            return RedirectResponse(
                _managed_users_redirect_url(
                    person_type,
                    page=return_page,
                    dni=return_dni,
                    name_query=return_name_query,
                    course_shortname=return_course_shortname,
                    course_query=return_course_query,
                    error="No se pudo resolver el curso en Moodle.",
                ),
                status_code=303,
            )

    result = moodle.unenrol_user(moodle_user_id, course_id, person_type)
    if result.status not in {"unenrolled", "missing"}:
        return RedirectResponse(
            _managed_users_redirect_url(
                person_type,
                page=return_page,
                dni=return_dni,
                name_query=return_name_query,
                course_shortname=return_course_shortname,
                course_query=return_course_query,
                error="Error al desmatricular en Moodle.",
            ),
            status_code=303,
        )

    _deactivate_local_enrollment_state(db, user, enrollment)
    db.add(
        OperationLog(
            actor=username,
            action="unenroll_course",
            status="ok",
            detail=f"{user.upn} desmatriculado de {enrollment.course_shortname}",
        )
    )
    db.commit()
    return RedirectResponse(
        _managed_users_redirect_url(
            person_type,
            page=return_page,
            dni=return_dni,
            name_query=return_name_query,
            course_shortname=return_course_shortname,
            course_query=return_course_query,
            message="Usuario desmatriculado del curso.",
        ),
        status_code=303,
    )


@app.post("/users/{person_type}/purge-local/{user_id}")
def managed_user_purge_local(
    person_type: str,
    user_id: int,
    request: Request,
    return_page: int = Form(1),
    return_dni: str = Form(""),
    return_name_query: str = Form(""),
    return_course_shortname: str = Form(""),
    return_course_query: str = Form(""),
    db: Session = Depends(get_db),
):
    username = require_user(request)
    if person_type not in {"student", "teacher"}:
        raise HTTPException(status_code=404, detail="Tipo invalido")

    user = (
        db.query(ManagedUser)
        .filter(
            ManagedUser.id == user_id,
            ManagedUser.person_type == person_type,
            ManagedUser.active.is_(True),
        )
        .first()
    )
    if not user:
        return RedirectResponse(
            _managed_users_redirect_url(
                person_type,
                page=return_page,
                dni=return_dni,
                name_query=return_name_query,
                course_shortname=return_course_shortname,
                course_query=return_course_query,
                error="Usuario no encontrado.",
            ),
            status_code=303,
        )

    active_enrollments = (
        db.query(ManagedEnrollment)
        .filter(ManagedEnrollment.user_id == user.id, ManagedEnrollment.active.is_(True))
        .count()
    )
    if active_enrollments > 0:
        return RedirectResponse(
            _managed_users_redirect_url(
                person_type,
                page=return_page,
                dni=return_dni,
                name_query=return_name_query,
                course_shortname=return_course_shortname,
                course_query=return_course_query,
                error="Primero desmatricula al usuario de sus cursos.",
            ),
            status_code=303,
        )

    _deactivate_local_user_state(db, user)
    db.add(
        OperationLog(
            actor=username,
            action="purge_local_user",
            status="ok",
            detail=f"{user.upn} removido solo de la base local",
        )
    )
    db.commit()
    return RedirectResponse(
        _managed_users_redirect_url(
            person_type,
            page=return_page,
            dni=return_dni,
            name_query=return_name_query,
            course_shortname=return_course_shortname,
            course_query=return_course_query,
            message="Usuario removido del listado local.",
        ),
        status_code=303,
    )


@app.get("/login", response_class=HTMLResponse)
def login_form(request: Request):
    azure_auth = AzureADAuth(settings)
    error_map = {
        "token": "No se pudo completar el login con Microsoft. Reintenta.",
        "userinfo": "Microsoft devolvio token, pero no se pudo leer el perfil del usuario.",
        "flow": "No se pudo iniciar el flujo de login con Microsoft.",
    }
    error_key = request.query_params.get("error", "")
    error_text = error_map.get(error_key, "")
    return templates.TemplateResponse(
        "login.html",
        {"request": request, "error": error_text, "azure_enabled": azure_auth.enabled},
    )


@app.get("/login/microsoft")
def login_microsoft(request: Request):
    azure_auth = AzureADAuth(settings)
    if not azure_auth.enabled:
        return RedirectResponse("/login", status_code=303)
    redirect_uri = f"{settings.crm_base_url}/auth/callback"
    flow = azure_auth.init_auth_code_flow(redirect_uri)
    auth_uri = flow.get("auth_uri")
    if not auth_uri:
        return RedirectResponse("/login?error=flow", status_code=303)
    request.session["msal_auth_flow"] = flow
    return RedirectResponse(url=auth_uri, status_code=303)


@app.post("/login")
def login(request: Request, username: str = Form(...), password: str = Form(...), db: Session = Depends(get_db)):
    username = username.strip().lower()
    user = db.query(LocalUser).filter(LocalUser.username == username, LocalUser.active.is_(True)).first()
    if not user or not verify_password(password, user.password_hash):
        azure_auth = AzureADAuth(settings)
        return templates.TemplateResponse(
            "login.html",
            {
                "request": request,
                "error": "Usuario o contrasena incorrectos.",
                "azure_enabled": azure_auth.enabled,
            },
            status_code=401,
        )
    request.session["username"] = user.username
    return RedirectResponse("/", status_code=303)


@app.get("/auth/callback")
def auth_callback(request: Request, code: str = None, db: Session = Depends(get_db)):
    if not code:
        return RedirectResponse("/login", status_code=303)

    azure_auth = AzureADAuth(settings)
    if not azure_auth.enabled:
        return RedirectResponse("/login", status_code=303)

    flow = request.session.pop("msal_auth_flow", None)
    token_result = azure_auth.acquire_token_by_auth_flow(flow, dict(request.query_params))
    if not token_result:
        return RedirectResponse("/login?error=token", status_code=303)

    user_info = azure_auth.get_user_info(token_result.get("access_token"))
    if not user_info:
        return RedirectResponse("/login?error=userinfo", status_code=303)

    is_valid, email_or_error = azure_auth.validate_user(user_info)
    if not is_valid:
        return templates.TemplateResponse(
            "login.html",
            {"request": request, "error": email_or_error, "azure_enabled": azure_auth.enabled},
            status_code=403,
        )

    user = azure_auth.sync_user_to_db(db, email_or_error, user_info)
    request.session["username"] = user.username
    return RedirectResponse("/", status_code=303)


@app.post("/logout")
def logout(request: Request):
    request.session.clear()
    return RedirectResponse("/login", status_code=303)


@app.get("/imports/{batch_id}", response_class=HTMLResponse)
def import_detail(batch_id: int, request: Request, db: Session = Depends(get_db)):
    username = require_user(request)
    batch = db.query(ImportBatch).filter(ImportBatch.id == batch_id).first()
    if not batch:
        raise HTTPException(status_code=404, detail="Importacion no encontrada.")
    people = (
        db.query(ImportedPerson)
        .filter(ImportedPerson.batch_id == batch_id)
        .order_by(ImportedPerson.id.asc())
        .all()
    )
    existing = [person for person in people if person.status == "exists" or person.azure_status == "exists"]
    return templates.TemplateResponse(
        "import_detail.html",
        {
            "request": request,
            "username": username,
            "batch": batch,
            "people": people,
            "existing": existing,
            "error": request.query_params.get("error", ""),
        },
    )


@app.get("/imports/{batch_id}/progress")
def import_progress(batch_id: int, request: Request, db: Session = Depends(get_db)):
    require_user(request)
    batch = db.query(ImportBatch).filter(ImportBatch.id == batch_id).first()
    if not batch:
        raise HTTPException(status_code=404, detail="Importacion no encontrada.")

    total = batch.total_rows or 0
    processed = batch.processed_rows or 0
    percent = int((processed / total) * 100) if total > 0 else 0
    if batch.run_status == "completed":
        percent = 100

    return JSONResponse(
        {
            "batch_id": batch.id,
            "run_status": batch.run_status,
            "total_rows": total,
            "processed_rows": processed,
            "success_rows": batch.success_rows or 0,
            "error_rows": batch.error_rows or 0,
            "percent": max(0, min(percent, 100)),
            "started_at": batch.started_at.isoformat() if batch.started_at else "",
            "finished_at": batch.finished_at.isoformat() if batch.finished_at else "",
        }
    )


@app.post("/imports/{batch_id}/delete")
def delete_import(batch_id: int, request: Request, db: Session = Depends(get_db)):
    username = require_user(request)
    batch = db.query(ImportBatch).filter(ImportBatch.id == batch_id).first()
    if not batch:
        return RedirectResponse("/imports", status_code=303)
    if batch.run_status in {"pending", "running"}:
        return RedirectResponse(
            f"/imports/{batch_id}?error={quote_plus('No se puede eliminar una importacion en ejecucion.')}",
            status_code=303,
        )

    db.query(ImportedPerson).filter(ImportedPerson.batch_id == batch_id).delete(synchronize_session=False)
    db.delete(batch)
    db.add(
        OperationLog(
            actor=username,
            action="delete_import",
            status="ok",
            detail=f"Importacion {batch_id} eliminada",
        )
    )
    db.commit()
    return RedirectResponse("/imports", status_code=303)


@app.get("/imports/{batch_id}/existing.csv")
def existing_csv(batch_id: int, request: Request, db: Session = Depends(get_db)):
    require_user(request)
    people = (
        db.query(ImportedPerson)
        .filter(ImportedPerson.batch_id == batch_id)
        .filter((ImportedPerson.status == "exists") | (ImportedPerson.azure_status == "exists"))
        .order_by(ImportedPerson.id.asc())
        .all()
    )
    lines = ["upn,nombre,apellido,documento,email,estado,mensaje"]
    for person in people:
        values = [
            person.upn,
            person.first_name,
            person.last_name,
            person.document,
            person.email,
            person.status,
            person.message.replace("\n", " "),
        ]
        escaped = ['"' + value.replace('"', '""') + '"' for value in values]
        lines.append(",".join(escaped))
    content = "\n".join(lines)
    return StreamingResponse(
        iter([content]),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="existentes-importacion-{batch_id}.csv"'},
    )


@app.post("/imports")
@app.post("/imports/preview")
async def preview_import(
    request: Request,
    person_type: str = Form(...),
    course_shortname: str = Form(""),
    moodle_group_name: str = Form(""),
    moodle_group_autocreate: bool = Form(False),
    moodle_group_mode: str = Form("tolerant"),
    email_subject: str = Form(""),
    email_message: str = Form(""),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    username = require_user(request)
    if person_type not in {"student", "teacher"}:
        return render_dashboard(request, db, username, "Tipo invalido.", 400)
    if not file.filename.lower().endswith(".xlsx"):
        return render_dashboard(request, db, username, "Solo se aceptan archivos .xlsx.", 400)

    content = await file.read()
    columns: list[str] = []
    sample_rows: list[dict[str, str]] = []
    try:
        columns, sample_rows = preview_people_excel(content)
    except ValueError as exc:
        return render_dashboard(request, db, username, str(exc), 400)
    except Exception as exc:
        return render_dashboard(request, db, username, f"No se pudo leer el Excel: {exc}", 400)

    PENDING_IMPORTS_DIR.mkdir(parents=True, exist_ok=True)
    token = secrets.token_urlsafe(24)
    pending_path = PENDING_IMPORTS_DIR / f"{token}.xlsx"
    pending_path.write_bytes(content)

    pending = request.session.get("pending_imports", {})
    pending[token] = {
        "path": str(pending_path),
        "filename": file.filename,
        "person_type": person_type,
        "course_shortname": course_shortname.strip(),
        "moodle_group_name": moodle_group_name.strip(),
        "moodle_group_autocreate": bool(moodle_group_autocreate),
        "moodle_group_mode": "strict" if str(moodle_group_mode).strip().lower() == "strict" else "tolerant",
        "email_subject": email_subject,
        "email_message": email_message,
    }
    request.session["pending_imports"] = pending

    return render_mapping(
        request,
        username,
        token,
        file.filename,
        person_type,
        course_shortname.strip(),
        moodle_group_name.strip(),
        bool(moodle_group_autocreate),
        "strict" if str(moodle_group_mode).strip().lower() == "strict" else "tolerant",
        columns,
        sample_rows,
        suggest_column_mapping(columns),
    )


@app.post("/imports/confirm")
async def create_import(
    request: Request,
    background_tasks: BackgroundTasks,
    token: str = Form(...),
    document_column: str = Form(""),
    first_name_column: str = Form(""),
    last_name_column: str = Form(""),
    full_name_column: str = Form(""),
    email_column: str = Form(""),
    db: Session = Depends(get_db),
):
    username = require_user(request)
    pending = request.session.get("pending_imports", {})
    item = pending.get(token)
    if not item:
        return render_dashboard(request, db, username, "La previsualizacion vencio. Subi el Excel nuevamente.", 400)

    pending_path = Path(item["path"])
    if not pending_path.exists():
        return render_dashboard(request, db, username, "No se encontro el archivo temporal. Subi el Excel nuevamente.", 400)

    content = pending_path.read_bytes()
    person_type = item["person_type"]
    course_shortname = item["course_shortname"]
    course_shortnames = parse_course_shortnames(course_shortname)
    moodle_group_name = str(item.get("moodle_group_name", "")).strip()
    moodle_group_autocreate = bool(item.get("moodle_group_autocreate", False))
    moodle_group_mode = "strict" if str(item.get("moodle_group_mode", "tolerant")).strip().lower() == "strict" else "tolerant"
    email_subject = item.get("email_subject", "")
    email_message = item.get("email_message", "")
    filename = item["filename"]
    mapping = {
        "document": document_column,
        "first_name": first_name_column,
        "last_name": last_name_column,
        "full_name": full_name_column,
        "email": email_column,
    }

    try:
        columns, sample_rows = preview_people_excel(content)
        rows = parse_people_excel(content, person_type, settings.crm_domain, mapping)
    except ValueError as exc:
        return render_mapping(
            request,
            username,
            token,
            filename,
            person_type,
            course_shortname,
            moodle_group_name,
            moodle_group_autocreate,
            moodle_group_mode,
            columns,
            sample_rows,
            mapping,
            str(exc),
            400,
        )
    except Exception as exc:
        return render_dashboard(request, db, username, f"No se pudo procesar el Excel: {exc}", 400)

    if not rows:
        return render_mapping(
            request,
            username,
            token,
            filename,
            person_type,
            course_shortname,
            moodle_group_name,
            moodle_group_autocreate,
            moodle_group_mode,
            columns,
            sample_rows,
            mapping,
            "El Excel no tiene filas para procesar.",
            400,
        )

    batch = ImportBatch(
        filename=filename,
        person_type=person_type,
        course_shortname=",".join(course_shortnames),
        moodle_group_name=moodle_group_name,
        moodle_group_autocreate=moodle_group_autocreate,
        moodle_group_mode=moodle_group_mode,
        created_by=username,
        dry_run=settings.crm_dry_run,
        email_subject=email_subject,
        email_message=email_message,
        run_status="pending",
        total_rows=len(rows),
        processed_rows=0,
        success_rows=0,
        error_rows=0,
    )
    db.add(batch)
    db.commit()
    db.refresh(batch)

    people: list[ImportedPerson] = []
    for row in rows:
        people.append(
            ImportedPerson(
                batch_id=batch.id,
                person_type=person_type,
                document=row.document,
                first_name=row.first_name,
                last_name=row.last_name,
                full_name=row.full_name,
                email=row.email,
                upn=row.upn,
            )
        )
    db.add_all(people)

    db.add(
        OperationLog(
            actor=username,
            action="import",
            status="running",
            detail=f"Importacion {batch.id}: lote iniciado con {len(rows)} filas desde {filename}",
        )
    )
    db.commit()
    pending.pop(token, None)
    request.session["pending_imports"] = pending
    pending_path.unlink(missing_ok=True)
    background_tasks.add_task(run_batch_in_background, batch.id, username, settings)
    return RedirectResponse(f"/imports/{batch.id}", status_code=303)


@app.get("/skus", response_class=HTMLResponse)
def skus(request: Request):
    username = require_user(request)
    return templates.TemplateResponse(
        "skus.html",
        {
            "request": request,
            "username": username,
            "settings": settings,
            "message": request.query_params.get("message", ""),
            "error": request.query_params.get("error", ""),
        },
    )


@app.get("/settings", response_class=HTMLResponse)
def integration_settings(request: Request):
    username = require_user(request)
    result = MoodleClient(settings).get_service_info()
    active_functions: set[str] = set()
    site_info: dict = {}

    if result.status == "ok" and isinstance(result.data, dict):
        site_info = result.data
        active_functions = {
            str(item.get("name", "")).strip()
            for item in result.data.get("functions", [])
            if isinstance(item, dict) and str(item.get("name", "")).strip()
        }

    expected_functions = {item["name"] for item in MOODLE_FUNCTION_CATALOG}
    function_rows = [
        {
            **item,
            "enabled": item["name"] in active_functions if result.status == "ok" else None,
        }
        for item in MOODLE_FUNCTION_CATALOG
    ]
    missing_functions = sorted(expected_functions - active_functions) if result.status == "ok" else []
    extra_functions = sorted(active_functions - expected_functions)
    audit_setup_required = (
        result.status == "error" and "control de acceso" in result.message.lower()
    )

    return templates.TemplateResponse(
        "settings.html",
        {
            "request": request,
            "username": username,
            "settings": settings,
            "result": result,
            "site_info": site_info,
            "function_rows": function_rows,
            "active_count": len(active_functions),
            "verified_count": len(active_functions & expected_functions),
            "expected_count": len(expected_functions),
            "missing_functions": missing_functions,
            "extra_functions": extra_functions,
            "audit_setup_required": audit_setup_required,
            "message": request.query_params.get("message", ""),
            "error": request.query_params.get("error", ""),
        },
    )


@app.post("/settings/matriculation")
def update_matriculation_settings(
    request: Request,
    course_alert_threshold: int = Form(5),
):
    require_user(request)
    threshold = max(1, min(int(course_alert_threshold or 5), 30))
    update_env_file({"MATRICULATION_COURSE_ALERT_THRESHOLD": str(threshold)})
    get_settings.cache_clear()
    global settings
    settings = get_settings()
    templates.env.globals["settings"] = settings
    return RedirectResponse(
        f"/settings?message=Umbral+de+materias+guardado+en+{threshold}",
        status_code=303,
    )


@app.get("/mail", response_class=HTMLResponse)
def mail_settings(request: Request):
    username = require_user(request)
    return templates.TemplateResponse(
        "mail_settings.html",
        {
            "request": request,
            "username": username,
            "settings": settings,
            "message": request.query_params.get("message", ""),
            "error": request.query_params.get("error", ""),
        },
    )


@app.post("/mail")
def update_mail_settings(
    request: Request,
    smtp_host: str = Form(""),
    smtp_port: int = Form(587),
    smtp_username: str = Form(""),
    smtp_password: str = Form(""),
    smtp_from_email: str = Form(""),
    smtp_from_name: str = Form(""),
    crm_batch_workers: int = Form(5),
    smtp_subject: str = Form(""),
    smtp_message_template: str = Form(""),
):
    require_user(request)
    normalized_template = smtp_message_template.replace("\r\n", "\n").replace("\r", "\n")
    update_env_file(
        {
            "SMTP_HOST": smtp_host.strip(),
            "SMTP_PORT": str(smtp_port),
            "SMTP_USERNAME": smtp_username.strip(),
            "SMTP_PASSWORD": smtp_password,
            "SMTP_FROM_EMAIL": smtp_from_email.strip(),
            "SMTP_FROM_NAME": smtp_from_name.strip(),
            "CRM_BATCH_WORKERS": str(max(1, min(crm_batch_workers, 20))),
            "SMTP_SUBJECT": smtp_subject.strip(),
            "SMTP_MESSAGE_TEMPLATE": normalized_template,
        }
    )
    get_settings.cache_clear()
    global settings
    settings = get_settings()
    templates.env.globals["settings"] = settings
    return RedirectResponse("/mail?message=Configuracion+guardada", status_code=303)


@app.post("/skus")
def update_azure_skus(
    request: Request,
    azure_student_sku_id: str = Form(""),
    azure_teacher_sku_id: str = Form(""),
):
    username = require_user(request)
    update_env_file(
        {
            "AZURE_STUDENT_SKU_ID": azure_student_sku_id.strip(),
            "AZURE_TEACHER_SKU_ID": azure_teacher_sku_id.strip(),
        }
    )
    get_settings.cache_clear()
    global settings
    settings = get_settings()
    return RedirectResponse("/skus?message=Licencias+guardadas", status_code=303)
