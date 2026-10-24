"""Read-only SIAL/Moodle/Microsoft audit; persists only to its own local snapshot."""

import argparse
import json
import os
from pathlib import Path
import sqlite3
import time
from datetime import datetime, timezone
from urllib.parse import quote, urlparse

import httpx
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.config import get_settings
from app.services.graph import GraphClient
from app.services.moodle import MoodleClient


def now():
    return datetime.now(timezone.utc).isoformat()


def progress(phase, **values):
    print(json.dumps({"fase": phase, **values}, ensure_ascii=True), flush=True)


def local_date(value):
    if value in (None, "", 0, "0"):
        return None
    try:
        dt = (datetime.fromtimestamp(float(value), timezone.utc)
              if isinstance(value, (float, int)) else datetime.fromisoformat(str(value).replace("Z", "+00:00")))
        return dt.astimezone().strftime("%Y-%m-%d %H:%M:%S %z")
    except (ValueError, TypeError, OverflowError):
        return str(value)


def open_db(folder):
    folder.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(folder / "auditoria.sqlite3")
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA foreign_keys=ON")
    db.executescript("""
        CREATE TABLE IF NOT EXISTS metadatos (clave TEXT PRIMARY KEY, valor TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS personas (
            dni TEXT PRIMARY KEY, upn TEXT NOT NULL UNIQUE, estados_sial TEXT NOT NULL,
            nombre TEXT, apellido TEXT, nombre_fuente TEXT,
            moodle_estado TEXT DEFAULT 'pendiente', moodle_id INTEGER,
            moodle_primer_acceso INTEGER, moodle_ultimo_acceso INTEGER,
            moodle_suspendido INTEGER, moodle_cursos_estado TEXT DEFAULT 'pendiente',
            moodle_cursos_cantidad INTEGER, moodle_consultado TEXT,
            microsoft_estado TEXT DEFAULT 'pendiente', microsoft_id TEXT,
            microsoft_habilitada INTEGER, microsoft_creada TEXT,
            microsoft_licencias INTEGER, microsoft_consultado TEXT,
            teams_estado TEXT DEFAULT 'pendiente', teams_cantidad INTEGER,
            teams_consultado TEXT,
            FOREIGN KEY (dni) REFERENCES revisiones(dni) DEFERRABLE INITIALLY DEFERRED
        );
        CREATE TABLE IF NOT EXISTS revisiones (
            dni TEXT PRIMARY KEY, decision_moodle TEXT DEFAULT '',
            decision_microsoft TEXT DEFAULT '', responsable TEXT DEFAULT '',
            fecha_revision TEXT DEFAULT '', observaciones TEXT DEFAULT ''
        );
        CREATE TABLE IF NOT EXISTS sial_inscripciones (
            dni TEXT NOT NULL REFERENCES personas(dni), estado TEXT NOT NULL,
            carrera_id TEXT, carrera TEXT, modalidad TEXT, comisiones TEXT,
            fecha_inscripcion TEXT, consultado TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS moodle_cursos (
            dni TEXT NOT NULL REFERENCES personas(dni), curso_id INTEGER NOT NULL,
            codigo TEXT, nombre TEXT, visible INTEGER, inicio INTEGER, fin INTEGER,
            ultimo_acceso INTEGER, consultado TEXT NOT NULL, PRIMARY KEY(dni,curso_id)
        );
        CREATE TABLE IF NOT EXISTS equipos (
            equipo_id TEXT PRIMARY KEY, nombre TEXT, archivado INTEGER, tenant_id TEXT
        );
        CREATE TABLE IF NOT EXISTS teams_pertenencias (
            dni TEXT NOT NULL REFERENCES personas(dni), equipo_id TEXT NOT NULL REFERENCES equipos(equipo_id),
            consultado TEXT NOT NULL, PRIMARY KEY(dni,equipo_id)
        );
        CREATE TABLE IF NOT EXISTS evidencias (
            origen TEXT NOT NULL, referencia TEXT NOT NULL, consultado TEXT NOT NULL,
            datos_json TEXT NOT NULL, PRIMARY KEY(origen,referencia)
        );
        CREATE TABLE IF NOT EXISTS incidencias (
            id INTEGER PRIMARY KEY, origen TEXT NOT NULL, dni TEXT, codigo TEXT,
            detalle TEXT, consultado TEXT NOT NULL
        );
    """)
    return db


def meta(db, key, value):
    db.execute("INSERT OR REPLACE INTO metadatos VALUES (?,?)", (key, str(value)))
    db.commit()


def evidence(db, source, ref, value):
    db.execute("INSERT OR REPLACE INTO evidencias VALUES (?,?,?,?)",
               (source, ref, now(), json.dumps(value, ensure_ascii=False)))


def incident(db, source, dni, code, detail):
    db.execute("INSERT INTO incidencias(origen,dni,codigo,detalle,consultado) VALUES (?,?,?,?,?)",
               (source, dni, str(code), detail, now()))


def update(db, dni, **fields):
    # Column names only come from this script; values are always parameters.
    db.execute("UPDATE personas SET " + ",".join(f"{k}=?" for k in fields) + " WHERE dni=?",
               (*fields.values(), dni))


def retry_get(client, url, **kwargs):
    for attempt in range(4):
        try:
            r = client.get(url, **kwargs)
        except httpx.HTTPError:
            if attempt == 3:
                raise RuntimeError("Fallo de transporte en consulta de lectura") from None
            time.sleep(2 ** attempt)
            continue
        if r.status_code not in (429, 500, 502, 503, 504) or attempt == 3:
            return r
        try:
            delay = float(r.headers.get("Retry-After", 2 ** attempt))
        except ValueError:
            delay = 2 ** attempt
        if delay > 60:
            raise RuntimeError("Servidor pide esperar mas de 60 segundos; reintentar luego")
        time.sleep(max(1, delay))


def sial_lists(since, settings=None, progress=progress, state_order=("aspirantes", "preinscriptos", "ingresantes")):
    if len(state_order) != 3 or set(state_order) != {"preinscriptos", "aspirantes", "ingresantes"}:
        raise ValueError("Se requieren los tres estados SIAL.")
    settings = settings or get_settings()
    auth = (settings.sial_username or os.environ.get("SIAL_AUDIT_USER", ""),
            settings.sial_password or os.environ.get("SIAL_AUDIT_PASSWORD", ""))
    base = settings.sial_base_url.rstrip("/")
    parsed = urlparse(base)
    if parsed.scheme != "https" or parsed.username or not all(auth):
        raise RuntimeError("Configurar URL HTTPS y credenciales de SIAL en Configuracion / Auditoria.")
    all_states = {}
    with httpx.Client(timeout=60, auth=auth) as client:
        for state in state_order:
            r = retry_get(client, f"{base}/{state}/listas/{since}")
            if r.status_code != 200:
                raise RuntimeError(f"SIAL {state}: HTTP {r.status_code}")
            data = r.json()
            if not isinstance(data, list) or any(not isinstance(p, dict) or not str(p.get("nrodoc", "")).strip().isdigit() for p in data):
                raise RuntimeError(f"Estructura SIAL inesperada: {state}")
            all_states[state] = data
            progress("SIAL", estado=state, filas=len(data))
    return all_states


def fetch_sial(db, since, settings=None, progress=progress, include_ingresantes=False):
    if db.execute("SELECT 1 FROM metadatos WHERE clave='sial_completo'").fetchone():
        return
    settings = settings or get_settings()
    all_states = sial_lists(since, settings, progress)
    for state, rows in all_states.items():
        meta(db, f"sial_total_{state}", len(rows))
    people = {}
    for state in (("aspirantes", "preinscriptos", "ingresantes") if include_ingresantes else ("aspirantes", "preinscriptos")):
        for row in all_states[state]:
            dni = str(row["nrodoc"]).strip()
            if not dni.isdigit():
                raise RuntimeError("SIAL devuelve documento no numerico; requiere revision de identidad")
            people.setdefault(dni, {})[state] = row
    # Detect another current condition without expanding the requested population.
    for row in all_states["ingresantes"]:
        dni = str(row["nrodoc"]).strip()
        if dni in people:
            people[dni]["ingresantes"] = row
    with db:
        for dni, states in people.items():
            db.execute("INSERT INTO revisiones(dni) VALUES (?)", (dni,))
            db.execute("INSERT INTO personas(dni,upn,estados_sial) VALUES (?,?,?)",
                       (dni, f"a{dni}@{settings.crm_domain}", "; ".join(states)))
            for state, row in states.items():
                evidence(db, "SIAL_" + state, dni, row)
                for c in row.get("carreras", []) or []:
                    commissions = "; ".join(str(x.get("nombre_comision", "")) for x in c.get("comisiones", []) or [])
                    db.execute("INSERT INTO sial_inscripciones VALUES (?,?,?,?,?,?,?,?)", (
                        dni, state, str(c.get("id_carrera", "")), c.get("nombre_carrera", ""),
                        str(c.get("modalidad", "")), commissions, str(c.get("fecha_inscri", "")), now()))
        db.execute("INSERT INTO metadatos VALUES ('sial_completo','1')")
    progress("poblacion", personas=len(people))


def fetch_moodle(db, settings=None, progress=progress):
    client = MoodleClient(settings or get_settings())
    people = db.execute("SELECT * FROM personas WHERE moodle_estado='pendiente' ORDER BY dni").fetchall()
    for start in range(0, len(people), 50):
        batch = people[start:start + 50]
        result = client.get_users_by_usernames([p["upn"] for p in batch])
        if result.status != "ok" or not isinstance(result.data, list):
            raise RuntimeError("Fallo la busqueda por lotes de usuarios Moodle; continuar luego")
        found = {str(u.get("username", "")).lower(): u for u in result.data}
        with db:
            for p in batch:
                u = found.get(p["upn"].lower())
                if u is None:
                    update(db, p["dni"], moodle_estado="sin_usuario", moodle_cursos_estado="sin_usuario",
                           moodle_consultado=now())
                    continue
                evidence(db, "Moodle_usuario", p["dni"], {k: u.get(k) for k in (
                    "id", "username", "firstname", "lastname", "firstaccess", "lastaccess", "suspended", "auth")})
                update(db, p["dni"], moodle_estado="encontrado", moodle_id=u["id"],
                       nombre=u.get("firstname"), apellido=u.get("lastname"), nombre_fuente="Moodle",
                       moodle_primer_acceso=u.get("firstaccess"), moodle_ultimo_acceso=u.get("lastaccess"),
                       moodle_suspendido=u.get("suspended"), moodle_consultado=now())
        progress("usuarios_moodle", procesados=min(start + 50, len(people)), total=len(people))
    people = db.execute("SELECT * FROM personas WHERE moodle_cursos_estado IN ('pendiente','error') AND moodle_id IS NOT NULL ORDER BY dni").fetchall()
    for i, p in enumerate(people, 1):
        result = client.get_user_courses(p["moodle_id"])
        with db:
            if result.status != "ok" or not isinstance(result.data, list):
                update(db, p["dni"], moodle_cursos_estado="error", moodle_consultado=now())
                incident(db, "Moodle_cursos", p["dni"], result.status, "No se obtuvo lista valida de cursos")
            else:
                db.execute("DELETE FROM moodle_cursos WHERE dni=?", (p["dni"],))
                for c in result.data:
                    db.execute("INSERT INTO moodle_cursos VALUES (?,?,?,?,?,?,?,?,?)", (
                        p["dni"], c["id"], c.get("shortname"), c.get("fullname"), c.get("visible"),
                        c.get("startdate"), c.get("enddate"), c.get("lastaccess"), now()))
                evidence(db, "Moodle_cursos", p["dni"], result.data)
                update(db, p["dni"], moodle_cursos_estado="consultado", moodle_cursos_cantidad=len(result.data), moodle_consultado=now())
        if i % 50 == 0 or i == len(people):
            progress("cursos_moodle", procesados=i, total=len(people))


class GraphReadOnly:
    def __init__(self, settings=None):
        self.graph = GraphClient(settings or get_settings())
        self.client = httpx.Client(timeout=45, headers=self.graph._headers())

    def url(self, path):
        url = path if path.startswith("https://") else self.graph.base_url + path
        parsed = urlparse(url)
        if parsed.scheme != "https" or parsed.hostname != "graph.microsoft.com":
            raise ValueError("Destino Graph de lectura inesperado")
        return url

    def get(self, path):
        return retry_get(self.client, self.url(path))

    def batch(self, paths):
        # POST is only the Graph batch envelope: every enclosed request is GET.
        requests = [{"id": str(i), "method": "GET", "url": path} for i, path in enumerate(paths)]
        r = self.client.post(self.graph.base_url + "/$batch", json={"requests": requests})
        if r.status_code != 200:
            raise RuntimeError(f"Lote de lectura Graph: HTTP {r.status_code}")
        data = {str(x["id"]): x for x in r.json().get("responses", [])}
        output = []
        for i, path in enumerate(paths):
            row = data.get(str(i), {})
            if not row or row.get("status") in (429, 500, 502, 503, 504):
                try:
                    delay = float(row.get("headers", {}).get("Retry-After", 2))
                except ValueError:
                    delay = 2
                if delay > 60:
                    output.append((429, {"error": {"code": "RetryLater"}}))
                    continue
                time.sleep(max(1, delay))
                retry = self.get(path)
                output.append((retry.status_code, retry.json()))
            else:
                output.append((row["status"], row.get("body", {})))
        return output

    def list_pages(self, first):
        rows = []
        body = first
        visited = set()
        while True:
            if not isinstance(body.get("value"), list):
                raise RuntimeError("Lista Graph invalida")
            rows.extend(body["value"])
            link = body.get("@odata.nextLink")
            if not link:
                return rows
            if link in visited:
                raise RuntimeError("Paginacion Graph repetida")
            visited.add(link)
            r = self.get(link)
            if r.status_code != 200:
                raise RuntimeError(f"Paginacion Graph HTTP {r.status_code}")
            body = r.json()


def fetch_graph(db, settings=None, progress=progress, probe_permissions=True):
    reader = GraphReadOnly(settings)
    select = "id,userPrincipalName,displayName,givenName,surname,accountEnabled,createdDateTime,assignedLicenses"
    try:
        people = db.execute("SELECT * FROM personas WHERE microsoft_estado IN ('pendiente','error') ORDER BY dni").fetchall()
        for start in range(0, len(people), 20):
            batch = people[start:start + 20]
            paths = [f"/users/{quote(p['upn'])}?$select={select}" for p in batch]
            results = reader.batch(paths)
            with db:
                for p, (status, body) in zip(batch, results):
                    if status == 404:
                        update(db, p["dni"], microsoft_estado="sin_usuario", teams_estado="sin_usuario", microsoft_consultado=now())
                    elif status != 200 or not body.get("id"):
                        update(db, p["dni"], microsoft_estado="error", microsoft_consultado=now())
                        incident(db, "Microsoft_usuario", p["dni"], status, body.get("error", {}).get("code", "Respuesta invalida"))
                    else:
                        evidence(db, "Microsoft_usuario", p["dni"], body)
                        fields = dict(microsoft_estado="encontrado", microsoft_id=body["id"],
                                      microsoft_habilitada=body.get("accountEnabled"), microsoft_creada=body.get("createdDateTime"),
                                      microsoft_licencias=len(body.get("assignedLicenses", [])), microsoft_consultado=now())
                        if not p["nombre"] and not p["apellido"]:
                            fields.update(nombre=body.get("givenName") or body.get("displayName"), apellido=body.get("surname"), nombre_fuente="Microsoft")
                        update(db, p["dni"], **fields)
            if start % 100 == 0 or start + 20 >= len(people):
                progress("usuarios_microsoft", procesados=min(start + 20, len(people)), total=len(people))
            time.sleep(0.5)
        people = db.execute("SELECT * FROM personas WHERE microsoft_id IS NOT NULL AND teams_estado IN ('pendiente','error') ORDER BY dni").fetchall()
        for start in range(0, len(people), 20):
            batch = people[start:start + 20]
            results = reader.batch([f"/users/{p['microsoft_id']}/joinedTeams" for p in batch])
            for p, (status, body) in zip(batch, results):
                with db:
                    if status != 200:
                        update(db, p["dni"], teams_estado="error", teams_consultado=now())
                        incident(db, "Teams_equipos", p["dni"], status, body.get("error", {}).get("code", "Respuesta invalida"))
                        continue
                    try:
                        teams = reader.list_pages(body)
                        if any(not t.get("id") for t in teams):
                            raise RuntimeError("Equipo sin identificador")
                    except RuntimeError:
                        update(db, p["dni"], teams_estado="error", teams_consultado=now())
                        incident(db, "Teams_equipos", p["dni"], "paginacion", "Consulta incompleta")
                        continue
                    db.execute("DELETE FROM teams_pertenencias WHERE dni=?", (p["dni"],))
                    for t in teams:
                        db.execute("""INSERT INTO equipos VALUES (?,?,?,?) ON CONFLICT(equipo_id)
                                   DO UPDATE SET nombre=excluded.nombre,archivado=excluded.archivado,tenant_id=excluded.tenant_id""",
                                   (t["id"], t.get("displayName"), t.get("isArchived"), t.get("tenantId")))
                        db.execute("INSERT OR IGNORE INTO teams_pertenencias VALUES (?,?,?)", (p["dni"], t["id"], now()))
                    evidence(db, "Teams_equipos", p["dni"], [{k: t.get(k) for k in ("id", "displayName", "isArchived", "tenantId")} for t in teams])
                    update(db, p["dni"], teams_estado="consultado", teams_cantidad=len({t["id"] for t in teams}), teams_consultado=now())
            if start % 100 == 0 or start + 20 >= len(people):
                progress("equipos_teams", procesados=min(start + 20, len(people)), total=len(people))
            time.sleep(0.5)
        sample = db.execute("SELECT microsoft_id FROM personas WHERE microsoft_id IS NOT NULL LIMIT 1").fetchone()
        if sample and probe_permissions:
            for label, endpoint in (
                ("microsoft_ultimo_ingreso", f"/users/{sample[0]}?$select=id,signInActivity"),
                ("teams_canales_compartidos", f"/users/{sample[0]}/teamwork/associatedTeams"),
                ("teams_informe_actividad", "/reports/getTeamsUserActivityUserDetail(period='D180')"),
            ):
                r = reader.get(endpoint)
                meta(db, label, f"HTTP {r.status_code}; no incluido en este informe")
                if r.status_code != 200:
                    with db:
                        incident(db, label, None, r.status_code, "No disponible con los permisos actuales")
    finally:
        reader.client.close()


def fetch_missing_names(db, settings=None):
    people = db.execute("SELECT dni FROM personas WHERE nombre IS NULL AND apellido IS NULL").fetchall()
    if not people:
        return
    settings = settings or get_settings()
    auth = (settings.sial_username or os.environ.get("SIAL_AUDIT_USER", ""),
            settings.sial_password or os.environ.get("SIAL_AUDIT_PASSWORD", ""))
    with httpx.Client(timeout=60, auth=auth) as client:
        for p in people:
            r = retry_get(client, settings.sial_base_url.rstrip("/") + "/alumnos/datospersonales/" + p["dni"])
            with db:
                data = r.json() if r.status_code == 200 else None
                if not isinstance(data, list) or len(data) != 1 or str(data[0].get("nrodoc", "")).strip() != p["dni"]:
                    incident(db, "SIAL_nombre", p["dni"], r.status_code, "No se obtuvo una identidad univoca")
                    continue
                person = data[0]
                update(db, p["dni"], nombre=person.get("nombre"), apellido=person.get("apellido"), nombre_fuente="SIAL")
                evidence(db, "SIAL_nombre", p["dni"], {k: person.get(k) for k in ("nrodoc", "nombre", "apellido")})


def recommendation(p):
    if "ingresantes" in p["estados_sial"]:
        return "Revisar: tambien aparece como ingresante"
    if p["moodle_cursos_cantidad"] and p["moodle_cursos_cantidad"] > 0:
        return "Conservar: tiene cursos Moodle"
    if p["moodle_estado"] in ("pendiente", "error") or p["moodle_cursos_estado"] in ("pendiente", "error"):
        return "Consulta Moodle incompleta"
    if p["microsoft_estado"] in ("pendiente", "error") or p["teams_estado"] in ("pendiente", "error"):
        return "Consulta Microsoft incompleta"
    if p["teams_cantidad"] and p["teams_cantidad"] > 0:
        return "Revisar equipos Teams; sin cursos Moodle"
    return "Revisar individualmente; sin cursos ni equipos directos detectados"


def add_sheet(wb, title, headers, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(list(row))
    # Imported values stay literal, including strings beginning with '='.
    for row in ws:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="165D85")
    ws.freeze_panes = "D2" if len(headers) > 3 else "A2"
    ws.sheet_view.zoomScale = 85
    if ws.max_row > 1:
        table = Table(displayName=f"Tabla{len(wb.worksheets)}", ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 24)
    return ws


def export(db, folder):
    people = db.execute("SELECT * FROM personas ORDER BY apellido,nombre,dni").fetchall()
    summary = []
    for state in ("aspirantes", "preinscriptos"):
        subset = [p for p in people if state in p["estados_sial"].split("; ")]
        zero = [p for p in subset if p["moodle_cursos_estado"] == "consultado" and p["moodle_cursos_cantidad"] == 0]
        summary.append([state, len(subset), sum(p["moodle_cursos_cantidad"] is not None and p["moodle_cursos_cantidad"] > 0 for p in subset),
                        len(zero), sum(p["moodle_estado"] == "sin_usuario" for p in subset),
                        sum(p["teams_cantidad"] is not None and p["teams_cantidad"] > 0 for p in subset),
                        sum(p["teams_estado"] == "consultado" and p["teams_cantidad"] == 0 for p in subset),
                        sum(p["microsoft_estado"] == "sin_usuario" for p in subset),
                        sum(p["teams_cantidad"] is not None and p["teams_cantidad"] > 0 for p in zero),
                        sum(p["teams_estado"] == "consultado" and p["teams_cantidad"] == 0 for p in zero)])
    wb = Workbook()
    wb.remove(wb.active)
    add_sheet(wb, "Resumen", ["Estado SIAL", "Personas", "Con cursos Moodle", "Usuario Moodle sin cursos", "Sin usuario Moodle", "Con equipos directos Teams", "Microsoft sin equipos directos", "Sin cuenta Microsoft", "Sin cursos pero con equipos", "Sin cursos ni equipos directos"], summary)
    metadata = dict(db.execute("SELECT clave,valor FROM metadatos"))
    notes = [
        ("Consulta iniciada", local_date(metadata.get("inicio"))),
        ("Consulta finalizada", local_date(metadata.get("fin"))),
        ("Campania desde", metadata.get("desde")),
        ("Poblacion", "DNI unicos de aspirantes y preinscriptos SIAL V2 de la campania consultada. Ingresantes solo se cruza para detectar superposiciones."),
        ("Identidad", "DNI -> aDNI@eco.unrc.edu.ar. Otros nombres de usuario no se buscaron."),
        ("Momento", "Foto del estado al consultar cada API; no es una transaccion simultanea entre sistemas."),
        ("Moodle cursos", "Cursos devueltos por core_enrol_get_users_courses con el token del CRM. Incluye cursos de otros periodos devueltos por la API. No mide actividad."),
        ("Moodle acceso", "firstaccess y lastaccess a nivel de sitio. Cero: no registra acceso; vacio: dato no disponible."),
        ("Equipos directos", "Membresia directa devuelta por joinedTeams; se recorren todas las paginas. Un equipo no archivado no demuestra actividad reciente."),
        ("Canales compartidos", metadata.get("teams_canales_compartidos", "No consultado")),
        ("Inicio sesion Microsoft", metadata.get("microsoft_ultimo_ingreso", "No consultado")),
        ("Actividad Teams", metadata.get("teams_informe_actividad", "No consultado")),
        ("Campos sin dato", "No se interpreta un error, permiso insuficiente o dato nulo como cero actividad o ausencia de matricula."),
        ("Revision", "Sugerencias para revision humana. No autoriza bajas. Las decisiones Moodle y Microsoft son independientes; documentar responsable y fundamento."),
        ("Uso Microsoft", "Este informe no evalua chats, archivos, OneDrive, correo ni actividad historica completa."),
        ("Datos guardados", "SQLite local separado del CRM con personas, inscripciones SIAL, cursos, equipos, pertenencias, evidencia, metadatos e incidencias."),
        ("Documentacion Teams", "https://learn.microsoft.com/en-us/graph/api/user-list-joinedteams?view=graph-rest-1.0"),
        ("Documentacion actividad", "https://learn.microsoft.com/en-us/graph/api/reportroot-getteamsuseractivityuserdetail?view=graph-rest-1.0"),
    ]
    add_sheet(wb, "Criterios_y_alcance", ["Dato", "Detalle"], notes, {1:32, 2:120})
    headers = ["DNI", "Apellido", "Nombre", "Condicion SIAL", "Usuario institucional", "Carreras SIAL", "Modalidades SIAL", "Comisiones SIAL", "Estado Moodle", "Cantidad cursos Moodle", "Cursos Moodle", "Primer acceso Moodle", "Ultimo acceso Moodle", "Moodle suspendido", "Estado Microsoft", "Microsoft habilitada", "Creacion Microsoft", "Cantidad licencias Microsoft", "Estado consulta Teams", "Cantidad equipos directos", "Equipos Teams", "Equipos archivados", "Equipos no archivados", "Equipos archivo desconocido", "Actividad Microsoft", "Orientacion revision", "Decision Moodle", "Decision Microsoft", "Responsable", "Fecha revision", "Observaciones", "Consulta Moodle", "Consulta Microsoft", "Consulta equipos", "Fuente nombre"]
    rows = []
    no_courses_rows = []
    no_both_rows = []
    for p in people:
        sial = db.execute("SELECT * FROM sial_inscripciones WHERE dni=?", (p["dni"],)).fetchall()
        courses = db.execute("SELECT * FROM moodle_cursos WHERE dni=? ORDER BY codigo", (p["dni"],)).fetchall()
        teams = db.execute("SELECT e.* FROM equipos e JOIN teams_pertenencias t USING(equipo_id) WHERE t.dni=? ORDER BY e.nombre", (p["dni"],)).fetchall()
        review = db.execute("SELECT * FROM revisiones WHERE dni=?", (p["dni"],)).fetchone()
        def access(value):
            return "Sin acceso registrado" if value == 0 else local_date(value)
        row = [p["dni"], p["apellido"], p["nombre"], p["estados_sial"], p["upn"],
               "; ".join(sorted({s["carrera"] for s in sial if s["carrera"]})),
               "; ".join(sorted({{"1":"Presencial", "2":"A distancia"}.get(s["modalidad"], s["modalidad"]) for s in sial})),
               "; ".join(sorted({s["comisiones"] for s in sial if s["comisiones"]})),
               p["moodle_estado"], p["moodle_cursos_cantidad"],
               "\n".join(f"{c['codigo']}: {c['nombre']}" for c in courses),
               access(p["moodle_primer_acceso"]), access(p["moodle_ultimo_acceso"]), p["moodle_suspendido"],
               p["microsoft_estado"], p["microsoft_habilitada"], local_date(p["microsoft_creada"]), p["microsoft_licencias"],
               p["teams_estado"], p["teams_cantidad"], "\n".join(t["nombre"] or t["equipo_id"] for t in teams),
               sum(t["archivado"] == 1 for t in teams) if p["teams_estado"] == "consultado" else None,
               sum(t["archivado"] == 0 for t in teams) if p["teams_estado"] == "consultado" else None,
               sum(t["archivado"] is None for t in teams) if p["teams_estado"] == "consultado" else None,
               "No disponible; permisos insuficientes", recommendation(p),
               review["decision_moodle"], review["decision_microsoft"], review["responsable"], review["fecha_revision"], review["observaciones"],
               local_date(p["moodle_consultado"]), local_date(p["microsoft_consultado"]), local_date(p["teams_consultado"]), p["nombre_fuente"]]
        rows.append(row)
        if p["moodle_cursos_estado"] == "consultado" and p["moodle_cursos_cantidad"] == 0:
            no_courses_rows.append(row)
            if p["teams_estado"] == "consultado" and p["teams_cantidad"] == 0:
                no_both_rows.append(row)
    widths = {1:15, 2:26, 3:28, 5:37, 6:55, 11:65, 21:65, 25:35, 26:55, 31:50}
    add_sheet(wb, "Personas", headers, rows, widths)
    add_sheet(wb, "Sin_cursos_Moodle", headers, no_courses_rows, widths)
    add_sheet(wb, "Sin_cursos_ni_equipos", headers, no_both_rows, widths)
    add_sheet(wb, "Inscripciones_SIAL", ["DNI", "Estado", "Carrera ID", "Carrera", "Modalidad codigo", "Comisiones", "Fecha inscripcion", "Consultado UTC"],
              db.execute("SELECT * FROM sial_inscripciones ORDER BY dni,estado,carrera_id"), {4:60, 6:35})
    add_sheet(wb, "Cursos_Moodle", ["DNI", "Curso ID", "Codigo", "Curso", "Visible", "Inicio", "Fin", "Ultimo acceso curso", "Consultado"],
              ([r["dni"],r["curso_id"],r["codigo"],r["nombre"],r["visible"],local_date(r["inicio"]),local_date(r["fin"]),local_date(r["ultimo_acceso"]),local_date(r["consultado"])]
               for r in db.execute("SELECT * FROM moodle_cursos ORDER BY dni,codigo")), {4:70})
    add_sheet(wb, "Equipos_Teams", ["DNI", "Equipo ID", "Equipo", "Archivado", "Consultado"],
              ([r["dni"],r["equipo_id"],r["nombre"],"Desconocido" if r["archivado"] is None else "Si" if r["archivado"] else "No",local_date(r["consultado"])]
               for r in db.execute("SELECT t.*,e.nombre,e.archivado FROM teams_pertenencias t JOIN equipos e USING(equipo_id) ORDER BY dni,nombre")), {2:40, 3:75})
    add_sheet(wb, "Incidencias", ["ID", "Origen", "DNI", "Codigo", "Detalle", "Consulta UTC"], db.execute("SELECT * FROM incidencias"), {2:36,5:75})
    output = folder / "auditoria_ingreso.xlsx"
    suffix = 2
    while output.exists():
        output = folder / f"auditoria_ingreso_{suffix}.xlsx"
        suffix += 1
    wb.save(output)
    progress("informe", excel=str(output), sqlite=str(folder / "auditoria.sqlite3"), resumen=summary)
    return output


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--desde", default="202512010000")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--export-only", action="store_true")
    args = parser.parse_args()
    db = open_db(args.output)
    if not db.execute("SELECT 1 FROM metadatos WHERE clave='inicio'").fetchone():
        meta(db, "inicio", now())
        meta(db, "desde", args.desde)
    elif dict(db.execute("SELECT clave,valor FROM metadatos"))["desde"] != args.desde:
        raise RuntimeError("La carpeta corresponde a otra campania")
    try:
        if not args.export_only:
            fetch_sial(db, args.desde)
            fetch_moodle(db)
            fetch_graph(db)
            fetch_missing_names(db)
            meta(db, "fin", now())
        export(db, args.output)
    except (RuntimeError, httpx.HTTPError):
        meta(db, "consulta_incompleta", now())
        export(db, args.output)
        raise
    finally:
        db.close()


if __name__ == "__main__":
    main()
